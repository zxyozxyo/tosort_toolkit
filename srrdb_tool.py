"""
srrdb.com Scene RAR Rebuilder
Downloads SRR files, reconstructs scene RARs, extracts NFO/SFV, creates samples.
Requires: pip install pyReScene
"""

# Python 3.13 removed nntplib; stub it out before rescene imports it
import sys as _sys
import types as _types
if "nntplib" not in _sys.modules:
    _sys.modules["nntplib"] = _types.ModuleType("nntplib")

import os
import sys
import json
import shutil
import re
import time
import tempfile
import threading
import subprocess
import zlib
import sqlite3
import hashlib
from pathlib import Path
from urllib.request import urlopen, Request
from urllib.error import HTTPError
from urllib.parse import quote

# rescene expects rar executables named YYYY-MM-DD_rar<MAJOR><MINOR>[b<N>].exe
# The date is the WinRAR release date and is used to sort versions for reconstruction.
_WINRAR_DATES: dict[str, str] = {
    "420": "2012-06-09",
    "411": "2012-03-15",
    "410": "2012-01-17",
    "401": "2011-06-14",
    "400": "2011-03-09",
    "393": "2010-12-23",
    "392": "2010-09-28",
    "391": "2010-07-19",
    "390": "2009-09-23",
    "380": "2008-09-22",
    "371": "2007-07-05",
    "370": "2007-06-07",
    "362": "2006-05-30",
    "361": "2006-04-05",
    "360": "2005-11-21",
    "351": "2005-09-26",
    "350": "2005-08-22",
    "342": "2005-02-22",
    "341": "2004-11-04",
    "340": "2004-09-10",
    "330": "2004-07-20",
    "320": "2004-05-20",
    "311": "2004-02-01",
    "310": "2003-11-27",
    "302": "2003-10-01",
    "301": "2003-06-04",
    "300": "2002-05-14",
    "293": "2002-08-14",
    "291": "2001-11-29",
    "290": "2001-08-29",
    "281": "2001-07-19",
    "280": "2001-06-01",
    "272": "2001-01-29",
    "271": "2001-01-11",
    "270": "2000-11-30",
    "260": "1999-10-21",
    "250": "1999-08-02",
    # RAR 5.x (released 2013–2020)
    "500": "2013-10-12", "501": "2013-12-05",
    "510": "2014-04-16", "511": "2014-05-21",
    "520": "2014-12-18", "521": "2015-06-11",
    "530": "2015-08-10", "531": "2016-04-21",
    "540": "2016-10-25", "550": "2017-05-16",
    "560": "2018-02-05", "561": "2018-06-05",
    "570": "2019-05-06", "571": "2019-08-15",
    "580": "2020-01-14", "590": "2020-05-07", "591": "2020-07-27",
    # RAR 6.x (released 2020–)
    "600": "2020-12-08", "601": "2021-01-25", "602": "2021-07-08",
    "610": "2021-12-27", "611": "2022-03-03", "620": "2023-02-28",
    "621": "2023-05-01", "622": "2023-08-01", "623": "2023-08-30",
    "624": "2023-10-04",
    # RAR 7.x — dates after mid-2025 are approximate; ordering is what matters
    "700": "2024-02-28", "701": "2024-06-24",
    "710": "2025-02-25", "711": "2025-04-08", "712": "2025-06-24",
    "713": "2025-09-01", "720": "2025-11-15", "721": "2026-01-20",
    "722": "2026-04-01", "723": "2026-06-15",
}
# regex for rescene-format rar executables
_RESCENE_RAR_RE = re.compile(r"^\d{4}-\d{2}-\d{2}_rar\d+(?:b\d)?\.(exe)?$", re.IGNORECASE)

# Observational only (never affects reconstruction): pull the -mt thread count
# rescene locked out of the rar command line it fires, and the stream name out
# of its "Compressing X..." message, so we can record the winning
# (version, -mt) per stream for diagnostics and the .srr2 idea.
_MT_RE       = re.compile(r"-mt(\d+)")
_COMPRESS_RE = re.compile(r"^Compressing\s+(.+?)\.\.\.\s*$")
# The exact WinRAR build (exe) rescene invoked, incl. any beta suffix — its
# __str__ drops the beta ("5.11" for both rar511.exe and rar511b1.exe), so we
# capture the filename off the command line to distinguish beta vs final builds.
_EXE_RE      = re.compile(r"(\d{4}-\d{2}-\d{2}_rar\d+(?:b\d)?\.exe)", re.I)

import webview

SRRDB_API    = "https://api.srrdb.com/v1"
SRRDB_DL_SRR = "https://www.srrdb.com/download/srr/{}"
# srrdb "adds" (extra files uploaded for a release — proof jpgs etc., sometimes
# still unconfirmed). Some are packed INSIDE the rars but not stored in the SRR,
# so they're needed as rebuild sources. Served by add-id, not the /download/file
# path used for SRR-stored files.
SRRDB_DL_ADD = "https://www.srrdb.com/download/temp/{release}/{id}/{name}"

_UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
)
HEADERS = {"User-Agent": _UA, "Accept": "application/json, */*"}
# srrdb's download host sits behind Anubis, which serves a proof-of-work HTML
# challenge to *browser-like* (Mozilla) User-Agents — so our normal UA gets a
# wall instead of the file. A non-browser UA is passed straight through (Anubis
# assumes only scrapers fake a browser). Used for direct file/add downloads.
DL_HEADERS = {"User-Agent": "srrdb_tool/1.0 (+https://www.srrdb.com)"}

# Persistent SRR cache — every SRR we ever fetch is kept here, keyed by release
# name, so re-tests never re-hit srrdb (its download host is rate-limited per
# 24h). Survives output-folder deletion/cleanup,
# unlike the copy dropped in each release's output dir. Small files (KB–MB), so
# no eviction needed.
SRR_CACHE_DIR = Path(__file__).parent / "srr_cache"

MEDIA_EXTS = {
    ".avi", ".mkv", ".mp4", ".m4v", ".mov", ".wmv",
    ".iso", ".img", ".bin", ".cue", ".nrg", ".vob", ".ts", ".m2ts",
}

# Largest-file candidates for content-CRC lookup — media plus ROM/disc images
# (3DS/NDS/Switch/GC/Wii etc.), so name-independent CRC search resolves games.
CONTENT_HASH_EXTS = MEDIA_EXTS | {
    ".3ds", ".cia", ".cci", ".nds", ".cxi", ".nsp", ".xci", ".rvz",
    ".wbfs", ".gcm", ".gcz", ".cso", ".wud", ".wux", ".nsz", ".xcz",
}

# Metadata files placed in the output folder before/besides reconstruction
META_EXTS = {".srr", ".nfo", ".sfv", ".nzb", ".jpg", ".jpeg", ".png", ".diz", ".txt"}

# Max wall-clock time for a single release's reconstruction before it is
# aborted so the batch can continue. Generous — only fires on a genuine stall.
_RECON_TIMEOUT_S = 1800  # 30 minutes

# Single-file thread-count near-miss rescue: rescene greedily locks the first
# -mt whose test piece passes, then does one full compress; if the size is a
# few bytes off ("Still not fine") it gives up, even though the ORIGINAL may
# have been packed with a higher thread count that only diverges on the full
# file. When that happens we re-run the full compress at other thread counts
# (1..CAP) restricted to the already-locked version, and let the outer SFV
# verify confirm the CRC. Scene machines of the 3DS era were ≤8 cores; going
# higher just burns time. The per-release deadline still bounds the total.
_MT_RETRY_CAP = 8

# Multi-file rescue sweeps a small embedded "extra" (proof jpg / nfo / diz)
# whose thread count produced the right size but wrong bytes. rescene's piece
# test locks the FIRST -mt whose piece size matches, which is a LOWER bound — the
# real count is usually at or above it — so the sweep must reach well beyond the
# locked value. Scene machines of this era ran up to ~32 cores (observed a jpg
# locked at -mt19), and WinRAR 5.x tops out at 32 threads. Cap at 32 so a
# high-core release isn't missed; the per-release deadline still bounds the total.
_MT_RETRY_CAP_SMALL = 32
_LARGE_STREAM_BYTES = 16 * 1024 * 1024

# Common thread counts scene packers actually use (rescene-forum tip: odd counts
# are rarely if ever used). Ordering prior for -mt values we have NO win history
# for yet — tried before the odd/rare fill so a high-core release (mt16/24/32) is
# reached sooner. Our own win-frequency (from the results DB) still ranks FIRST;
# this only orders the not-yet-seen tail. NOT a filter — every value still runs.
# -mt0 is a DISTINCT algorithm rescene otherwise skips (issue #173); include it
# late in the prior so it's tried after the mainstream counts but before the odd
# tail. Sweeps must add 0 to their candidate set for it to be reached.
_MT_COMMON = (1, 2, 4, 6, 8, 12, 16, 24, 32, 0)

# Cross-version proof-jpg sweep: a scene group packs with a CONTEMPORARY WinRAR,
# so the proof jpg's build is almost always within a few years of the game's
# locked build. Cap the sweep to that era (± window) plus a hard build count, so
# a GENUINE wall gives up in minutes instead of grinding all 232 builds to the
# 30-min deadline. Group-history builds are always kept regardless of era. A
# winnable different-build jpg is era-adjacent and tried first anyway, so this
# costs ~zero real rebuilds while hugely improving batch throughput on walls.
_XVER_ERA_DAYS = 3 * 365          # ± ~3 years around the game's locked build
_XVER_MAX_BUILDS = 48             # hard backstop when dates can't be parsed

# Version-wall cache: a release whose main content file matches NO pack version
# (rescene tried every one, none reproduced it) is a pure version wall — only a
# bigger WinRAR pack can ever fix it, and re-grinding all 232 versions wastes
# ~30 min every re-run. Record such walls tagged with the pack signature; on a
# re-run, skip them instantly UNLESS the pack grew (new versions may crack it).
# _WALL_CACHE_GEN is bumped only if version-hunt logic changes materially, which
# auto-invalidates the cache so every wall gets one fresh attempt.
_WALL_CACHE_GEN = 5   # 2026-08-02: gen 4's sweep could exhaust its candidates
                      # for reasons that were NOT the archive's fault (a
                      # calibration probe that picked a build unable to run the
                      # recipe; single-file sets refused outright; a confirm pass
                      # that looked for the wrong basename when a source came
                      # from the extras store). Every one of those recorded a
                      # clean "no build reproduces this" miss. Bumping re-opens
                      # them. (gen 4 = SRR-driven recipe sweep; gen 3 =
                      # stored-extra method2 rescue; gen 2 was a reverted
                      # detection-rescue experiment.)

# ── SRR-driven recipe sweep ──────────────────────────────────────────────────
# rescene detects a build by compressing each packed file IN ISOLATION. When a
# set was packed by ONE `rar a <extra> <content>` command, WinRAR's -mt pipeline
# makes every stream depend on the OTHER files in the same command — measured on
# Art_Academy…EXiMiUS: the .nds alone is 11 bytes off at the TRUE build+-mt, and
# the .jpg alone matches NO build × -mt (0/2295). Isolated detection is therefore
# doomed for the whole class, and rescene's own "size == packed_size" test for a
# file that fits in one volume is size-only, so it locks a WRONG build and then
# grinds the pack to the 30-min deadline.
#
# The sweep replaces that with a direct measurement: compress ALL the set's files
# together (exactly as the group did) at each distinct-output build × -mt, and
# check the result against the per-volume stream CRCs the SRR already carries.
# For a file split over volumes, every non-final block's file_crc IS the CRC32 of
# that volume's slice of the COMPRESSED stream — hard, byte-level evidence that
# needs no copy of the original RARs.
_RECIPE_SWEEP_BUDGET_S = 900      # 15 min ceiling; a hit normally lands in <60 s
# Sources bigger than this are TRUNCATED for the sweep: only enough input to
# produce the first volume's compressed slice is needed, and that prefix is
# byte-identical to the full file's (verified across 8/12/16/24/33 MB cuts). Keeps
# the sweep seconds-cheap on multi-GB releases instead of minutes per candidate.
_RECIPE_TRUNC_MIN = 48 * 1024 * 1024

# Release-date version cap: a scene group can't pack with a WinRAR newer than the
# release date — so on the main version hunt, drop far-future builds and try those
# NEAREST the release date first. The real build is found fast and a wall exhausts
# far fewer versions. A generous margin AFTER the date is the safety net (repacks,
# slightly-off folder dates, a group on a marginally newer build); group-history
# builds are ALWAYS kept regardless. If a capped run still fails, the re-run widens
# to the whole pack — so nothing is permanently excluded.
_VERSION_CAP_MARGIN_DAYS = 3 * 365   # ~3 years after the release date (2–4 yr band)

# A release normally packs 1–2 "extras" (proof jpg / file_id.diz) that aren't in
# the content folder and get fetched from srrdb adds. Far more than this means a
# wrong release match or a release whose loose files simply aren't present (e.g.
# a PS5 asset dump with hundreds of .anim/.skel) — not an adds case. Bail with a
# single summary instead of iterating/logging each (which floods the UI).
_MAX_FETCH_ADDS = 8

# Local extras store: files larger than this are skipped by the scan (extras —
# nfo/diz/proof jpg/sfv/sample — are small; this avoids CRCing a stray game dump
# that happens to sit in a source folder). Content is indexed by CRC32 so the
# store is name-independent: two packs can hold same-named files with different
# bytes and both are kept, matched to a rebuild by the SRR's exact packed CRC.
_EXTRAS_MAX_FILE_BYTES = 200 * 1024 * 1024

# Auto-harvest: small extras (esp. file_id.diz) that we CRC-verify as real packed
# sources get copied into a managed folder + registered in the extras store, so
# the SAME bytes resolve offline for a group's OTHER releases. A file_id.diz is
# group-constant and NOT stored in the SRR (it's compressed content), so a group's
# later release often can't get it ("need the diz to rescene but it's inside the
# rar" catch-22) — harvesting one release's diz breaks it for the rest.
_HARVEST_EXTS = (".diz", ".nfo", ".sfv", ".jpg", ".jpeg", ".png", ".txt")
_HARVEST_MAX_BYTES = 4 * 1024 * 1024

# Scene "fix" releases (DIRFIX/NFOFIX) are metadata-only follow-ups — a
# corrected NFO or directory-name note, with NO packed content. There is
# nothing for rescene to rebuild, so a content-CRC / name match legitimately
# finds nothing; we detect these to report a clean "nothing to do" instead of
# a scary "not in srrdb" error.
_METADATA_ONLY_RE = re.compile(r'(?:^|[._\-\s])(dir|nfo)fix(?:[._\-\s]|$)',
                               re.IGNORECASE)

def _metadata_only_tag(*names: str):
    """Return 'DIRFIX'/'NFOFIX' if any name is a metadata-only scene follow-up
    (nothing to rebuild), else None."""
    for name in names:
        m = _METADATA_ONLY_RE.search(name or "")
        if m:
            return m.group(1).upper() + "FIX"
    return None


def _normalize_name(name: str) -> str:
    """Convert folder/file names to scene dot-notation for better srrdb search."""
    # Replace spaces, underscores, hyphens-surrounded-by-spaces with dots
    name = re.sub(r"[\s_]+", ".", name)
    # Collapse multiple dots
    name = re.sub(r"\.{2,}", ".", name)
    return name.strip(".")


def _canon_release(s: str) -> str:
    """Canonical form for exact release-name comparison: unify dots/underscores/
    spaces, lowercase — but KEEP the group hyphen. Lets a folder named with
    underscores match the same release stored with dots on srrdb, while still
    telling siblings apart (Pac_World_2 ≠ Pac_World)."""
    return re.sub(r"[._\s]+", ".", s or "").strip(".").lower()


def _strip_date_prefix(name: str) -> str:
    """Drop a leading dats.site / pre pre-date (YYYY-MM-DD-Release → Release).
    Scene names never start with such a date, so it is pure folder noise — and
    if left on it corrupts canon comparisons (folder canon gets a '2009-05-10-'
    prefix the srrdb release name lacks, so an otherwise-exact match is missed).
    Mirrors the strip search_srrdb_progressive already does for the query."""
    m = re.match(r"^\d{4}[-._]\d{2}[-._]\d{2}[-._](.+)$", name or "")
    return m.group(1) if m and len(m.group(1)) >= 4 else (name or "")


def _find_script(base_name: str) -> str | None:
    """Find srr.py / srs.py in the current Python's Scripts directory."""
    scripts_dir = Path(sys.executable).parent / "Scripts"
    for candidate in (base_name + ".py", base_name + ".exe", base_name):
        p = scripts_dir / candidate
        if p.exists():
            return str(p)
    # Also check PATH
    found = shutil.which(base_name) or shutil.which(base_name + ".exe")
    return found or None


class SrrdbToolAPI:
    def __init__(self):
        self._window  = None
        self._stop    = threading.Event()   # hard stop: abort release + halt batch
        self._skip    = threading.Event()   # soft skip: abort release, keep batch going
        self._running = False
        self._app_dir = Path(__file__).parent / "apps"
        self._srr_script = None
        self._srs_script = None
        self._live_procs = []      # rar.exe subprocesses of the current rebuild
        self._recon_deadline = 0   # wall-clock abort time for current rebuild
        # ── "run it fresh" prompt (see _use_db_history) ──
        self._fresh_reply  = threading.Event()  # GUI answered the prompt
        self._fresh_choice = None    # "fresh" | "history" | "always"
        self._fresh_all    = False   # latched: ignore history for the whole run
        self._fresh_release = None   # this release's decision (one prompt each)

    def set_window(self, w):
        self._window = w

    # ── Internal ──────────────────────────────────────────────────────────────

    def _emit(self, event: str, data: dict):
        if not self._window:
            return
        try:
            # Escape for the enclosing single-quoted JS string literal:
            # backslashes first (Windows paths!), then quotes. Without this,
            # any message containing a backslash fails JSON.parse in the GUI
            # and the log line is silently dropped.
            payload = (json.dumps(data, ensure_ascii=True)
                       .replace("\\", "\\\\").replace("'", "\\'"))
            self._window.evaluate_js(
                f"window.srrEvent('{event}', JSON.parse('{payload}'))"
            )
        except Exception:
            pass

    def _log(self, msg: str, cls: str = "info"):
        self._emit("log", {"msg": msg, "cls": cls})

    # Be a polite API client: srrdb publishes no rate limits but runs anti-bot
    # protection, so throttle to ~1 req/s, retry once on 429/503 with backoff,
    # and cache GET responses for the session (auto-match re-tests the same
    # releases across runs).
    _API_MIN_INTERVAL = 1.0
    _api_lock = threading.Lock()
    _api_last_call = 0.0
    _api_cache: dict[str, dict] = {}

    def clear_session_cache(self) -> dict:
        """Drop cached API responses — used by the GUI Rescan button so a
        re-test reflects current disk + srrdb state, not remembered results."""
        n = len(SrrdbToolAPI._api_cache)
        SrrdbToolAPI._api_cache.clear()
        return {"ok": True, "cleared": n}

    def _api_get(self, url: str) -> dict:
        cached = self._api_cache.get(url)
        if cached is not None:
            return cached
        with SrrdbToolAPI._api_lock:
            wait = SrrdbToolAPI._api_last_call + self._API_MIN_INTERVAL - time.time()
            if wait > 0:
                time.sleep(wait)
            SrrdbToolAPI._api_last_call = time.time()
        req = Request(url, headers=HEADERS)
        try:
            with urlopen(req, timeout=15) as r:
                data = json.loads(r.read().decode("utf-8", errors="replace"))
        except HTTPError as e:
            if e.code in (429, 503):
                retry_after = min(int(e.headers.get("Retry-After", 10) or 10), 60)
                self._log(
                    f"  srrdb rate limit (HTTP {e.code}) — waiting {retry_after}s…",
                    "warn",
                )
                time.sleep(retry_after)
                with urlopen(Request(url, headers=HEADERS), timeout=15) as r:
                    data = json.loads(r.read().decode("utf-8", errors="replace"))
            else:
                raise
        if len(self._api_cache) > 500:
            self._api_cache.clear()
        # Never cache empty search results — a transient srrdb hiccup would
        # otherwise make every retry return the same miss for the session.
        empty_search = (
            isinstance(data, dict)
            and "/search/" in url
            and int(data.get("resultsCount") or 0) == 0
        )
        if not empty_search:
            self._api_cache[url] = data
        return data

    # ── rescene / script detection ────────────────────────────────────────────

    def check_rescene(self) -> dict:
        """Check whether pyReScene is importable and srr/srs scripts are findable."""
        import_ok = False
        version   = None
        try:
            import rescene.main  # type: ignore
            version   = getattr(rescene, "__version__", "installed")
            import_ok = True
        except Exception:
            pass

        self._srr_script = _find_script("srr")
        self._srs_script = _find_script("srs")

        srr_ok = import_ok or bool(self._srr_script)
        srs_ok = bool(self._srs_script)

        winrar_pack = self._app_dir / "winrar_pack-4.20"
        rar_versions = []
        if winrar_pack.is_dir():
            rar_versions = [f.name for f in winrar_pack.iterdir()
                            if f.is_file() and _RESCENE_RAR_RE.match(f.name)]
        has_pack = winrar_pack.is_dir() and any(winrar_pack.rglob("wrar*.exe"))

        return {
            "rescene":      srr_ok,
            "srs":          srs_ok,
            "import_ok":    import_ok,
            "version":      version,
            "srr_script":   self._srr_script,
            "srs_script":   self._srs_script,
            "hint":         "pip install pyReScene" if not srr_ok else None,
            "rar_versions": rar_versions,
            "has_pack":     has_pack,
        }

    def setup_rar_executables(self) -> dict:
        """
        Extract rar.exe from each wrar*.exe installer in apps/winrar_pack-4.20/ and
        save as YYYY-MM-DD_rar<MAJOR><MINOR>.exe (the format rescene requires).
        Originals (wrar*.exe) and any other files are never touched.
        Old wrongly-named rar_X.YY.exe files from a previous run are removed.
        """
        winrar_pack = self._app_dir / "winrar_pack-4.20"
        if not winrar_pack.is_dir():
            return {"ok": False, "error": "apps/winrar_pack-4.20/ not found"}

        # Find 7z
        seven_zip = None
        for name in ("7z.exe", "7za.exe", "7zr.exe"):
            p = self._app_dir / name
            if p.exists():
                seven_zip = str(p)
                break
        if not seven_zip:
            seven_zip = shutil.which("7z") or shutil.which("7za")
        if not seven_zip:
            return {"ok": False, "error": "7z.exe not found in apps/ — needed to unpack installers"}

        installers = sorted(winrar_pack.rglob("wrar*.exe")) + sorted(
            winrar_pack.rglob("winrar-x*.exe"))
        if not installers:
            return {"ok": False, "error": "No wrar*.exe / winrar-x64-*.exe installers found in apps/winrar_pack-4.20/ (searched recursively)"}

        done = skipped = failed = 0
        messages = []

        # Remove old wrongly-named files from previous (broken) runs
        old_re = re.compile(r"^rar[_\-]\d+[\._]\d+\.exe$", re.IGNORECASE)
        for f in winrar_pack.iterdir():
            if f.is_file() and old_re.match(f.name):
                try:
                    f.unlink()
                    messages.append(f"  Removed old {f.name}")
                except Exception:
                    pass

        # RAR 5.00–6.24 ARE useful: post-2013 scene releases were made with
        # modern WinRAR in RAR4 mode (-ma4); the rebuilder injects -ma4 into
        # every 5.x+ invocation. WinRAR 7.00 REMOVED RAR4 creation entirely
        # ("Unknown option: ma4" — verified against the real binaries), so
        # 7.x must stay out of the pack. 6.24 (2023) is the last RAR4-capable.
        rar7_re = re.compile(r"^\d{4}-\d{2}-\d{2}_rar[7-9]\d\d(b\d)?\.exe$",
                             re.IGNORECASE)
        for f in list(winrar_pack.iterdir()):
            if f.is_file() and rar7_re.match(f.name):
                try:
                    f.unlink()
                    messages.append(f"  Removed {f.name} (7.x cannot create RAR4)")
                except Exception:
                    pass

        for installer in installers:
            # 3-DIGIT naming (wrar420.exe / winrar-x64-620b1.exe) — minor is two
            # digits. OLD 2-DIGIT naming (wrar50b3.exe = 5.00 beta 3, wrar26b8.exe
            # = 2.60 beta 8) — minor is ONE digit with an implied trailing zero.
            # Scene betas are very often shipped 2-digit, so accept both or the
            # crucial 5.x betas get silently skipped. Localised/dupe names
            # (wrar520fr.exe, wrar500_2.exe) fail the anchored $ and are ignored.
            m = (re.match(r"wrar(\d)(\d{2})(b\d)?\.exe$", installer.name, re.I)
                 or re.match(r"winrar-x\d+-(\d)(\d{2})(b\d)?\.exe$",
                             installer.name, re.I))
            if m:
                major, minor, beta = m.group(1), m.group(2), (m.group(3) or "")
            else:
                m = (re.match(r"wrar(\d)(\d)(b\d)?\.exe$", installer.name, re.I)
                     or re.match(r"winrar-x\d+-(\d)(\d)(b\d)?\.exe$",
                                 installer.name, re.I))
                if not m:
                    continue
                major, minor = m.group(1), m.group(2) + "0"  # 5.0 -> 500
                beta = m.group(3) or ""
            ver_key = f"{major}{minor}"
            if int(ver_key) >= 700:
                skipped += 1
                continue  # WinRAR 7.00+ removed RAR4 creation (-ma4)
            date = _WINRAR_DATES.get(ver_key, f"200{major}-01-01")  # fallback date
            target_name = f"{date}_rar{ver_key}{beta}.exe"
            target = winrar_pack / target_name

            # Remove any existing file for this version with a WRONG date
            stale_re = re.compile(
                rf"^\d{{4}}-\d{{2}}-\d{{2}}_rar{re.escape(ver_key)}{re.escape(beta)}\.exe$",
                re.IGNORECASE,
            )
            for stale in list(winrar_pack.iterdir()):
                if stale.is_file() and stale_re.match(stale.name) and stale != target:
                    try:
                        stale.unlink()
                        messages.append(f"  Removed stale {stale.name}")
                    except Exception:
                        pass

            if target.exists():
                messages.append(f"  {target_name} — already exists, skipped")
                skipped += 1
                continue

            try:
                with tempfile.TemporaryDirectory() as tmpdir:
                    subprocess.run(
                        [seven_zip, "e", str(installer), "rar.exe", f"-o{tmpdir}", "-y"],
                        capture_output=True, text=True, timeout=30,
                    )
                    extracted = Path(tmpdir) / "rar.exe"
                    if extracted.exists():
                        shutil.copy2(str(extracted), str(target))
                        messages.append(f"  {target_name} ✓  (from {installer.name})")
                        done += 1
                    else:
                        messages.append(f"  {installer.name} — rar.exe not found inside archive")
                        failed += 1
            except subprocess.TimeoutExpired:
                messages.append(f"  {installer.name} — timed out")
                failed += 1
            except Exception as e:
                messages.append(f"  {installer.name} — ERROR: {e}")
                failed += 1

        return {"ok": True, "done": done, "skipped": skipped, "failed": failed, "messages": messages}

    # ── Browse ────────────────────────────────────────────────────────────────

    def browse_folder(self) -> str:
        try:
            import tkinter as tk
            from tkinter import filedialog
            root = tk.Tk()
            root.withdraw()
            root.attributes("-topmost", True)
            path = filedialog.askdirectory()
            root.destroy()
            return path or ""
        except Exception:
            return ""

    # ── Folder scanning ───────────────────────────────────────────────────────

    def scan_folder(self, folder: str) -> dict:
        """Detect release name and content files in a single folder.
        Also checks immediate subdirectories so double-nested releases are found."""
        base = Path(folder)
        if not base.is_dir():
            return {"ok": False, "error": "Not a folder"}

        def _glob_with_sub(pattern: str) -> list[Path]:
            """Glob top-level, then fall back to immediate subfolders if empty."""
            hits = sorted(base.glob(pattern))
            if not hits:
                for sub in base.iterdir():
                    if sub.is_dir():
                        hits.extend(sub.glob(pattern))
                hits.sort()
            return hits

        nfos  = _glob_with_sub("*.nfo") or _glob_with_sub("*.NFO")
        sfvs  = _glob_with_sub("*.sfv") or _glob_with_sub("*.SFV")
        media = sorted(
            f for f in base.rglob("*")
            if f.is_file() and f.suffix.lower() in MEDIA_EXTS
        )
        if nfos:
            raw = nfos[0].stem
        elif sfvs:
            raw = sfvs[0].stem
        else:
            raw = base.name
        guess = _normalize_name(raw)
        return {
            "ok":            True,
            "release_guess": guess,
            "nfos":          [f.name for f in nfos],
            "sfvs":          [f.name for f in sfvs],
            "media":         [f.name for f in media],
        }

    def scan_subfolders(self, folder: str) -> list:
        """Return a scan dict for each immediate subfolder."""
        base = Path(folder)
        if not base.is_dir():
            return []
        results = []
        for sub in sorted(base.iterdir()):
            if not sub.is_dir():
                continue
            info = self.scan_folder(str(sub))
            results.append({
                "path":          str(sub),
                "name":          sub.name,
                "release_guess": info.get("release_guess", _normalize_name(sub.name)),
                "media_count":   len(info.get("media", [])),
                "has_nfo":       bool(info.get("nfos")),
                "has_sfv":       bool(info.get("sfvs")),
            })
        return results

    # ── srrdb.com API ─────────────────────────────────────────────────────────

    def _do_search(self, q: str) -> dict:
        """Raw srrdb search for exactly the query string q."""
        try:
            data = self._api_get(f"{SRRDB_API}/search/{quote(q)}")
            results = [
                {
                    "release": r.get("release", ""),
                    "date":    (r.get("date") or "")[:10],
                    "hasNFO":  r.get("hasNFO") == "yes",
                    "hasSRS":  r.get("hasSRS") == "yes",
                }
                for r in data.get("results", [])[:100]
            ]
            return {
                "ok":      True,
                "count":   int(data.get("resultsCount", 0)),
                "results": results,
                "query":   q,
            }
        except HTTPError as e:
            return {"ok": False, "error": f"HTTP {e.code}"}
        except Exception as e:
            return {"ok": False, "error": str(e)}

    def search_srrdb(self, query: str) -> dict:
        """Search srrdb.com — normalises spacing to dots, falls back to underscores."""
        q = _normalize_name(query.strip())
        if not q:
            return {"ok": False, "error": "Empty query"}
        res = self._do_search(q)
        if res.get("ok") and res["count"] == 0 and "." in q:
            q_alt = q.replace(".", "_")
            res2 = self._do_search(q_alt)
            if res2.get("ok") and res2["count"] > 0:
                return res2
        return res

    def search_srrdb_progressive(self, query: str) -> dict:
        """
        Name search with progressive fallback for poorly-named folders.
        Tries the full name first, then strips trailing dot-tokens one at a time
        (removes platform/region suffixes), then tries without the group tag.
        Returns the first query that gets results, or the 0-result response.
        """
        q = _normalize_name(query.strip())
        if not q:
            return {"ok": False, "error": "Empty query"}

        candidates: list[str] = []

        # Archive/pre folders are often prefixed with the pre date
        # (2011-06-05-Release_Name). Scene names never start with a date like
        # that, so strip it and search the clean name FIRST.
        m = re.match(r"^\d{4}[-._]\d{2}[-._]\d{2}[-._](.+)$", q)
        if m and len(m.group(1)) >= 4:
            candidates.append(m.group(1))
        candidates.append(q)

        # When multiple results come back, an EXACT name match to the query is
        # the answer — collapse to it so a sibling release (Pac_World vs
        # Pac_World_2, ...NFOFIX vs base) doesn't make an easy hit "ambiguous".
        def _collapse_exact(res: dict, attempt: str) -> dict:
            want = _canon_release(attempt)
            exact = [r for r in res.get("results", [])
                     if _canon_release(r.get("release", "")) == want]
            if len(exact) == 1 and res.get("count", 0) > 1:
                return {**res, "count": 1, "results": exact}
            return res

        # Try with underscores variant inline
        def _try(attempt: str) -> dict | None:
            res = self._do_search(attempt)
            if res.get("ok") and res["count"] > 0:
                res = _collapse_exact(res, attempt)
                res["query"] = attempt
                return res
            if "." in attempt:
                res2 = self._do_search(attempt.replace(".", "_"))
                if res2.get("ok") and res2["count"] > 0:
                    res2 = _collapse_exact(res2, attempt)
                    res2["query"] = attempt.replace(".", "_")
                    return res2
            return None

        # Later variants build on the date-stripped form when one exists
        base = candidates[0]

        # Strip the group tag (everything after the last hyphen) as a high-priority variant
        m = re.match(r"^(.+)-([A-Za-z0-9]{2,12})$", base)
        if m:
            candidates.append(m.group(1))  # without -GROUP

        # Progressively remove trailing dot-tokens (handles region/platform suffixes)
        parts = base.split(".")
        for trim in range(1, min(5, len(parts))):
            shorter = ".".join(parts[:-trim])
            if shorter and shorter not in candidates:
                candidates.append(shorter)

        for attempt in candidates:
            hit = _try(attempt)
            if hit:
                hit["query_trimmed"] = (attempt != q)  # flag that we simplified
                return hit

        # Nothing found — return zero-result response with original query
        return {"ok": True, "count": 0, "results": [], "query": q, "query_trimmed": False}

    def search_by_crc(self, crc32_hex: str) -> dict:
        """Search srrdb.com by archived-file CRC32 — the CRC of the content file
        INSIDE the RARs (e.g. the movie file), not the RAR volume CRCs from an
        SFV (those are not searchable). Exact-match, name-independent."""
        crc = crc32_hex.strip().upper().zfill(8)
        try:
            data = self._api_get(f"{SRRDB_API}/search/archive-crc:{crc}")
            results = [
                {
                    "release": r.get("release", ""),
                    "date":    (r.get("date") or "")[:10],
                    "hasNFO":  r.get("hasNFO") == "yes",
                    "hasSRS":  r.get("hasSRS") == "yes",
                }
                for r in data.get("results", [])[:10]
            ]
            return {
                "ok":      True,
                "count":   int(data.get("resultsCount", 0)),
                "results": results,
                "query":   f"archive-crc:{crc}",
                "by_crc":  True,
            }
        except HTTPError as e:
            return {"ok": False, "error": f"HTTP {e.code}"}
        except Exception as e:
            return {"ok": False, "error": str(e)}

    @staticmethod
    def _parse_sfv(sfv_path: str) -> list:
        """Parse an SFV file and return list of (filename, crc32_hex) tuples."""
        entries = []
        try:
            with open(sfv_path, "r", encoding="utf-8", errors="replace") as f:
                for line in f:
                    line = line.rstrip()
                    if not line or line.startswith(";"):
                        continue
                    parts = line.rsplit(" ", 1)
                    if len(parts) == 2 and len(parts[1]) == 8:
                        try:
                            int(parts[1], 16)
                            entries.append((parts[0].strip(), parts[1].upper()))
                        except ValueError:
                            pass
        except Exception:
            pass
        return entries

    def search_by_content_hash(self, folder: str) -> dict:
        """Hash the largest media file in the folder (CRC32) and look it up on
        srrdb via archive-crc. Works no matter how files/folders are named —
        the CRC of the original content file is stored in the database."""
        base = Path(folder)
        media = sorted(
            (f for f in base.rglob("*")
             if f.is_file() and f.suffix.lower() in CONTENT_HASH_EXTS),
            key=lambda f: f.stat().st_size, reverse=True,
        )
        if not media:
            return {"ok": True, "count": 0, "results": [], "query": "(no media file)"}
        target = media[0]
        self._log(
            f"  Hashing {target.name} ({target.stat().st_size:,} B) "
            "for exact srrdb CRC lookup…", "dim",
        )
        t0 = time.time()
        crc = self._crc32_file(str(target))
        self._log(f"  CRC32 = {crc:08X} ({time.time() - t0:.0f}s)", "dim")
        return self.search_by_crc(f"{crc:08X}")

    def _release_platform(self, name: str) -> str:
        """Platform token of a release/folder name ('' if none) — used to reject
        a content-CRC hit whose platform contradicts the folder's. Mirrors the
        detection in _parse_meta; tokens split on the usual scene separators."""
        toks = set(re.split(r"[._\-\s]+", (name or "").upper()))
        p = next((p for p in self._PLATFORMS if p in toks), "")
        if not p and toks & self._VIDEO_TOKENS:
            p = "video"
        return p

    def search_from_folder(self, folder: str) -> dict:
        """
        Progressive name search (NFO/SFV stem preferred over folder name), then
        content-hash CRC fallback when the name finds nothing.
        Looks for SFV/NFO in the folder AND one level of subfolders.
        Returns the same structure as search_srrdb plus 'method' field.
        """
        base = Path(folder)

        sfv_paths: list[Path] = (
            sorted(base.glob("*.sfv")) + sorted(base.glob("*.SFV"))
            + sorted(p for sub in base.iterdir() if sub.is_dir()
                     for p in list(sub.glob("*.sfv")) + list(sub.glob("*.SFV")))
        )
        nfo_paths: list[Path] = (
            sorted(base.glob("*.nfo")) + sorted(base.glob("*.NFO"))
            + sorted(p for sub in base.iterdir() if sub.is_dir()
                     for p in list(sub.glob("*.nfo")) + list(sub.glob("*.NFO")))
        )
        # The FOLDER name is the true release name; NFO/SFV stems are often
        # group codenames (hr-prrl, contrast-madden) that srrdb FUZZY-matches to
        # unrelated releases (e.g. hr-prrl → a vinyl album with cat# PRRLP001).
        # So: try the folder name first, and only ACCEPT a result that is an
        # EXACT name match to the hint or folder — a fuzzy single hit is never
        # trusted over CRC.
        folder_canon = _canon_release(_normalize_name(_strip_date_prefix(base.name)))
        hint_candidates: list[str] = [base.name]
        if nfo_paths and nfo_paths[0].stem not in hint_candidates:
            hint_candidates.append(nfo_paths[0].stem)
        if sfv_paths and sfv_paths[0].stem not in hint_candidates:
            hint_candidates.append(sfv_paths[0].stem)

        self._log(f"  [DIAG] folder={base.name!r} folder_canon={folder_canon!r} "
                  f"hints={hint_candidates}", "dim")
        best: dict | None = None
        best_trimmed = False
        for hint in hint_candidates:
            res = self.search_srrdb_progressive(hint)
            trimmed = res.pop("query_trimmed", False)
            names = [r.get("release", "") for r in res.get("results", [])]
            self._log(f"  [DIAG] hint={hint!r} → query={res.get('query')!r} "
                      f"count={res.get('count')} trimmed={trimmed} "
                      f"results={names[:6]}", "dim")
            if not (res.get("ok") and res.get("count", 0) > 0):
                continue
            hint_canon = _canon_release(_normalize_name(_strip_date_prefix(hint)))
            exact = next(
                (r for r in res["results"]
                 if _canon_release(r.get("release", "")) in (hint_canon, folder_canon)),
                None,
            )
            if exact:
                self._log(f"  [DIAG] EXACT-path pin → {exact.get('release')!r}", "dim")
                return {"ok": True, "count": 1, "results": [exact],
                        "query": res.get("query", hint),
                        "method": ("name (simplified)" if trimmed else "name")}
            if best is None:
                best, best_trimmed = res, trimmed

        # No EXACT name match — an exact content CRC beats any fuzzy name guess.
        hres = self.search_by_content_hash(folder)
        if hres.get("ok") and hres.get("count", 0) > 0:
            results = hres.get("results", [])
            # One inner-file CRC can be shared by SEVERAL releases (the same cart
            # dumped by different groups/regions — Jackass EUR: PUPPA + LiTE both
            # carry the identical .nds). So never blind-pick results[0]:
            #   1. prefer the release whose (date-stripped) name matches THIS
            #      folder — that's unambiguously the right one;
            #   2. else drop any result whose PLATFORM contradicts the folder's
            #      (an NDS folder must not resolve to a PS3 release — guards a
            #      CRC collision / wrong hashed file, e.g. Yu-Gi-Oh KOR NDS →
            #      Biohazard JPN PS3);
            #   3. only if nothing disambiguates do we keep the raw hits.
            name_match = [r for r in results
                          if _canon_release(r.get("release", "")) == folder_canon]
            fp = self._release_platform(_strip_date_prefix(base.name))
            if name_match:
                results = name_match
            elif fp:
                same_plat = [r for r in results
                             if self._release_platform(r.get("release", "")) == fp]
                results = same_plat  # may be [] → cross-platform collision, drop
            self._log(f"  [DIAG] CONTENT-CRC path → raw="
                      f"{[r.get('release','') for r in hres.get('results',[])][:6]} "
                      f"folder_plat={fp!r} → kept="
                      f"{[r.get('release','') for r in results][:6]}", "dim")
            if results:
                hres = {**hres, "results": results, "count": len(results)}
                hres["method"] = f"content CRC32 ({hres.get('query', '')})"
                return hres

        # Last resort: a fuzzy name hit — flag it so it isn't blindly trusted
        if best:
            self._log(f"  [DIAG] FUZZY-best path → "
                      f"{[r.get('release','') for r in best.get('results',[])][:6]}", "dim")
            best["method"] = "name (fuzzy — verify)"
            return best

        return {"ok": True, "count": 0, "results": [],
                "query": hint_candidates[0], "method": "name"}

    def get_release_details(self, release_name: str) -> dict:
        try:
            data = self._api_get(f"{SRRDB_API}/details/{quote(release_name)}")
            return {"ok": True, "details": data}
        except Exception as e:
            return {"ok": False, "error": str(e)}

    def score_release_against_folder(self, release_name: str, content_dir: str) -> float:
        """
        Score how well a release matches a content folder (0.0–1.0).
        Uses srrdb /details API to get expected archived files — no SRR download needed.
        Matches by filename (case-insensitive) and file size (within 1%).
        """
        try:
            data = self._api_get(f"{SRRDB_API}/details/{quote(release_name)}")
            archived = data.get("archived-files", [])
            if not archived:
                return 0.0
            content = {
                f.name.lower(): f.stat().st_size
                for f in Path(content_dir).rglob("*")
                if f.is_file()
            }
            content_sizes = set(content.values())
            score = 0.0
            for af in archived:
                name = af.get("name", "").lower()
                size = int(af.get("size", -1))
                if name in content:
                    local_size = content[name]
                    # Size within 1% → strong hit; name only → weak hit
                    if size <= 0 or abs(local_size - size) / max(size, 1) < 0.01:
                        score += 1.0
                    else:
                        score += 0.3
                elif size > 0 and size in content_sizes:
                    # Renamed content file — an exact byte-size match is a
                    # near-unique fingerprint for large media files.
                    score += 0.7
            return score / len(archived)
        except Exception:
            return 0.0

    def find_best_match(self, candidates: list, content_dir: str, max_test: int = 10) -> dict:
        """
        Score up to max_test candidates against content_dir using the details API.
        Returns {release, score} for the best match, or {release: None} if nothing scores > 0.
        """
        best_release = None
        best_score   = 0.0
        tested = 0
        for c in candidates[:max_test]:
            release = c.get("release", "") if isinstance(c, dict) else str(c)
            if not release:
                continue
            score = self.score_release_against_folder(release, content_dir)
            self._log(f"    [{score:.0%}] {release}", "dim")
            tested += 1
            if score > best_score:
                best_score   = score
                best_release = release
            if best_score >= 0.999:
                break  # perfect match — skip remaining API calls
        return {"release": best_release, "score": best_score, "tested": tested}

    # ── SRR download ──────────────────────────────────────────────────────────

    def _do_download_srr(self, release_name: str, dest: Path) -> dict:
        """Single attempt — download SRR for exact release_name into dest folder."""
        url      = SRRDB_DL_SRR.format(quote(release_name))
        out_path = dest / f"{release_name}.srr"
        self._log(f"  → {url}", "dim")
        try:
            req = Request(url, headers={**HEADERS, "Referer": "https://www.srrdb.com/"})
            with urlopen(req, timeout=60) as r:
                data = r.read()
            if not data:
                return {"ok": False, "not_found": False, "error": "Empty response"}
            if data[:2] == b"PK":
                import zipfile, io as _io
                with zipfile.ZipFile(_io.BytesIO(data)) as z:
                    srr_names = [n for n in z.namelist() if n.lower().endswith(".srr")]
                    if not srr_names:
                        return {"ok": False, "not_found": False, "error": "ZIP contained no SRR"}
                    out_path.write_bytes(z.read(srr_names[0]))
            else:
                out_path.write_bytes(data)
            return {"ok": True, "srr_path": str(out_path), "size": out_path.stat().st_size}
        except HTTPError as e:
            return {"ok": False, "not_found": e.code == 404, "error": f"HTTP {e.code}"}
        except Exception as e:
            return {"ok": False, "not_found": False, "error": str(e)}

    def download_srr(self, release_name: str, dest_dir: str) -> dict:
        """
        Download SRR, automatically retrying with dots/underscores swapped if the
        first attempt 404s (srrdb stores some releases with underscores, some with dots).
        """
        dest = Path(dest_dir)
        dest.mkdir(parents=True, exist_ok=True)

        names_to_try = [release_name]
        alt = release_name.replace(".", "_") if "." in release_name else release_name.replace("_", ".")
        if alt != release_name:
            names_to_try.append(alt)

        # 0 — serve from the persistent SRR cache if we've ever fetched this
        # release. This is the guard against re-hitting srrdb's rate-limited
        # download host on re-tests, cleared output folders, or a fresh batch.
        try:
            SRR_CACHE_DIR.mkdir(parents=True, exist_ok=True)
        except Exception:
            pass
        for name in names_to_try:
            cached = SRR_CACHE_DIR / f"{name}.srr"
            if cached.is_file() and cached.stat().st_size > 0:
                out_path = dest / f"{name}.srr"
                try:
                    if out_path.resolve() != cached.resolve():
                        shutil.copy2(str(cached), str(out_path))
                except Exception:
                    out_path = cached  # fall back to using the cached copy in place
                return {"ok": True, "srr_path": str(out_path),
                        "size": Path(out_path).stat().st_size, "cached": True}

        for name in names_to_try:
            result = self._do_download_srr(name, dest)
            if result["ok"]:
                if name != release_name:
                    self._log(f"  (srrdb uses '{name}' — saved under that name)", "dim")
                # Populate the persistent cache so this release is never fetched
                # again, even if its output folder is later deleted.
                try:
                    src = Path(result["srr_path"])
                    dst = SRR_CACHE_DIR / src.name
                    if dst.resolve() != src.resolve():
                        shutil.copy2(str(src), str(dst))
                except Exception:
                    pass
                return result
            if not result.get("not_found"):
                return result  # non-404 error, don't retry
        tried = " / ".join(names_to_try)
        return {"ok": False, "error": f"Release not found on srrdb.com (tried: {tried})"}

    def _srrdb_details(self, release_name: str):
        """Fetch a release's metadata from the srrdb JSON API (api.srrdb.com —
        NOT behind the download host's bot-wall). Returns the parsed dict with a
        '_resolved_name' key, or None. Retries dots/underscores like download."""
        names = [release_name]
        alt = (release_name.replace(".", "_") if "." in release_name
               else release_name.replace("_", "."))
        if alt != release_name:
            names.append(alt)
        for name in names:
            try:
                req = Request(f"{SRRDB_API}/details/{quote(name)}",
                              headers={**HEADERS, "Referer": "https://www.srrdb.com/"})
                with urlopen(req, timeout=30) as r:
                    data = json.loads(r.read().decode("utf-8", "replace"))
                if isinstance(data, dict) and (data.get("archived-files") or data.get("adds")):
                    data["_resolved_name"] = name
                    return data
            except Exception:
                continue
        return None

    def _download_add(self, release: str, add_id, filename: str,
                      dest_path: Path, expect_crc=None) -> dict:
        """Download one srrdb 'add' by id, verify its CRC32 against the API
        value, and write it to dest_path. Rejects HTML bot-wall responses so a
        blocked download never masquerades as a source file."""
        # filename may carry a subfolder (e.g. "[for reconstruction]/x.jpg") —
        # keep the "/" so the path resolves; only the segments get encoded.
        url = SRRDB_DL_ADD.format(release=quote(release), id=add_id,
                                  name=quote(filename, safe="/"))
        try:
            # Non-browser UA to bypass the Anubis challenge (see DL_HEADERS).
            req = Request(url, headers={**DL_HEADERS,
                                        "Referer": "https://www.srrdb.com/"})
            with urlopen(req, timeout=120) as r:
                data = r.read()
        except HTTPError as e:
            return {"ok": False, "error": f"HTTP {e.code}"}
        except Exception as e:
            return {"ok": False, "error": str(e)}
        head = data[:256].lstrip().lower()
        if not data or head.startswith(b"<!doctype") or head.startswith(b"<html"):
            return {"ok": False, "error": "srrdb returned an HTML page "
                    "(bot-wall / no network / not downloadable)"}
        if expect_crc:
            got = "%08X" % (zlib.crc32(data) & 0xffffffff)
            if got.upper() != str(expect_crc).upper():
                return {"ok": False,
                        "error": f"CRC mismatch (got {got}, expected {expect_crc})"}
        try:
            dest_path.parent.mkdir(parents=True, exist_ok=True)
            dest_path.write_bytes(data)
        except Exception as e:
            return {"ok": False, "error": f"write failed: {e}"}
        return {"ok": True, "size": len(data)}

    # ── rescene operations ────────────────────────────────────────────────────

    def _srr_extract(self, srr_path: str, out_dir: str) -> dict:
        """Extract stored files (NFO, SFV, SRS…) from an SRR using the rescene Python API."""
        Path(out_dir).mkdir(parents=True, exist_ok=True)
        try:
            import rescene.main as rm  # type: ignore
            rm.extract_files(
                srr_file=str(srr_path),
                out_folder=str(out_dir),
                extract_paths=True,
            )
            files = [str(f.relative_to(out_dir)) for f in Path(out_dir).rglob("*") if f.is_file()]
            return {"ok": True, "files": files}
        except ImportError:
            return {"ok": False, "error": "rescene not available — pip install pyReScene"}
        except Exception as e:
            return {"ok": False, "error": str(e)}

    def _find_rar_dir(self) -> str | None:
        """Find directory with rescene-format rar executables (YYYY-MM-DD_rar*.exe)."""
        for candidate in (self._app_dir / "winrar_pack-4.20", self._app_dir):
            if candidate.is_dir() and any(
                _RESCENE_RAR_RE.match(f.name) for f in candidate.iterdir() if f.is_file()
            ):
                return str(candidate)
        return None

    @staticmethod
    def _rar_set_prefix(vol_name: str) -> str:
        """Group RAR volume names into sets: name.rar/name.r00 → 'name';
        name.part01.rar → 'name'. Keeps any stored path (e.g. 'Subs/name')."""
        base = re.sub(r"\.[^.]+$", "", vol_name)
        base = re.sub(r"\.part\d+$", "", base, flags=re.IGNORECASE)
        return base

    def _srr_rar_sets(self, srr_path: str) -> dict:
        """Map each RAR set in the SRR to the content files packed inside it.
        Returns {set_prefix: {"volumes": [...], "packed": [...]}}."""
        from rescene.rar import RarReader, BlockType  # type: ignore
        sets: dict[str, dict] = {}
        current = None
        for block in RarReader(str(srr_path)).read_all():
            if block.rawtype == BlockType.SrrRarFile:
                current = self._rar_set_prefix(getattr(block, "file_name", ""))
                entry = sets.setdefault(current, {"volumes": [], "packed": []})
                entry["volumes"].append(getattr(block, "file_name", ""))
            elif block.rawtype == BlockType.RarPackedFile and current is not None:
                fname = getattr(block, "file_name", "")
                if fname and fname not in sets[current]["packed"]:
                    sets[current]["packed"].append(fname)
        return sets

    def _srr_packed_sizes(self, srr_path: str) -> dict:
        """Map packed-file basename (lower) -> expected UNPACKED size, read from
        the SRR's RarPackedFile blocks. Used to detect line-ending-mismatched
        text sources."""
        from rescene.rar import RarReader, BlockType  # type: ignore
        sizes: dict = {}
        for block in RarReader(str(srr_path)).read_all():
            if block.rawtype == BlockType.RarPackedFile:
                fname = getattr(block, "file_name", "")
                sz = getattr(block, "unpacked_size", None)
                if fname and sz is not None:
                    sizes.setdefault(Path(fname).name.lower(), sz)
        return sizes

    def _srr_packed_info(self, srr_path: str) -> dict:
        """Map packed-file basename (lower) -> (unpacked_size, crc_hex) from the
        SRR's RarPackedFile blocks. crc_hex is the 8-digit lowercase CRC32 of
        the packed file's *content* (None if the SRR doesn't carry it, e.g. very
        old RAR that stores the CRC only in a trailing block). Used to resolve a
        needed source from the local extras DB by exact content."""
        from rescene.rar import RarReader, BlockType  # type: ignore
        info: dict = {}
        for block in RarReader(str(srr_path)).read_all():
            if block.rawtype == BlockType.RarPackedFile:
                fname = getattr(block, "file_name", "")
                sz = getattr(block, "unpacked_size", None)
                crc = getattr(block, "file_crc", None)
                if fname and sz is not None:
                    key = Path(fname).name.lower()
                    if key not in info:
                        crc_hex = (f"{crc:08x}"
                                   if isinstance(crc, int) and crc != 0xFFFFFFFF
                                   else None)
                        info[key] = (sz, crc_hex)
        return info

    # ── Local extras store (content-addressed by CRC32) ───────────────────────
    # A user-supplied folder (or several) of scene "extras" — the exact packed
    # nfo/diz/proof-jpg copies that some groups pack DIFFERENTLY from the loose
    # copy the SRR stored, and that srrdb has no add for. Indexed once into a
    # single SQLite file in the tosort folder (easy backup), keyed by content
    # CRC32 so it's name-independent and can hold same-named files from two packs
    # with different bytes. During a rebuild, a source we can't otherwise produce
    # is looked up by the SRR's EXACT packed CRC+size; on a match the file is
    # COPIED (never moved) into the release and re-verified. Entirely inert when
    # no folders are configured, so the normal rebuild path is untouched.

    @property
    def _extras_db_path(self) -> Path:
        return Path(__file__).parent / "srrdb_extras.db"

    def _extras_connect(self, create: bool = True):
        """Open the extras SQLite DB (schema-created on demand). Returns a
        connection, or None when create=False and the DB doesn't exist yet.
        Default rollback journal (no WAL) so the store stays a single .db file
        for backup. Bound to the calling thread (SQLite default)."""
        p = self._extras_db_path
        if not create and not p.exists():
            return None
        con = sqlite3.connect(str(p))
        con.execute("""CREATE TABLE IF NOT EXISTS files(
            path TEXT PRIMARY KEY, crc TEXT, size INTEGER, name TEXT,
            folder TEXT, mtime REAL, sha TEXT)""")
        con.execute("CREATE INDEX IF NOT EXISTS ix_files_crc ON files(crc, size)")
        # Migrate stores created before the sha column existed. Old rows keep
        # sha=NULL until the next scan backfills them (see _extras_scan).
        if "sha" not in [r[1] for r in con.execute("PRAGMA table_info(files)")]:
            con.execute("ALTER TABLE files ADD COLUMN sha TEXT")
        con.execute("CREATE TABLE IF NOT EXISTS folders(path TEXT PRIMARY KEY, added REAL)")
        con.execute("CREATE TABLE IF NOT EXISTS meta(key TEXT PRIMARY KEY, value TEXT)")
        con.commit()
        return con

    @staticmethod
    def _file_crc32(path: str):
        """Streaming CRC32 of a file → 8-digit lowercase hex, or None on error."""
        try:
            crc = 0
            with open(path, "rb") as f:
                for chunk in iter(lambda: f.read(1 << 20), b""):
                    crc = zlib.crc32(chunk, crc)
            return f"{crc & 0xffffffff:08x}"
        except OSError:
            return None

    @staticmethod
    def _file_hashes(path: str):
        """Streaming CRC32 + SHA-256 in ONE read pass → (crc8hex, sha64hex),
        or (None, None) on error. SHA-256 disambiguates the astronomically rare
        case of two DIFFERENT files sharing a CRC32 + size (see extras store)."""
        try:
            crc = 0
            h = hashlib.sha256()
            with open(path, "rb") as f:
                for chunk in iter(lambda: f.read(1 << 20), b""):
                    crc = zlib.crc32(chunk, crc)
                    h.update(chunk)
            return f"{crc & 0xffffffff:08x}", h.hexdigest()
        except OSError:
            return None, None

    @staticmethod
    def _extras_collision_groups(con):
        """(crc, size) keys that hold TWO OR MORE distinct sha values — genuine
        collisions where DIFFERENT files share a CRC32 + size. Returns a list of
        (crc, size, [paths]); empty when the store is clean (the normal case).
        Rows with no sha yet (un-rescanned) are ignored, so a collision is only
        ever reported from real SHA-256 evidence."""
        from collections import defaultdict
        by_key = defaultdict(list)
        for crc, size, sha, path in con.execute(
                "SELECT crc, size, sha, path FROM files WHERE sha IS NOT NULL"):
            by_key[(crc, size)].append((sha, path))
        out = []
        for (crc, size), items in by_key.items():
            if len({s for s, _ in items}) > 1:
                out.append((crc, size, [p for _, p in items]))
        return out

    def extras_collisions(self) -> list:
        """GUI/audit helper: list the store's genuine CRC+size collisions as
        {crc,size,files:[…]} — normally empty."""
        con = self._extras_connect(create=False)
        if not con:
            return []
        try:
            return [{"crc": c, "size": s, "files": ps}
                    for c, s, ps in self._extras_collision_groups(con)]
        finally:
            con.close()

    # --- GUI-facing config + scan API ---
    def extras_get_folders(self) -> dict:
        """Configured extras folders (with per-folder file counts) + store stats,
        for the GUI. Persisted in the DB, so it survives a restart."""
        con = self._extras_connect(create=False)
        if not con:
            return {"folders": [], "total_files": 0, "last_scan": None}
        try:
            folders = [r[0] for r in con.execute("SELECT path FROM folders ORDER BY path")]
            counts = {r[0]: r[1] for r in con.execute(
                "SELECT folder, COUNT(*) FROM files GROUP BY folder")}
            total = con.execute("SELECT COUNT(*) FROM files").fetchone()[0]
            last = con.execute("SELECT value FROM meta WHERE key='last_scan'").fetchone()
            return {
                "folders": [{"path": f, "files": counts.get(f, 0)} for f in folders],
                "total_files": total,
                "last_scan": last[0] if last else None,
                "collisions": len(self._extras_collision_groups(con)),
                "scanning": bool(getattr(self, "_extras_scanning", False)),
            }
        finally:
            con.close()

    def extras_add_folder(self) -> dict:
        """Browse for a folder and add it to the extras source list."""
        folder = self.browse_folder()
        if not folder:
            return {"ok": False}
        con = self._extras_connect()
        try:
            con.execute("INSERT OR IGNORE INTO folders(path, added) VALUES(?,?)",
                        (folder, time.time()))
            con.commit()
        finally:
            con.close()
        return {"ok": True, "folder": folder}

    def extras_remove_folder(self, folder: str) -> dict:
        """Remove a folder from the list and drop its indexed files (the folder
        on disk is never touched)."""
        con = self._extras_connect(create=False)
        if not con:
            return {"ok": True}
        try:
            con.execute("DELETE FROM folders WHERE path=?", (folder,))
            con.execute("DELETE FROM files WHERE folder=?", (folder,))
            con.commit()
        finally:
            con.close()
        return {"ok": True}

    def extras_rescan(self) -> dict:
        """Kick off an incremental (re)scan of all configured folders in a
        background thread — emits log + progress events, then an 'extras' event
        when done."""
        if getattr(self, "_extras_scanning", False):
            return {"ok": False, "error": "a scan is already running"}
        self._extras_scanning = True
        threading.Thread(target=self._extras_scan_thread, daemon=True).start()
        return {"ok": True}

    def _extras_scan_thread(self):
        try:
            self._extras_scan()
        except Exception as e:
            self._log(f"Extras scan error: {e}", "err")
        finally:
            self._extras_scanning = False
            self._emit("extras", {"scanning": False})

    def _extras_scan(self):
        con = self._extras_connect()
        try:
            folders = [r[0] for r in con.execute("SELECT path FROM folders")]
            if not folders:
                self._log("Extras: no source folders configured — nothing to scan.",
                          "warn")
                return
            # Incremental: keep CRCs of files whose size+mtime are unchanged.
            have = {r[0]: (r[1], r[2], r[3]) for r in
                    con.execute("SELECT path, size, mtime, sha FROM files")}
            allfiles = []
            for folder in folders:
                if not os.path.isdir(folder):
                    self._log(f"Extras: folder not found, skipping — {folder}", "warn")
                    continue
                for root, _dirs, names in os.walk(folder):
                    for n in names:
                        allfiles.append((folder, os.path.join(root, n)))
            total = len(allfiles)
            self._log(f"Extras: scanning {total:,} file(s) across "
                      f"{len(folders)} folder(s)…", "info")
            seen = set()
            added = updated = skipped = big = 0
            for i, (folder, fp) in enumerate(allfiles, 1):
                if i % 200 == 0:
                    self._emit("progress", {
                        "pct": int(i * 100 / max(total, 1)),
                        "label": f"Scanning extras {i:,}/{total:,}"})
                try:
                    st = os.stat(fp)
                except OSError:
                    continue
                seen.add(fp)
                if st.st_size > _EXTRAS_MAX_FILE_BYTES:
                    big += 1
                    continue
                prev = have.get(fp)
                # Skip unchanged files — but only once they carry a sha, so a
                # store built before the sha column gets backfilled on this pass.
                if (prev and prev[0] == st.st_size
                        and abs(prev[1] - st.st_mtime) < 1e-6 and prev[2]):
                    skipped += 1
                    continue
                crc, sha = self._file_hashes(fp)
                if crc is None:
                    continue
                con.execute(
                    "INSERT OR REPLACE INTO files"
                    "(path,crc,size,name,folder,mtime,sha) VALUES(?,?,?,?,?,?,?)",
                    (fp, crc, st.st_size, os.path.basename(fp).lower(), folder,
                     st.st_mtime, sha))
                if prev:
                    updated += 1
                else:
                    added += 1
                if (added + updated) % 500 == 0:
                    con.commit()
            # Prune rows whose file vanished (deleted / folder removed).
            removed = 0
            for (path,) in list(con.execute("SELECT path FROM files")):
                if path not in seen:
                    con.execute("DELETE FROM files WHERE path=?", (path,))
                    removed += 1
            con.execute("INSERT OR REPLACE INTO meta(key,value) VALUES('last_scan',?)",
                        (time.strftime("%Y-%m-%d %H:%M:%S"),))
            con.commit()
            total_now = con.execute("SELECT COUNT(*) FROM files").fetchone()[0]
            coll = self._extras_collision_groups(con)
            self._log(
                f"Extras scan done: {added:,} new, {updated:,} updated, "
                f"{skipped:,} unchanged, {removed:,} removed"
                + (f", {big:,} too big (skipped)" if big else "")
                + f" — {total_now:,} file(s) indexed.", "ok")
            if coll:
                self._log(
                    f"  ⚠ {len(coll)} CRC+size collision(s) — DIFFERENT files "
                    "sharing a checksum (SHA-256 differs). When a release needs "
                    "one of these the store surfaces all candidates so you can "
                    "tell which is the real proof:", "warn")
                for crc, size, paths in coll[:20]:
                    self._log(
                        f"    CRC {crc} · {size:,} B → "
                        + "; ".join(os.path.basename(p) for p in paths), "dim")
            self._emit("progress", {"pct": 100, "label": ""})
            self._emit("extras", {"scanning": False, "total_files": total_now})
        finally:
            con.close()

    def _extras_lookup(self, crc_hex, size):
        """Return an on-disk extras path whose content CRC+size match, else None.
        Read-only and safe with no DB configured (returns None)."""
        if not crc_hex:
            return None
        con = self._extras_connect(create=False)
        if not con:
            return None
        try:
            for (path,) in con.execute(
                    "SELECT path FROM files WHERE crc=? AND size=?", (crc_hex, size)):
                if os.path.isfile(path):
                    return path
        except sqlite3.Error:
            return None
        finally:
            con.close()
        return None

    @property
    def _harvest_dir(self) -> Path:
        return Path(__file__).parent / "srrdb_harvest"

    def _harvest_extra(self, src_path, name=None):
        """Copy a CRC-verified small extra into the managed harvest folder and
        register it in the extras store, so the SAME file (by content CRC+size)
        resolves OFFLINE for other releases — chiefly a group's constant
        file_id.diz fetched as a srrdb add on one release, breaking the catch-22
        on the group's later releases that srrdb has no add for.

        Best-effort and fully guarded: any failure is swallowed, and it only ADDS
        a resolution source — never removes or changes existing behaviour. Skips
        when the exact content is already resolvable from the store."""
        try:
            src = Path(src_path)
            nm = name or src.name
            if src.suffix.lower() not in _HARVEST_EXTS or not src.is_file():
                return
            size = src.stat().st_size
            if size == 0 or size > _HARVEST_MAX_BYTES:
                return
            crc, sha = self._file_hashes(str(src))
            if not crc:
                return
            con = self._extras_connect(create=True)
            if not con:
                return
            try:
                # Already resolvable from a LIVE store file? Then nothing to do.
                for (pp,) in con.execute(
                        "SELECT path FROM files WHERE crc=? AND size=?",
                        (crc, size)):
                    if os.path.isfile(pp):
                        return
                hdir = self._harvest_dir
                hdir.mkdir(parents=True, exist_ok=True)
                dest = hdir / f"{crc}_{size}_{nm}"
                if not dest.exists():
                    shutil.copyfile(str(src), str(dest))
                con.execute(
                    "INSERT OR REPLACE INTO files"
                    "(path,crc,size,name,folder,mtime,sha) VALUES(?,?,?,?,?,?,?)",
                    (str(dest), crc, size, nm, str(hdir),
                     dest.stat().st_mtime, sha))
                con.execute(
                    "INSERT OR IGNORE INTO folders(path, added) VALUES(?, ?)",
                    (str(hdir), time.time()))
                con.commit()
                self._log(f"  Harvested {nm} → extras store (reusable for this "
                          "group's other releases).", "dim")
            finally:
                con.close()
        except Exception:
            pass

    def _resolve_from_extras(self, items, is_wrong, hints, stored_pool, packed_info):
        """For each (packed_path, name) in items, look the SRR's exact packed
        CRC+size up in the local extras store; on a content match COPY it in
        (never move), re-verify the CRC, and point the hint at it. Returns the
        sublist it resolved. Inert when the store is empty/absent.

        Collision-aware: if TWO OR MORE distinct files (different SHA-256) share
        the needed CRC+size, that's a genuine checksum collision — we can't tell
        the true packed copy from CRC alone, so we surface every candidate by
        name and use the first (the downstream SFV verify is the real arbiter)."""
        resolved = []
        con = self._extras_connect(create=False)
        try:
            for p, nm in items:
                exp = packed_info.get(nm)
                if not exp:
                    continue
                size, crc_hex = exp
                if not crc_hex or not con:
                    continue
                # On-disk candidates for this exact CRC+size, one per DISTINCT
                # content (sha). An un-hashed old row keys on its own path, so it
                # can't masquerade as a collision with a hashed row.
                cands, seen = [], set()
                try:
                    for sha, path in con.execute(
                            "SELECT sha, path FROM files WHERE crc=? AND size=?",
                            (crc_hex, size)):
                        if not os.path.isfile(path):
                            continue
                        key = sha or path
                        if key in seen:
                            continue
                        seen.add(key)
                        cands.append(path)
                except sqlite3.Error:
                    continue
                if not cands:
                    continue
                if len(cands) > 1:
                    self._log(
                        f"  ⚠ Extras store: {len(cands)} DIFFERENT files share "
                        f"CRC {crc_hex} + {size:,} B for {Path(p).name} — a real "
                        "SHA-256 collision. Using the first; the SFV verify "
                        "decides. Alternates: "
                        + "; ".join(os.path.basename(c) for c in cands[1:]),
                        "warn")
                hit = cands[0]
                dest = (stored_pool / "_adds" / Path(p).name if is_wrong
                        else stored_pool / Path(p).name)
                try:
                    dest.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copyfile(hit, dest)
                except OSError:
                    continue
                if self._file_crc32(str(dest)) != crc_hex:
                    try:
                        dest.unlink()
                    except OSError:
                        pass
                    continue
                hints[p] = str(dest)
                resolved.append((p, nm))
                self._log(f"  Extras store: matched {Path(p).name} by CRC "
                          f"{crc_hex} — copied from your local pack.", "ok")
        finally:
            if con:
                con.close()
        return resolved

    def _fix_text_source_endings(self, hints: dict, srr_path: str, out_dir: str):
        """Some groups pack an nfo/diz with different line endings than the copy
        stored in the SRR, so the extracted source is the WRONG SIZE (e.g. CRLF
        5027 B vs the packed LF 5008 B) and rescene bails with 'Data file is not
        the correct size'. When a text source's size doesn't match the packed
        size, write a CRLF<->LF variant that DOES match and point the hint at it.

        Only ever touches text sources whose size is already off, so releases
        that rebuild fine are untouched. The SFV/CRC verify stays the final
        guard — a size that matches but isn't byte-exact still fails cleanly.

        Returns the list of packed-file names that remain UNRECONSTRUCTABLE — a
        wrong-size text source that is neither a clean CRLF/LF variant nor
        fetchable as a srrdb add. The caller fast-fails on those (rescene would
        just recompress the big content file, then bail on the text block)."""
        declined: list[str] = []
        try:
            expected = self._srr_packed_sizes(srr_path)
        except Exception:
            return declined
        if not expected:
            return declined
        conv_dir = Path(out_dir) / "_stored" / "_leconv"
        for p, src in list(hints.items()):
            nm = Path(p).name.lower()
            if not nm.endswith((".nfo", ".diz", ".txt")):
                continue
            exp = expected.get(nm)
            if exp is None:
                continue
            try:
                data = Path(src).read_bytes()
            except Exception:
                continue
            if len(data) == exp:
                continue  # already the right size — leave it alone
            lf = data.replace(b"\r\n", b"\n").replace(b"\r", b"\n")
            crlf = lf.replace(b"\n", b"\r\n")
            matched = False
            for variant, kind in ((lf, "LF"), (crlf, "CRLF")):
                if len(variant) == exp:
                    try:
                        conv_dir.mkdir(parents=True, exist_ok=True)
                        dest = conv_dir / Path(p).name
                        dest.write_bytes(variant)
                    except Exception:
                        break
                    hints[p] = str(dest)
                    self._log(f"  Line-ending fix: {Path(p).name} "
                              f"{len(data):,}→{exp:,} B ({kind}) to match the "
                              "packed copy.", "dim")
                    matched = True
                    break
            if not matched:
                self._log(f"  ⚠ {Path(p).name}: stored copy is {len(data):,} B "
                          f"but the packed copy is {exp:,} B, and it's NOT a clean "
                          "CRLF/LF difference — can't reconstruct it from the "
                          "stored copy (needs the exact packed file).", "warn")
                declined.append(Path(p).name)
        return declined

    def _fresh_rescene(self):
        """Import a pristine rescene.main (purging any cached copy) and apply our
        patches. Called once per reconstruction so no module-level global state
        (archived_files, temp dirs, repository, event subscribers) can leak from
        one release into the next during a long batch."""
        for name in [n for n in list(sys.modules)
                     if n == "rescene" or n.startswith("rescene.")]:
            del sys.modules[name]
        import rescene.main as rm  # type: ignore

        # --- -ma4 + dictionary-switch injection for RAR 5.x/6.x binaries ---
        _orig_popen = rm.custom_popen
        _r5plus = re.compile(r"\d{4}-\d{2}-\d{2}_rar[5-9]\d\d(b\d)?\.exe$",
                             re.IGNORECASE)
        _MD_LETTERS = {"a": "64", "b": "128", "c": "256", "d": "512",
                       "e": "1024", "f": "2048", "g": "4096"}
        def _fix_md(arg):
            m = re.match(r"^-md([a-gA-G]|\d+)$", arg)
            if not m:
                return arg
            v = m.group(1)
            return f"-md{_MD_LETTERS.get(v.lower(), v)}k"
        def _inject(cmd):
            if (len(cmd) >= 3 and str(cmd[1]).lower() == "a"
                    and _r5plus.search(str(cmd[0])) and "-ma4" not in cmd):
                return [cmd[0], cmd[1], "-ma4"] + [_fix_md(str(a)) for a in cmd[2:]]
            return cmd
        def _popen_ma4(cmd, *a, **kw):
            # User stop/skip: refuse new rar.exe spawns immediately so an
            # in-flight reconstruction unwinds and returns at once instead of
            # churning to the next version. stop_process() also kills the
            # running rar.exe, so the current compress dies and this blocks the
            # follow-up spawn — the reconstruction aborts within a second.
            if self._stop.is_set() or self._skip.is_set():
                raise RuntimeError("stopped by user")
            # Hard deadline: a stuck release (e.g. rescene spiralling on an
            # incompressible embedded jpg) must never hang the whole batch.
            # Once past the deadline, refuse new rar.exe spawns so reconstruct
            # exhausts and returns; the heartbeat also kills the running one.
            deadline = getattr(self, "_recon_deadline", 0)
            if deadline and time.time() > deadline:
                # Mark that the hunt was CUT OFF (not a clean exhaustion) so the
                # wall-cache won't treat a deadline-truncated run as "tried
                # everything" — some betas may be untested.
                self._recon_hit_deadline = True
                raise RuntimeError("reconstruction deadline exceeded")
            try:
                cmd = _inject(cmd)
            except Exception:
                pass
            proc = _orig_popen(cmd, *a, **kw)
            try:
                self._live_procs.append(proc)
            except Exception:
                pass
            return proc
        rm.custom_popen = _popen_ma4
        _orig_full = rm.RarExecutable.full
        def _full_ma4(rar_self):
            try:
                return _inject(_orig_full(rar_self))
            except Exception:
                return _orig_full(rar_self)
        rm.RarExecutable.full = _full_ma4

        # --- compressed-stream cache for the multi-file cross-version sweep ----
        # The per-stream version sweep (see _rescue_multifile_crc) re-runs the
        # WHOLE reconstruction once per candidate WinRAR build to hunt a proof
        # jpg's true version. Every pass otherwise recompresses the big content
        # file identically — it's pinned to one build+thread-count throughout —
        # so a 256 MB .3ds burns ~40 s per pass for nothing. Cache that one
        # expensive output (rescene runs it via the module's subprocess.Popen,
        # NOT custom_popen) and replay it. Strictly gated on self._compress_cache
        # being a live dict, which ONLY the sweep sets; when it's None the shim
        # is a pure passthrough, so the normal reconstruction path is byte-for-
        # byte unchanged. Correctness never rests on the cache: a hit reproduces
        # exactly what rar.exe emits for identical (exe, args, source) inputs,
        # and the final SFV CRC verify still guards every accepted rebuild.
        _api = self
        _real_sp = subprocess

        def _cache_probe(cmd):
            # A cacheable full compress: `rar a … OUT SRC` where OUT is the
            # pyReScene_compressed.rar output and SRC is a real ≥16 MB content
            # file (never a pyReScene_data_piece, never the small jpg we sweep).
            if len(cmd) < 3 or str(cmd[1]).lower() != "a":
                return None
            out_path = src_path = None
            for a in cmd[2:]:
                s = str(a)
                base = os.path.basename(s).lower()
                if base == "pyrescene_compressed.rar":
                    out_path = s
                elif "pyrescene" not in base and os.path.isfile(s):
                    try:
                        if os.path.getsize(s) >= _LARGE_STREAM_BYTES:
                            src_path = s
                    except OSError:
                        pass
            return (src_path, out_path) if (src_path and out_path) else None

        def _cache_key(cmd, src_path):
            flags = tuple(str(a) for a in cmd
                          if str(a) == "a" or str(a).startswith("-"))
            st = os.stat(src_path)
            return (os.path.basename(str(cmd[0])).lower(), flags,
                    os.path.basename(src_path).lower(),
                    st.st_size, int(st.st_mtime))

        def _cache_dir():
            d = getattr(_api, "_ccache_dir", None)
            if not d:
                d = tempfile.mkdtemp(prefix="srrdb_ccache_")
                _api._ccache_dir = d
            return d

        def _cache_restore(entry, out_path):
            out_dir = os.path.dirname(out_path)
            for fname, cfile in entry:
                dst = os.path.join(out_dir, fname)
                try:
                    if os.path.exists(dst):
                        os.unlink(dst)
                except OSError:
                    pass
                try:
                    os.link(cfile, dst)         # instant, same volume
                except OSError:
                    shutil.copy2(cfile, dst)

        def _cache_store(out_path, key, cache):
            out_dir = os.path.dirname(out_path)
            stem = os.path.splitext(os.path.basename(out_path))[0].lower()
            cdir = _cache_dir()
            entry = []
            for f in os.listdir(out_dir):
                if not f.lower().startswith(stem):
                    continue
                src = os.path.join(out_dir, f)
                cf = os.path.join(cdir, f"{len(cache)}_{f}")
                try:
                    if os.path.exists(cf):
                        os.unlink(cf)
                    os.link(src, cf)
                except OSError:
                    try:
                        shutil.copy2(src, cf)
                    except OSError:
                        return              # skip caching this one, no harm
                entry.append((f, cf))
            if entry:
                cache[key] = entry

        class _FakeProc(object):
            returncode = 0
            def communicate(self, *a, **k):
                return (b"", b"")
            def wait(self, *a, **k):
                return 0
            def poll(self):
                return 0
            def kill(self):
                pass
            def terminate(self):
                pass

        class _CachingProc(object):
            def __init__(self, proc, out_path, key, cache):
                self._p = proc
                self._out = out_path
                self._key = key
                self._cache = cache
            def communicate(self, *a, **k):
                res = self._p.communicate(*a, **k)
                try:
                    if self._p.returncode == 0 and self._key not in self._cache:
                        _cache_store(self._out, self._key, self._cache)
                except Exception:
                    pass
                return res
            def __getattr__(self, n):
                return getattr(self._p, n)

        class _SPShim(object):
            def __getattr__(self, n):
                return getattr(_real_sp, n)
            def Popen(self, cmd, *a, **k):
                cache = getattr(_api, "_compress_cache", None)
                if cache is None:
                    return _real_sp.Popen(cmd, *a, **k)
                try:
                    probe = _cache_probe(cmd)
                except Exception:
                    probe = None
                if not probe:
                    return _real_sp.Popen(cmd, *a, **k)
                src_path, out_path = probe
                try:
                    key = _cache_key(cmd, src_path)
                except Exception:
                    return _real_sp.Popen(cmd, *a, **k)
                entry = cache.get(key)
                if entry:
                    try:
                        _cache_restore(entry, out_path)
                        return _FakeProc()
                    except Exception:
                        pass                # fall through to a real compress
                return _CachingProc(_real_sp.Popen(cmd, *a, **k),
                                    out_path, key, cache)
        rm.subprocess = _SPShim()

        # --- case-insensitive volume-set grouping for mixed-case release names ---
        # Some groups (e.g. LiGHTFORCE) pack a release whose volumes carry
        # INCONSISTENT case — LFC-BFYP.RAR, lfc-bfyp.r00, LFC-BFYP.R02… On a
        # case-insensitive filesystem these are one archive. rescene groups the
        # volume blocks into "sets" via get_set(), and get_archived_file_blocks
        # stops collecting a file's blocks the moment the set name changes.
        # get_set derives the set from the volume name WITHOUT folding case
        # ('LFC-BFYP' vs 'lfc-bfyp'), so the case flip after volume 1 splits one
        # archive in two: rescene then thinks the compressed file fits in a
        # single ~50 MB volume while the source is the full multi-GB file → a
        # permanent, unfixable "size a few bytes off" near-miss (and a futile
        # version/-mt sweep chasing a phantom). Fold case so every volume of one
        # archive shares a set. No-op for a normal consistent-case release.
        _orig_get_set = rm.get_set
        def _get_set_ci(srr_rar_block, _orig=_orig_get_set):
            try:
                return _orig(srr_rar_block).lower()
            except Exception:
                return _orig(srr_rar_block)
        rm.get_set = _get_set_ci

        # --- known-good version cache fronts the candidate order ---
        # Only for the FIRST file's version hunt (archived_files empty). Once
        # rescene has found a good version it puts that first itself for the
        # remaining files — overriding that with our prefs can promote slow
        # RAR5 binaries into per-file hunts (e.g. on an embedded jpg).
        _orig_get = rm.RarRepository.get_rar_executables
        def _get_pref(repo_self, date, _orig=_orig_get, _rm=rm):
            order = list(_orig(repo_self, date))
            # Per-stream version-sweep rescue: force the ONE stream currently
            # being compressed (a proof jpg) to a candidate build while the rest
            # of the set stays pinned to its locked build. Keyed by the source
            # basename (_current_src, set at the top of _crf_init). Checked
            # FIRST so it wins over the set-wide _version_force/_set_good_rar
            # locks — the big content file can stay on the pinned version while
            # the swept extra tries another. Only ever set by the multi-file
            # cross-version tier, so every other path skips this untouched.
            sv = getattr(self, "_stream_ver_override", None)
            if sv:
                want = sv.get(getattr(self, "_current_src", None))
                if want:
                    only = [r for r in order if str(r) == want]
                    if only:
                        return only
            # Build-sweep rescue: restrict the hunt to ONE exact exe (by file
            # name, so a beta can be told apart from its final — rescene's
            # __str__ collapses both to e.g. "2014-05-21 5.11"). Used to retry a
            # CRC near-miss under the sibling build.
            bforce = getattr(SrrdbToolAPI, "_build_force", None)
            if bforce:
                return [r for r in order
                        if getattr(r, "file_name", None) == bforce]
            # Version-sweep rescue: restrict the hunt to ONE forced version so a
            # near-miss re-run tries exactly that build (rescene otherwise stops
            # at the first piece-CRC match and never full-verifies the rest).
            force = getattr(SrrdbToolAPI, "_version_force", None)
            if force:
                return [r for r in order if str(r) == force]
            # Capture the full ordered version list once (first-file hunt, no
            # restriction active) so the version-sweep rescue knows every build
            # available in the pack, in rescene's own naming.
            if not getattr(_rm, "archived_files", None) \
                    and not getattr(self, "_set_good_rar", None):
                self._all_versions = [str(r) for r in order]
            # Fast-fail the embedded-jpg wall: once a version reproduced the
            # first stream of THIS set, every other stream in the same set was
            # produced by the same single WinRAR invocation — and the whole set
            # is reassembled with one version — so only that version can ever
            # reproduce a later stream. If it can't, no pack version will (it's
            # the thread-count/settings wall, not a missing version). Restrict
            # the hunt to it so rescene rejects in seconds instead of churning
            # all 84 versions. Set-scoped: _set_good_rar is cleared per
            # reconstruct() call, so it never leaks across RAR sets that may
            # legitimately use different versions, and it's only ever set AFTER
            # a genuine match — so it can't hurt a rebuild that would succeed.
            good = getattr(self, "_set_good_rar", None)
            if good:
                only = [r for r in order if str(r) == good]
                if only:
                    # A SMALL embedded file (proof jpg) locked FIRST and set the
                    # set version. For another SMALL file, RESTRICT to it (the
                    # fast-fail that stops an embedded-jpg wall wandering all 232
                    # versions). But for the BIG content file, only FRONT that
                    # version — never restrict — so a two-stage PUSSYCAT pack
                    # whose game was made by a DIFFERENT build than its proof jpg
                    # (jpg 5.11 / game 4.11) can still fall through and find the
                    # game's true version instead of dying on "No good version".
                    if not getattr(self, "_current_is_big", False):
                        return only
                    # Big file: front the set version, then this group's
                    # known-good builds, then the rest by date — so a game made
                    # by a DIFFERENT build is found fast, not after grinding the
                    # oldest 2.x versions first.
                    hints = [good] + [v for v in
                                      (getattr(self, "_recon_prefs", None) or [])
                                      if v and v != good]
                    front = [r for r in order if str(r) in hints]
                    front.sort(key=lambda r: hints.index(str(r)))
                    return front + [r for r in order if str(r) not in hints]
            prefs = getattr(SrrdbToolAPI, "_pref_versions", None) or []
            if prefs and not getattr(_rm, "archived_files", None):
                front = [r for r in order if str(r) in prefs]
                front.sort(key=lambda r: prefs.index(str(r)))
                order = front + [r for r in order if str(r) not in prefs]
            # Release-date cap (main first-file hunt only): drop builds newer than
            # release_date + margin and try the rest NEAREST the release date
            # first. Group-history builds (prefs + _recon_prefs) are always kept.
            # Skipped when widened (a prior capped run failed) or no date known.
            cap = getattr(self, "_version_date_cap", None)
            if (cap and not getattr(self, "_version_cap_widen", False)
                    and not getattr(_rm, "archived_files", None)
                    and not getattr(self, "_set_good_rar", None)):
                import datetime
                cutoff = cap + datetime.timedelta(days=_VERSION_CAP_MARGIN_DAYS)
                keepset = set(prefs) | set(getattr(self, "_recon_prefs", None)
                                           or [])
                capped, dropped = [], 0
                for r in order:
                    d = self._version_date(str(r))
                    if str(r) in keepset or d is None or d <= cutoff:
                        capped.append(r)
                    else:
                        dropped += 1
                if dropped:
                    pref_part = [r for r in capped if str(r) in keepset]
                    rest = [r for r in capped if str(r) not in keepset]
                    rest.sort(key=lambda r: -( (self._version_date(str(r))
                                                or datetime.date.min).toordinal()))
                    order = pref_part + rest
                    # Distinct version strings actually searched (betas collapse) —
                    # the target the wall-cache compares _versions_tried against.
                    self._capped_count = len({str(r) for r in order})
            # Resume a hunt a previous run's deadline cut off: builds that run
            # already tested move to the BACK, so this pass spends its 30 minutes
            # on NEW ground instead of repeating the same head of the list. They
            # stay in the order (a truncated run can leave a version partly
            # tested), so nothing is permanently excluded. First-file hunt only —
            # once a version is locked the set-scoped ordering owns the sequence.
            done = getattr(self, "_retry_tried_versions", None)
            if (done and not getattr(_rm, "archived_files", None)
                    and not getattr(self, "_set_good_rar", None)):
                fresh = [r for r in order if str(r) not in done]
                if fresh:
                    order = fresh + [r for r in order if str(r) in done]
            return order
        rm.RarRepository.get_rar_executables = _get_pref

        # --- skip the 1 GB "test with previous file" fallback on non-solid sets ---
        # When a stream's direct match fails, rescene retries by PREPENDING the
        # previous stream and recompressing (main.py:2411-2429) — here the 1 GB
        # .3ds, per thread-count try. In a NON-SOLID archive there is no
        # cross-file dictionary, so the prepend produces byte-identical output
        # for the target file: it can never turn a fail into a success, it just
        # burns minutes. Neutralise it, but ONLY once a version is locked for the
        # set AND every detected stream is provably non-solid (positive .solid
        # check; if any object lacks the flag we leave it untouched). Solid sets,
        # which genuinely need the prepend, are never affected.
        _orig_sefb = rm.RarArguments.set_extra_files_before
        def _sefb(args_self, file_list, _orig=_orig_sefb, _rm=rm):
            if getattr(self, "_set_good_rar", None):
                vals = list((getattr(_rm, "archived_files", None) or {}).values())
                if vals and all(hasattr(v, "solid") and not v.solid for v in vals):
                    return _orig(args_self, [])  # drop the prepend
            return _orig(args_self, file_list)
        rm.RarArguments.set_extra_files_before = _sefb

        # --- drop rescene's inner "more files" append during a -mt rescue ---
        # Inside CompressedRarFile.__init__, a failed solo hunt is retried once
        # with the NEXT file appended (main.py:2156-2159 → set_extra_files_after).
        # For a NON-SOLID member that append can't change the target's bytes
        # (same reasoning as _sefb), so during a forced-mt rescue it can't turn a
        # miss into a hit — drop it. Scoped to rescue mode (_mt_override set) so
        # the normal path is completely unchanged; solid sets keep the append.
        _orig_sefa = rm.RarArguments.set_extra_files_after
        def _sefa(args_self, file_list, _orig=_orig_sefa, _rm=rm):
            if getattr(self, "_mt_override", None):
                vals = list((getattr(_rm, "archived_files", None) or {}).values())
                if vals and all(hasattr(v, "solid") and not v.solid for v in vals):
                    return _orig(args_self, [])  # drop the append
            return _orig(args_self, file_list)
        rm.RarArguments.set_extra_files_after = _sefa

        # --- global -mt PIN + explicit -mt0 (rescene issue #173) ------------
        # Two forced-thread paths share this wrapper:
        #  • _mt_pin (PUSSYCAT jpg+content class): pin EVERY stream — method1
        #    AND method2's all-files pack — to the ONE forced -mt, so the whole
        #    set is built at (forced build, forced -mt), exactly like the recipe
        #    probe's single `rar a jpg nds` command. Crucially it does NOT
        #    disable method2 (unlike _mt_override) — this class NEEDS the
        #    all-files pass to reproduce the in-context content. pin==0 ⇒ -mt0.
        #  • -mt0: rescene's increase_thread_count floors the thread count at 1,
        #    so -mt0 — a DISTINCT algorithm, not -mt1 — is silently skipped and
        #    any release packed with it can't be reconstructed. When our sweep
        #    forces exactly [0], set "-mt0" directly (bypassing the floor).
        # Both are gated on flags ONLY our forced sweeps set, so the normal auto
        # hunt is completely unchanged.
        _orig_inc = rm.RarArguments.increase_thread_count
        def _inc_mt0(args_self, rarbin, _orig=_orig_inc, _rm=rm):
            pin = getattr(self, "_mt_pin", None)
            if pin is not None:
                if not args_self.threads:
                    args_self.threads = "-mt%d" % pin
                    return True
                return False        # only the pinned value; nothing higher
            if list(_rm.RarArguments.mt_settings.mt_set or []) == [0]:
                if not args_self.threads:
                    args_self.threads = "-mt0"
                    return True
                return False        # 0 tried, nothing higher in a [0] set
            return _orig(args_self, rarbin)
        rm.RarArguments.increase_thread_count = _inc_mt0

        # --- disable the "method2" ALL-files fallback during a -mt rescue ---
        # THE expensive path: when a stream's CompressedRarFile fails, rescene
        # falls back to CompressedRarFileAll (main.py:2032), which recompresses
        # EVERY file together ("Compressing ALL files", prepending the 1 GB .3ds)
        # and sweeps thread counts via try_again. For a NON-SOLID member this can
        # never reproduce the target's bytes, so during a forced-mt rescue it's
        # pure waste — fail fast instead so the outer sweep moves to the next mt.
        # Gated on _mt_override, so the normal path keeps method2 untouched.
        _orig_all_init = rm.CompressedRarFileAll.__init__
        def _all_init(all_self, *a, _orig=_orig_all_init, _rm=rm, **kw):
            if getattr(self, "_mt_override", None):
                raise _rm.RarNotFound("method2 disabled during -mt rescue")
            return _orig(all_self, *a, **kw)
        rm.CompressedRarFileAll.__init__ = _all_init

        # --- stream rescene's internal events to the GUI log ---
        _noisy = {rm.MsgCode.BLOCK, rm.MsgCode.RBLOCK, rm.MsgCode.FBLOCK,
                  rm.MsgCode.STORING}
        def _on_rescene_event(e, _self=self, _noisy=_noisy):
            try:
                msg = str(getattr(e, "message", "") or "").strip()
                if not msg or getattr(e, "code", None) in _noisy:
                    return
                if "Cannot create a file when that file already exists" in msg:
                    return
                if msg.startswith("Good RAR version detected"):
                    ver = msg.split(":", 1)[-1].strip()
                    _self._last_good_rar = ver
                    # Lock this version for the rest of the CURRENT set so the
                    # next compressed stream's hunt can fast-fail (see _get_pref).
                    _self._set_good_rar = ver
                # Track every version rescene actually TESTS ("Trying <ver>.")
                # so the version-wall cache can tell a genuine whole-pack miss
                # from a deadline that struck before all versions were tried.
                elif msg.startswith("Trying ") and msg.endswith("."):
                    vt = getattr(_self, "_versions_tried", None)
                    if vt is not None:
                        vt.add(msg[len("Trying "):-1].strip())
                # --- observational: record the winning (version, -mt) per
                # stream. Read-side only — parses rescene's own log messages
                # and touches no reconstruction state. rescene fires the rar
                # command line (carrying the -mt it locked) immediately before
                # its "Compressing X..." line, so latch the mt, then flush a
                # record when the compress line names the stream.
                mt_m = _MT_RE.search(msg)
                if mt_m and ".exe" in msg.lower():
                    _self._pending_mt = int(mt_m.group(1))
                # Latch the exact build (exe) rescene is invoking so the
                # build-sweep rescue can skip it and try the beta/final sibling.
                if ".exe" in msg.lower():
                    em = _EXE_RE.search(msg)
                    if em:
                        _self._last_good_exe = em.group(1)
                cm = _COMPRESS_RE.match(msg)
                if cm and hasattr(_self, "_recon_streams"):
                    _self._recon_streams.append((
                        cm.group(1).strip(),
                        getattr(_self, "_last_good_rar", None),
                        getattr(_self, "_pending_mt", None),
                    ))
                    _self._pending_mt = None
                if len(msg) > 300:
                    msg = msg[:300] + " …[truncated]"
                _self._log(f"    rescene: {msg}", "dim")
            except Exception:
                pass
        rm.subscribe(_on_rescene_event)

        # --- single-file thread-count near-miss rescue ---
        # Wrap CompressedRarFile.__init__ so that, ONLY when rescene has already
        # declared a non-solid stream unrebuildable with "Still not fine :(."
        # (its locked -mt reproduced the version but not the exact packed size),
        # we retry the full compress at other thread counts before giving up.
        #
        # Strictly additive & safe:
        #   • Runs ONLY after the original __init__ raised "Still not fine" — so
        #     any release that rebuilds today is completely untouched (we return
        #     immediately on the first successful __init__).
        #   • Solid streams are left exactly as-is (raise re-propagates), so the
        #     prepend/solid path near the accidental fix is never entered here.
        #   • Only drives RarArguments.mt_settings (thread count) — it never
        #     touches _sefb, _get_pref, or the version lock. The per-set version
        #     restriction still applies, so retries only re-test the ONE locked
        #     version at a different -mt.
        #   • Bounded by the existing per-release deadline + Stop/Skip.
        _orig_crf_init = rm.CompressedRarFile.__init__

        def _crf_init(crf_self, first_block, blocks, src,
                      next_block=None, next_src=None, solid=False,
                      _orig=_orig_crf_init, _rm=rm, _self=self):
            # Record which source file is about to be compressed so _get_pref's
            # per-stream version sweep can restrict THIS stream's version hunt,
            # and whether it's a BIG content file (so the set-wide version lock
            # only FRONTS its version for the big file, never restricts it — a
            # small proof jpg locked FIRST must not pin the game to its build).
            _self._current_src = os.path.basename(src).lower()
            try:
                _self._current_is_big = (
                    os.path.getsize(src) >= _LARGE_STREAM_BYTES)
            except OSError:
                _self._current_is_big = False
            # Method2 all-files rescue: once the big content stream has rebuilt
            # (archived_files non-empty), force the NEXT file's per-file rebuild
            # to raise so rescene's factory (compressed_rar_file_factory) falls
            # through to CompressedRarFileAll — a single `rar a` of ALL files
            # together at the big stream's locked -mt, reproducing trailing
            # embedded extras exactly as the scene pack did (an isolated
            # per-file recompress can't). Gated on the rescue flag so the normal
            # path is untouched; solid sets keep their own prepend path.
            if (getattr(_self, "_force_method2", False) and not solid
                    and len(getattr(_rm, "archived_files", None) or {}) > 0):
                raise ValueError("srrdb: forcing method2 all-files rebuild")
            # Multi-file CRC rescue: if the outer sweep has pinned a specific
            # -mt for THIS stream (keyed by source basename), force it and skip
            # the single-file rescue — a forced-mt failure just means "wrong
            # thread count for this stream", which the outer loop handles by
            # trying the next value. Empty override (the normal case) is a no-op.
            override = getattr(_self, "_mt_override", None) or {}
            forced = override.get(os.path.basename(src).lower())
            if forced is not None:
                _rm.RarArguments.mt_settings = _rm.RarMtSettings()
                _rm.RarArguments.mt_settings.mt_set = [forced]
                try:
                    _orig(crf_self, first_block, blocks, src,
                          next_block, next_src, solid)
                finally:
                    _rm.RarArguments.mt_settings = _rm.RarMtSettings()
                return

            streams = getattr(_self, "_recon_streams", None)
            base = len(streams) if streams is not None else None
            try:
                _orig(crf_self, first_block, blocks, src,
                      next_block, next_src, solid)
                return
            except ValueError as e:
                # Only the NON-SOLID size near-miss qualifies. Solid stays
                # untouched; the Dragon_Ball-style CRC near-miss does NOT raise
                # here (__init__ succeeds, the outer SFV verify catches it), so
                # it never reaches this rescue.
                if solid or "still not fine" not in str(e).lower():
                    raise
                # During a version sweep the whole reconstruction is re-run per
                # forced version; the -mt rescue would multiply that by CAP for
                # every candidate, so skip it and let this version fail fast.
                if getattr(_self, "_in_version_sweep", False):
                    raise
            _self._rescue_mt_near_miss(
                _orig, crf_self,
                (first_block, blocks, src, next_block, next_src, solid),
                _rm, base)

        rm.CompressedRarFile.__init__ = _crf_init

        # --- skip the redundant "more_files" second version hunt on non-solid ---
        # When the first version hunt fails, rescene retries by appending the NEXT
        # source and re-hunting EVERY pack version (main.py:2156-2159). For a
        # NON-SOLID archive the appended file can't change the target's compressed
        # bytes, so that second full grind is guaranteed to reproduce the first's
        # result — pure waste that DOUBLES the wall-clock on a version wall (seen:
        # LEGO_Rock_Band NDS grinding 140 versions twice to the 30-min deadline).
        # Skip it (return no match) for non-solid; solid archives genuinely need
        # the cross-file dictionary, so they're untouched.
        _orig_smre = rm.CompressedRarFile.search_matching_rar_executable
        def _smre(crf_self, block, blocks, thread_count, more_files=False,
                  _orig=_orig_smre):
            if more_files and not getattr(crf_self, "solid", False):
                return None
            return _orig(crf_self, block, blocks, thread_count, more_files)
        rm.CompressedRarFile.search_matching_rar_executable = _smre

        # --- stored-extra method2 seed (jpg-wall rescue) ---------------------
        # A release packed in ONE `rar a jpg nds` command compresses the .nds
        # IN-CONTEXT of the stored proof jpg; WinRAR's -mt pipeline makes that
        # differ from the .nds compressed alone. rescene detects the .nds in
        # ISOLATION (a stored file is never a CompressedRarFile, so it's not in
        # archived_files and the "previous file" test can't run) → no build
        # matches → "No good RAR version found". Its method2 fallback (compress
        # ALL files together, which WOULD reproduce the in-context .nds) refuses
        # to engage because it needs a prior compressed file for its version
        # (main.py: `assert len(archived_files) != 0`, and the factory only
        # falls to method2 when archived_files is non-empty). When the sole
        # compressed file IS the one that failed, that never happens.
        #
        # Fix: wrap the factory so that — ONLY when our rescue armed it with a
        # candidate build (_m2_seed_build) — a detection failure with nothing
        # rebuilt yet seeds archived_files with that build and drives method2.
        # method2 then compresses all files together (in archive order) and
        # sweeps the thread count itself; the outer SFV verify is the arbiter.
        # Pure pass-through when _m2_seed_build is unset, so the normal path and
        # every other rescue are byte-for-byte untouched.
        _orig_factory = rm.compressed_rar_file_factory
        def _factory(block, blocks, src, in_folder, hints,
                     auto_locate_renamed, _orig=_orig_factory, _rm=rm):
            build = getattr(self, "_m2_seed_build", None)
            # Armed rescue: the normal isolated version hunt is ALREADY known to
            # fail for this release (that's what triggered the rescue), so skip
            # it and seed method2 up-front — saving a full redundant pack grind
            # (~10-15 min per big file). Non-solid + nothing rebuilt yet only.
            # Returns None (→ fall through to the untouched normal factory) when
            # the set isn't the stored-extra shape. Pure no-op when unarmed.
            if (build and not (block.flags & block.SOLID)
                    and len(_rm.archived_files) == 0):
                m2 = self._seed_method2(_rm, build, block, blocks, src,
                                        in_folder, hints, auto_locate_renamed)
                if m2 is not None:
                    return m2
            return _orig(block, blocks, src, in_folder, hints,
                         auto_locate_renamed)
        rm.compressed_rar_file_factory = _factory
        return rm

    def _rescue_mt_near_miss(self, orig_init, crf_self, ctor_args, rm, base):
        """Retry a non-solid 'Still not fine' full compress at other thread
        counts (1..CAP), restricted to the already-locked version. Called ONLY
        after orig_init has raised 'Still not fine' — so it never affects a
        release that rebuilds normally. Raises ValueError('Still not fine :(.')
        if no thread count reproduces the exact packed size; on success it
        leaves crf_self fully constructed for the caller.

        base: len(self._recon_streams) captured BEFORE the first (failed)
        __init__, so the per-attempt combo records can be collapsed to just the
        winning one (or the original near-miss restored on total failure)."""
        first_block, blocks, src, next_block, next_src, solid = ctor_args
        streams = getattr(self, "_recon_streams", None)
        orig_records = list(streams[base:]) if base is not None else []
        try:
            cur = crf_self.good_rar.args.thread_count()
        except Exception:
            cur = 1
        # A version that predates -mt (added in RAR 3.60, 2006) IGNORES the
        # thread switch, so every "retry at another -mt" recompresses the file
        # identically — pure churn that can never succeed. Fail fast: this
        # near-miss is from some other setting we can't vary, not a thread count.
        supports_mt = True
        try:
            supports_mt = bool(crf_self.good_rar.supports_setting_threads())
        except Exception:
            pass
        if not supports_mt:
            self._log(
                "    rescene: locked version predates -mt (RAR <3.60) — a "
                "thread-count sweep can't vary anything, so this near-miss is "
                "from another setting. Not rebuildable.", "warn")
            if base is not None:
                del streams[base:]
                streams.extend(orig_records)
            raise ValueError("Still not fine :(.")
        # PIN the hunt to the already-locked build for every retry. This is a
        # THREAD-COUNT near-miss (the version was detected good; only -mt is
        # wrong), so the version is fixed — matching this method's contract.
        # Without the pin each wrong -mt sends rescene wandering the WHOLE pack
        # again: the locked build's data-piece no longer matches at the changed
        # -mt, so it falls through to trying all 232 builds — and _set_good_rar's
        # big-file branch even fronts-then-lists every build, sailing past the
        # release-date cap. A 2009 release then grinds builds up to 2023 × every
        # -mt → the 30-min deadline (seen on Dragon_Quest_9…NDS-Caravan). Pinned,
        # a wrong -mt fails against the one build in ~1s. Same lever the
        # multi-file rescue uses (_version_force); restored in the finally.
        locked_ver = None
        try:
            locked_ver = str(crf_self.good_rar)
        except Exception:
            locked_ver = getattr(self, "_last_good_rar", None)
        _prev_force = getattr(SrrdbToolAPI, "_version_force", None)
        if locked_ver:
            SrrdbToolAPI._version_force = locked_ver
        self._log(
            f"    rescene: near-miss at -mt{cur} — retrying other thread "
            f"counts (up to -mt{_MT_RETRY_CAP}) before giving up…"
            + (f" (pinned to {locked_ver})" if locked_ver else ""), "dim")
        try:
            for n in self._order_mts(n for n in range(0, _MT_RETRY_CAP + 1)
                                     if n != cur):
                if self._stop.is_set() or self._skip.is_set():
                    break
                if time.time() > getattr(self, "_recon_deadline", float("inf")):
                    self._log("    rescene: deadline reached — stopping mt "
                              "retries.", "dim")
                    break
                # Fresh temp dir each attempt: the failed __init__ called close()
                # which rmtree'd the previous working_temp_dir.
                rm.working_temp_dir = rm.get_temp_directory()
                rm.RarArguments.mt_settings = rm.RarMtSettings()
                rm.RarArguments.mt_settings.mt_set = [n]
                try:
                    orig_init(crf_self, first_block, blocks, src,
                              next_block, next_src, solid)
                    self._log(f"    rescene: exact size matched at -mt{n} ✓ "
                              "(final CRC still checked by the SFV verify)", "ok")
                    # Collapse per-attempt combo records to just the winning one.
                    if base is not None:
                        winning = streams[-1] if len(streams) > base else None
                        del streams[base:]
                        if winning is not None:
                            streams.append(winning)
                    return
                except rm.RarNotFound:
                    continue  # this thread count didn't even match the piece
                except ValueError as e2:
                    if "still not fine" not in str(e2).lower():
                        raise
                    continue  # size still off; try the next thread count
                finally:
                    rm.RarArguments.mt_settings = rm.RarMtSettings()
        finally:
            SrrdbToolAPI._version_force = _prev_force
        # Nothing matched — restore the original near-miss record and fail
        # exactly as rescene would have.
        if base is not None:
            del streams[base:]
            streams.extend(orig_records)
        raise ValueError("Still not fine :(.")

    def _clear_produced_volumes(self, out_root: Path) -> int:
        """Delete the produced RAR volumes (the SFV-listed files, excluding the
        SFV/NFO themselves) so a retry reconstruction won't hit rescene's
        'Operation aborted. Archive already exists.' NEVER touches _stored (which
        holds the packed sources) or any other working dir / metadata."""
        names: set = set()
        for pat in ("*.sfv", "*.SFV"):
            for s in out_root.rglob(pat):
                for name, _crc in self._parse_sfv(str(s)):
                    k = name.lower()
                    if not k.endswith((".sfv", ".nfo")):
                        names.add(k)
        removed = 0
        for f in list(out_root.rglob("*")):
            if not f.is_file():
                continue
            if self._VERIFY_SKIP_DIRS & set(f.relative_to(out_root).parts[:-1]):
                continue
            if f.name.lower() in names:
                try:
                    f.unlink(); removed += 1
                except OSError:
                    pass
        return removed

    def _mt_freq_rank(self) -> list:
        """Winning -mt values across all past OK rebuilds, most-frequent first.
        A data-driven prior for the sweeps: our own results DB shows the real
        distribution (mt8 dominates, then 1, 3, 4, …), and the ReScene community
        likewise notes packers use a small set of common thread counts. Ordering
        the -mt sweep by this rank finds the answer far sooner on average — pure
        reorder, coverage unchanged. Empty (natural order) until wins exist."""
        try:
            cnt: dict = {}
            for r in self._load_results():
                if not r.get("ok"):
                    continue
                for c in (r.get("combos") or []):
                    if len(c) >= 3 and c[2] is not None:
                        cnt[int(c[2])] = cnt.get(int(c[2]), 0) + 1
            return [mt for mt, _ in
                    sorted(cnt.items(), key=lambda kv: (-kv[1], kv[0]))]
        except Exception:
            return []

    def _order_mts(self, candidates, front=()) -> list:
        """Order candidate -mt ints by: `front` values first (in given order),
        then our own win-frequency (_mt_freq_rank), then the common scene counts
        (_MT_COMMON), then the odd/rare tail ascending. De-duped. A pure
        reordering of whatever set is passed in — never adds or drops a value, so
        every sweep keeps its exact coverage and just tries the likeliest first."""
        rank = getattr(self, "_mt_rank_cache", None)
        if rank is None:
            rank = self._mt_freq_rank()
            self._mt_rank_cache = rank
        ri = {mt: i for i, mt in enumerate(rank)}
        ci = {mt: i for i, mt in enumerate(_MT_COMMON)}
        cand = list(dict.fromkeys(int(x) for x in candidates))
        out, seen = [], set()
        for mt in front:
            if mt in cand and mt not in seen:
                out.append(mt); seen.add(mt)

        def _key(mt):
            if mt in ri:                       # proven winner — our data first
                return (0, ri[mt])
            if mt in ci:                       # common scene count, not yet won
                return (1, ci[mt])
            return (2, mt)                     # odd/rare — ascending
        rest = sorted((mt for mt in cand if mt not in seen), key=_key)
        return out + rest

    def _rescue_multifile_crc(self, srr_file: str, content_dir: str,
                              out_root: Path):
        """Rescue a multi-file NON-SOLID CRC near-miss: one or more compressed
        streams reproduced the right SIZE but wrong BYTES, so rescene's
        __init__ (which only size-checks, main.py:2216) accepted them and only
        the final volume SFV caught the wrong CRC.

        Three culprits, all handled here:
          • an embedded jpg whose true thread count is ABOVE the one rescene
            greedily locked (the first -mt whose test piece matched the size);
          • a small trailing TEXT file (nfo/diz/txt) packed AFTER the big
            content file. rescene seeds each file's thread hunt from the MAX
            thread count of files already done in the set (main.py:2143-2148),
            so the tiny file INHERITS the big file's high -mt and locks a
            spurious one — while the scene original packed it single-threaded.
          • a trailing extra whose bytes NO isolated thread count reproduces,
            because WinRAR's -mt pipeline compressed it in-context of the whole
            multi-file `rar a` command. For this we drive rescene's method2
            (all files compressed together) as the first tier — see below.

        We sweep the SMALL -mt>1 suspects (the embedded extras), smallest first,
        and NEVER sweep a big content file when a smaller suspect exists — a big
        file's compressed size is reproduced by essentially the one thread count
        already locked, so every other value just fast-fails the size test after
        a minutes-long recompress. The big file is left to re-lock naturally
        (the proven single-suspect path). Text metadata climbs from -mt1 (its
        true count is almost always single-threaded); a binary extra climbs from
        cur_mt+1 (values below already failed the size test) up to a wider cap
        (small files are cheap to retry). Version is pinned to the one the first
        run locked so the re-runs don't drift to a different version that also
        reproduces the big stream but never the suspect. Any OTHER small suspect
        is held at its locked -mt so a passing SFV is unambiguously the swept
        stream.

        With a SINGLE suspect there are no others to hold, so behaviour matches
        the original single-suspect rescue. 0 suspects → no-op.

        Returns the winning verify-result dict on success, else None. Only ever
        called after the SFV verify already failed, so a release that rebuilds
        (or fails) normally today is unaffected."""
        streams = list(getattr(self, "_recon_streams", None) or [])
        if len(streams) < 2:
            return None                 # nothing to compress together / sweep
        suspects = [s for s in streams if s[2] and s[2] > 1]
        # Pin the version the original run locked (e.g. 2014-05-21 5.11). Without
        # this the sweep re-hunts from scratch and can lock a DIFFERENT version
        # (e.g. 5.50) that also reproduces the big stream — at which point the
        # suspect can never match (it was packed with the original version) and
        # every attempt collapses into rescene's method2 prepend. Fronting the
        # pinned version means the big stream locks it first and the fast-fail
        # then restricts the suspect's hunt to that same version.
        pinned = getattr(self, "_last_good_rar", None)
        sizes = self._srr_packed_sizes(srr_file)

        def _sz(name: str) -> int:
            return sizes.get(os.path.basename(name).lower(), 1 << 62)

        def _is_text_meta(name: str) -> bool:
            return os.path.splitext(name)[1].lower() in (
                ".nfo", ".diz", ".txt", ".sfv", ".ini")

        def _is_big(name: str) -> bool:
            return _sz(name) >= _LARGE_STREAM_BYTES

        # A big content stream followed by ≥1 extra means the set was packed
        # with ONE `rar a` command whose trailing extras rescene rebuilds in
        # isolation — the case the method2 tier below repairs. The -mt sweeps
        # additionally need at least one -mt>1 suspect. If neither applies there
        # is nothing this rescue can do, so bail (keeps the no-op guarantee).
        has_big = any(_is_big(s[0]) for s in streams)
        if not suspects and not has_big:
            return None

        # Sweep the likeliest culprit first — the BIGGEST small embedded extra
        # (proof jpg before a tiny nfo/diz). Each attempt costs the same (it
        # recompresses the big content file regardless of which extra is being
        # swept), and a bigger extra is far more likely to be the -mt culprit: a
        # tiny nfo's locked count is usually already correct (WinRAR caps threads
        # by data size), whereas a mid-size jpg often locked a spurious LOWER
        # value than the archive's real -mt. A big CONTENT file (.3ds etc.) is
        # NEVER swept when a smaller suspect exists: its compressed size is
        # reproduced by essentially the one thread count already locked, so every
        # other value fast-fails the size test after a minutes-long recompress.
        have_small = any(not _is_big(s[0]) for s in suspects)
        ordered = sorted(suspects, key=lambda s: _sz(s[0]), reverse=True)
        # The whole set was packed with ONE `rar a -mt<N>` command, so the real
        # thread count is the highest any stream locked (the big content file
        # reveals it — smaller files may cap lower). Try that value FIRST for
        # every suspect: a jpg that piece-locked a low count almost always needs
        # exactly this. Turns a Captain-Toad-style 20-attempt sweep into ~1.
        dominant_mt = max((s[2] for s in streams if s[2]), default=0)
        # Data-driven -mt ordering for every sweep below (computed once here).
        self._mt_rank_cache = self._mt_freq_rank()
        # The sweep needs its OWN wall-clock budget: each attempt calls
        # _srr_reconstruct, which resets self._recon_deadline to 0 on exit, so
        # we can't lean on that here. Budget the whole sweep (each attempt still
        # has its own per-reconstruction deadline inside _srr_reconstruct).
        # Every attempt recompresses the big content stream in full, so a
        # multi-GB .3ds needs proportionally more wall-clock or the -mt sweep
        # can't reach a trailing extra's true count (a 2 GB .3ds is ~3-4 min per
        # attempt). Scale from the base timeout by size, capped so a pathological
        # release can't hang the batch indefinitely.
        big_bytes = max((_sz(s[0]) for s in streams if _is_big(s[0])),
                        default=0)
        budget = _RECON_TIMEOUT_S
        if big_bytes > (1 << 30):     # > 1 GB
            budget = min(int(_RECON_TIMEOUT_S * (big_bytes / (1 << 30))),
                         4 * _RECON_TIMEOUT_S)
        sweep_deadline = time.time() + budget

        def _attempt(ov, label, force_method2=False, build_force=None):
            """One reconstruct+verify pass. `ov` is the -mt override map
            (basename→thread count); `force_method2` instead drives an
            all-files-together rebuild; `build_force` pins one exact exe (by
            file name) so a beta/final sibling can be tried. Returns the passing
            verify dict, or None. The caller owns the Stop/Skip/deadline guards."""
            self._clear_produced_volumes(out_root)
            self._mt_override = dict(ov)
            self._force_method2 = force_method2
            SrrdbToolAPI._build_force = build_force
            self._recon_streams = []          # fresh combos for this attempt
            if pinned and not build_force:
                # RESTRICT the hunt to the set's locked version (not just front
                # it). Every file in the set shares one WinRAR version, so once
                # it's locked the correct version for a swept stream is ALWAYS
                # `pinned`; if `pinned` at the forced -mt doesn't match, no other
                # version will either (it's the wrong thread count). Fronting
                # alone let rescene wander all 84 versions — twice — per failed
                # attempt (brutal on jpg-first sets where nothing pins the hunt
                # early). `_version_force` makes a wrong -mt fail in ~1 s. Not
                # set during the build sweep, which deliberately tries a sibling
                # build (`_build_force` takes precedence in _get_pref anyway).
                SrrdbToolAPI._pref_versions = [pinned]
                SrrdbToolAPI._version_force = pinned
            self._log(f"    {label}…", "dim")
            try:
                rc = self._srr_reconstruct(
                    srr_file, content_dir, str(out_root), log_rar_pack=False)
            finally:
                self._mt_override = {}
                self._force_method2 = False
                SrrdbToolAPI._build_force = None
                SrrdbToolAPI._version_force = None
                SrrdbToolAPI._pref_versions = []
            if not rc.get("ok"):
                return None
            v2 = self._verify_rebuilt_sfv(out_root)
            if v2["checked"] and not v2["bad"]:
                return v2
            return None

        try:
            # ── Method2: all-files-together rebuild ────────────────────────
            # The scene pack ran ONE `rar a -mt<N> f1 f2 …` command. rescene
            # rebuilds each file in ISOLATION, which reproduces the first big
            # stream but often NOT a trailing embedded extra: WinRAR's -mt
            # pipeline compresses a small file differently in-context than
            # alone, so NO isolated thread count ever matches (the whole
            # Thomas/Harvest/Moco/Rayman trailing-jpg family). Drive rescene's
            # own method2 (CompressedRarFileAll): once the big stream locks its
            # version+mt, the NEXT file's per-file rebuild is forced to raise
            # (see _crf_init / self._force_method2) and the factory recompresses
            # ALL files together at that mt — the faithful reproduction. One
            # attempt (the big stream already revealed the count), CRC-verified
            # before we accept it, so a miss simply falls through to the -mt
            # sweeps below and NOTHING that rebuilds today changes.
            #
            # SKIP when the whole set is single-threaded (dominant_mt <= 1):
            # with one thread, in-context == isolated compression, so method2
            # reproduces the EXACT bytes the plain rebuild already failed on —
            # it can never help, and on a multi-GB .3ds it wastes many minutes
            # that the -mt sweep below needs (the trailing extra was a separate
            # `rar a -mtX` add, e.g. Rayman -mt3 / Metroid -mt6).
            if has_big and dominant_mt > 1:
                if self._stop.is_set() or self._skip.is_set():
                    return None
                if time.time() <= sweep_deadline:
                    self._log(
                        "  Multi-file near-miss: trying an all-files-together "
                        "rebuild — rescene method2, the exact single pack "
                        "command" + (f", pinned to {pinned}" if pinned else "")
                        + "…", "dim")
                    r = _attempt({}, "all files together at the archive -mt",
                                 force_method2=True)
                    if r:
                        self._log(
                            "  ✓ Multi-file rescue: all-files-together rebuild "
                            f"CRC-matches the SFV — all {r['checked']} "
                            "volume(s) verified.", "ok")
                        return r

            # ── Unified shared-mt phase ────────────────────────────────────
            # When TWO OR MORE binary extras (e.g. two proof jpgs) share the
            # failing volume, the whole set was packed with ONE `rar a -mt<N>`
            # command, so their true thread count is the SAME value — almost
            # always the archive-wide dominant. rescene locks each small file
            # independently on a piece-SIZE match, which is near-arbitrary for
            # sub-MB files (two ~550 KB jpgs have locked -mt7 and -mt1 in the
            # same set!), so the per-suspect sweep below — which holds the OTHER
            # extra at its mis-locked value, and never even sees an extra that
            # locked -mt1 (excluded from `suspects`) — can never converge. Pin
            # every small binary extra to one shared count and sweep that single
            # value, dominant first; text metadata is pinned to -mt1 (threading
            # never engages on it, so its bytes are thread-count-independent).
            bin_small = [s for s in streams
                         if not _is_big(s[0]) and not _is_text_meta(s[0])]
            if len(bin_small) >= 2:
                cap = _MT_RETRY_CAP_SMALL
                # Dominant first (the whole set's likely shared count), then by
                # global win-frequency — pure reorder of 1..cap.
                shared_order = self._order_mts(
                    range(0, cap + 1),
                    front=([dominant_mt] if dominant_mt and dominant_mt <= cap
                           else ()))
                self._log(
                    f"  Multi-file near-miss: {len(bin_small)} binary extras "
                    "share the volume — trying one shared -mt for all of them "
                    f"(archive -mt{dominant_mt} first"
                    + (f", pinned to {pinned}" if pinned else "") + ")…", "dim")
                for n in shared_order:
                    if self._stop.is_set() or self._skip.is_set():
                        return None
                    if time.time() > sweep_deadline:
                        self._log("  Multi-file rescue: deadline reached — "
                                  "stopping.", "dim")
                        return None
                    ov = {s[0].lower(): (1 if _is_text_meta(s[0]) else n)
                          for s in streams if not _is_big(s[0])}
                    r = _attempt(ov, f"all extras at -mt{n}")
                    if r:
                        self._log(
                            f"  ✓ Multi-file rescue: all binary extras rebuilt "
                            f"at a shared -mt{n} — all {r['checked']} volume(s) "
                            "now CRC-match the SFV.", "ok")
                        return r

            # ── Per-suspect sweep ──────────────────────────────────────────
            if len(suspects) > 1:
                self._log(
                    f"  Multi-file near-miss: {len(suspects)} streams at -mt>1 "
                    "— sweeping the biggest extra first (archive -mt"
                    + (f"{dominant_mt} " if dominant_mt else " ")
                    + "tried first); the big content file re-locks naturally"
                    + (f", pinned to {pinned}" if pinned else "") + "…", "dim")
            for suspect_file, _ver, cur_mt in ordered:
                skey = suspect_file.lower()
                big = _is_big(suspect_file)
                if big and have_small:
                    continue
                # Hold any OTHER SMALL suspect at its locked -mt so a pass is
                # unambiguously the swept stream; big neighbours are left to
                # re-lock naturally (the proven single-suspect path — forcing a
                # 100 MB+ file through the -mt override is both fragile and slow).
                base_pins = {s[0].lower(): s[2] for s in suspects
                             if s[0].lower() != skey and not _is_big(s[0])}
                cap = _MT_RETRY_CAP if big else _MT_RETRY_CAP_SMALL
                # Front the archive's real thread count (dominant_mt) — the
                # single likeliest value for a mid-size extra that piece-locked
                # too low — then order the rest by global win-frequency.
                front = ([dominant_mt] if dominant_mt and dominant_mt != cur_mt
                         and dominant_mt <= cap else [])
                if _is_text_meta(suspect_file):
                    # Inherited a spurious high count; the truth is almost always
                    # single-threaded, so keep -mt1 near the front too.
                    order = self._order_mts(
                        (n for n in range(1, cap + 1) if n != cur_mt),
                        front=front + ([1] if 1 != cur_mt else []))
                else:
                    # The original hunt locks the FIRST -mt whose piece matched
                    # the size (ascending), so every value BELOW cur_mt already
                    # failed the size test — keep those LAST. Above cur_mt is the
                    # live range; order both halves by win-frequency.
                    above = self._order_mts(range(cur_mt + 1, cap + 1),
                                            front=front)
                    below = self._order_mts(range(0, cur_mt))
                    order = above + [n for n in below if n not in above]
                self._log(
                    f"  sweeping -mt for {suspect_file} (locked -mt{cur_mt}, "
                    f"trying -mt{min(order)}–{max(order)})"
                    + (f"; holding {', '.join(sorted(base_pins))}"
                       if base_pins else "") + "…", "dim")
                for n in order:
                    if self._stop.is_set() or self._skip.is_set():
                        return None
                    if time.time() > sweep_deadline:
                        self._log("  Multi-file rescue: deadline reached — "
                                  "stopping.", "dim")
                        return None
                    ov = dict(base_pins)
                    ov[skey] = n
                    r = _attempt(ov, f"trying {suspect_file} at -mt{n}")
                    if r:
                        self._log(
                            f"  ✓ Multi-file rescue: {suspect_file} rebuilt at "
                            f"-mt{n} — all {r['checked']} volume(s) now "
                            "CRC-match the SFV.", "ok")
                        return r

            # ── Sweep -mt1-locked small extras ─────────────────────────────
            # A small extra that locked -mt1 was excluded from `suspects`, but
            # its piece-SIZE lock is UNRELIABLE — many thread counts share a
            # small file's compressed size, so a proof jpg genuinely packed at a
            # higher count (e.g. added after an -mt1 .3ds via a separate
            # `rar a -mtX`) still locks a spurious -mt1 and is never swept. This
            # is the PROVEN fix for the -mt1 .3ds family (Rayman -mt3, Metroid
            # -mt6), so it runs BEFORE the rarer build sweep — and, for an -mt1
            # set, method2 above was skipped, so this is the primary lever and
            # gets the full budget. Sweep each such extra across -mt2..cap (mt1
            # already failed), holding the OTHER small extras at their locked
            # count; the big content file re-locks naturally. CRC-verified.
            mt1_extras = sorted(
                (s for s in streams if not _is_big(s[0])
                 and s[2] == 1 and not _is_text_meta(s[0])),
                key=lambda s: _sz(s[0]), reverse=True)
            for extra_file, _v, _m in mt1_extras:
                skey = extra_file.lower()
                base_pins = {s[0].lower(): s[2] for s in streams
                             if not _is_big(s[0]) and s[0].lower() != skey}
                cap = _MT_RETRY_CAP_SMALL
                self._log(
                    f"  Sweeping -mt for {extra_file} (locked -mt1, unreliable "
                    f"for a small file — trying -mt2–{cap})"
                    + (f"; holding {', '.join(sorted(base_pins))}"
                       if base_pins else "") + "…", "dim")
                for n in self._order_mts([0, *range(2, cap + 1)]):
                    if self._stop.is_set() or self._skip.is_set():
                        return None
                    if time.time() > sweep_deadline:
                        self._log("  Multi-file rescue: deadline reached — "
                                  "stopping.", "dim")
                        return None
                    ov = dict(base_pins)
                    ov[skey] = n
                    r = _attempt(ov, f"trying {extra_file} at -mt{n}")
                    if r:
                        self._log(
                            f"  ✓ Multi-file rescue: {extra_file} rebuilt at "
                            f"-mt{n} — all {r['checked']} volume(s) now "
                            "CRC-match the SFV.", "ok")
                        return r

            # ── Alternate-build sweep ──────────────────────────────────────
            # A CRC near-miss can also mean rescene locked the right version
            # NUMBER but the wrong BUILD of it — a beta vs the final (both are
            # e.g. "2014-05-21 5.11", so nothing above could tell them apart).
            # The big stream can reproduce under both builds while a trailing
            # extra differs. Retry each SIBLING build (same date+major.minor,
            # different exe, minus the one already used); if a plain rebuild
            # misses AND the set is multithreaded, try method2 under it too (a
            # single-threaded set gains nothing from method2). CRC-verified, so
            # a miss simply falls through.
            siblings = self._sibling_builds(
                self._find_rar_dir(), getattr(self, "_last_good_rar", None),
                getattr(self, "_last_good_exe", None))
            if siblings:
                self._log(
                    f"  Multi-file near-miss: locked version has "
                    f"{len(siblings)} sibling build(s) ({', '.join(siblings)}) "
                    "— retrying under each…", "dim")
            for build in siblings:
                if self._stop.is_set() or self._skip.is_set():
                    return None
                if time.time() > sweep_deadline:
                    self._log("  Multi-file rescue: deadline reached — "
                              "stopping.", "dim")
                    return None
                r = _attempt({}, f"rebuild under {build}", build_force=build)
                if not r and dominant_mt > 1 and any(_is_big(s[0])
                                                     for s in streams):
                    if time.time() > sweep_deadline:
                        return None
                    r = _attempt({}, f"all files together under {build}",
                                 force_method2=True, build_force=build)
                if r:
                    self._log(
                        f"  ✓ Multi-file rescue: rebuilt under sibling build "
                        f"{build} — all {r['checked']} volume(s) now CRC-match "
                        "the SFV.", "ok")
                    return r

            # ── Per-stream cross-version sweep (cached big stream) ──────────
            # Last resort for the proof-jpg wall: the big content file rebuilt
            # perfectly (only the volume holding a small binary extra is wrong),
            # NO thread count matched at the locked version, and no sibling build
            # helped. Remaining hypothesis — the extra (a proof jpg) was added by
            # a DIFFERENT WinRAR build than the game, a separate `rar a` the group
            # ran with whatever WinRAR was on the box. Force JUST that extra to
            # each pack build in turn while every big content stream stays pinned
            # to the locked build + its thread count; the big recompress is
            # byte-identical every pass, so the compress cache runs it once (~40 s
            # on a 256 MB .3ds) and replays it for free thereafter. rescene still
            # hunts the extra's own -mt at each candidate build and CRC-checks it,
            # and the full SFV verify guards the result — a miss just falls
            # through. Bounded by the sweep deadline + Stop/Skip.
            bin_extras = [s for s in streams
                          if not _is_big(s[0]) and not _is_text_meta(s[0])]
            allv = [v for v in (getattr(self, "_all_versions", None) or [])
                    if v and v != pinned]
            if has_big and pinned and bin_extras and allv:
                # Order candidates: this group's known-good history first, then
                # release dates nearest the locked build (they share its
                # compression era), then the rest — same priors as the
                # single-file version sweep.
                cand_prefs = [v for v in (getattr(self, "_recon_prefs", None)
                                          or []) if v in allv]
                rest = [v for v in allv if v not in cand_prefs]
                d0 = self._version_date(pinned)
                if d0 is not None:
                    # Keep only era-adjacent builds (±_XVER_ERA_DAYS), nearest
                    # first — a genuine wall then exhausts in minutes instead of
                    # grinding implausibly-distant builds to the deadline.
                    rest = sorted(
                        (v for v in rest
                         if self._version_date(v) is not None
                         and abs((self._version_date(v) - d0).days)
                         <= _XVER_ERA_DAYS),
                        key=lambda v: abs((self._version_date(v) - d0).days))
                cand = (cand_prefs + rest)[:_XVER_MAX_BUILDS]
                # Pin every big content stream to the locked build + its locked
                # thread count so its compress is identical (cache-hittable)
                # across the whole sweep; text metadata rides the pinned build.
                big_mt = {s[0].lower(): s[2] for s in streams
                          if _is_big(s[0]) and s[2]}
                base_ver = {s[0].lower(): pinned for s in streams}
                extra_names = ", ".join(sorted(s[0] for s in bin_extras))
                self._log(
                    f"  Multi-file near-miss: sweeping {extra_names} across "
                    f"{len(cand)} era-adjacent pack build(s) — big stream pinned "
                    f"to {pinned} and cached, nearest release date first…", "dim")
                self._compress_cache = {}
                try:
                    for v in cand:
                        if self._stop.is_set() or self._skip.is_set():
                            return None
                        if time.time() > sweep_deadline:
                            self._log("  Multi-file rescue: deadline reached — "
                                      "stopping.", "dim")
                            return None
                        ver_map = dict(base_ver)
                        for s in bin_extras:
                            ver_map[s[0].lower()] = v
                        self._stream_ver_override = ver_map
                        try:
                            r = _attempt(big_mt,
                                         f"trying {extra_names} under {v}")
                        finally:
                            self._stream_ver_override = None
                        if r:
                            self._log(
                                f"  ✓ Multi-file rescue: {extra_names} rebuilt "
                                f"under {v} — all {r['checked']} volume(s) now "
                                "CRC-match the SFV.", "ok")
                            return r
                finally:
                    self._compress_cache = None
        finally:
            self._mt_override = {}
            self._force_method2 = False
            SrrdbToolAPI._build_force = None
            self._stream_ver_override = None
            cdir = getattr(self, "_ccache_dir", None)
            if cdir:
                shutil.rmtree(cdir, ignore_errors=True)
                self._ccache_dir = None
        self._log("  Multi-file rescue exhausted — no thread count reproduced "
                  "the near-miss stream(s) exactly. Kept as FAILED.", "warn")
        return None

    @staticmethod
    def _sibling_builds(rar_dir, locked_ver, used_exe):
        """Exe file names in the pack that share the locked version's
        date+major.minor but are a DIFFERENT build (beta vs final, or another
        beta), excluding the one already used. Empty when the version has no
        sibling — which is the normal case, so this rescue is a rare no-op."""
        if not rar_dir or not locked_ver:
            return []
        m = re.match(r"\s*(\d{4}-\d{2}-\d{2})\s+(\d+)\.(\d+)", locked_ver)
        if not m:
            return []
        prefix = f"{m.group(1)}_rar{m.group(2)}{m.group(3)}"
        pat = re.compile(re.escape(prefix) + r"(b\d)?\.exe$", re.I)
        try:
            return [f for f in sorted(os.listdir(rar_dir))
                    if pat.match(f) and f != used_exe]
        except OSError:
            return []

    @staticmethod
    def _version_date(ver: str):
        """Parse the leading YYYY-MM-DD out of a rescene version string
        ('2004-07-20 3.30') → date, or None."""
        import datetime
        m = re.match(r"\s*(\d{4})-(\d{2})-(\d{2})", ver or "")
        if not m:
            return None
        try:
            return datetime.date(int(m[1]), int(m[2]), int(m[3]))
        except ValueError:
            return None

    def _release_date(self, release: str, queue_path: str):
        """Best-effort scene RELEASE date for the version cap. Primary source is
        the dats.site dated folder prefix (YYYY-MM-DD-Release…), which is exactly
        the release date; a plain YYYY year prefix maps to that year-end (so the
        +3yr margin still spans the era). Returns a date, or None when no reliable
        date is present (→ no cap, the hunt tries the whole pack as before)."""
        import datetime
        name = Path(queue_path).name if queue_path else ""
        m = re.match(r"^(\d{4})[-._](\d{2})[-._](\d{2})", name)
        if m:
            try:
                return datetime.date(int(m[1]), int(m[2]), int(m[3]))
            except ValueError:
                pass
        m = re.match(r"^(\d{4})[-._]", name)
        if m and 1990 <= int(m[1]) <= 2099:
            return datetime.date(int(m[1]), 12, 31)
        return None

    def _srr_pack_date(self, srr_path: str):
        """Fallback release date from the SRR's own RAR file timestamps — every
        packed-file block carries the pack-time datetime, so a release with no
        dated folder/name can STILL be capped. Returns the LATEST plausible file
        date (an upper bound for 'no WinRAR newer than this'), or None. A noisy
        signal (some groups set rounded/placeholder times), but with the +3yr
        margin and widen-on-re-run it only ever helps — never permanently
        excludes the true version."""
        import datetime
        try:
            from rescene.rar import RarReader, BlockType
            today = datetime.date.today()
            best = None
            for b in RarReader(srr_path):
                if b.rawtype != BlockType.RarPackedFile:
                    continue
                dt = getattr(b, "file_datetime", None)
                if not dt or len(dt) < 3:
                    continue
                try:
                    d = datetime.date(int(dt[0]), int(dt[1]), int(dt[2]))
                except (ValueError, TypeError):
                    continue
                # WinRAR era only, and never a future date (guards junk stamps).
                if 1998 <= d.year <= today.year and (best is None or d > best):
                    best = d
            return best
        except Exception:
            return None

    def _srr_recipe(self, srr_file: str):
        """(level:int, md_param:str, solid:bool) from the SRR's first COMPRESSED
        packed block — the `rar a` recipe to replicate for a probe/dedup. None
        if unreadable or all-stored."""
        try:
            from rescene.rar import RarReader, BlockType  # type: ignore
            level = None
            md = "-mdG"
            solid = False
            for b in RarReader(str(srr_file)).read_all():
                if b.rawtype != BlockType.RarPackedFile:
                    continue
                if (getattr(b, "flags", 0) or 0) & 0x10:
                    solid = True
                if level is None:
                    cp = b.get_compression_parameter()          # '-m5'
                    if cp and cp != "-m0":
                        level = int(cp[2:])
                        dp = b.get_dictionary_size_parameter()   # '-mdG' etc.
                        if dp:
                            md = dp
            return (level, md, solid) if level is not None else None
        except Exception:
            return None

    def _disc_source(self, content_dir: str, out_root: Path):
        """A small-ish COMPRESSIBLE source file to fingerprint pack builds with —
        the largest file ≤4MB (a better discriminator than a tiny nfo, still
        cheap to compress 232×). None if every source is >4MB (skip dedup)."""
        cands = []
        for d in (Path(out_root) / "_stored", Path(content_dir)):
            if d.is_dir():
                for f in d.rglob("*"):
                    if (f.is_file() and f.stat().st_size > 256
                            and f.suffix.lower() not in
                            (".sfv", ".srr", ".srs", ".nfo")):
                        cands.append(f)
        small = [f for f in cands if f.stat().st_size <= 4 * 1024 * 1024]
        if not small:
            return None
        return str(max(small, key=lambda f: f.stat().st_size))

    def _pack_family_reps(self, srr_file: str, content_dir: str, out_root: Path,
                          rar4_only: bool = False):
        """EXE FILENAMES of the pack builds that emit DISTINCT compressed output
        at this release's recipe — one representative per family, keeping a final
        and its betas SEPARATE (they can differ, and the sweep must be able to
        force the EXACT winning exe, not just the version string — 1001's winner
        is rar360.exe while rescene's string sweep grabs rar360b8). RAR3/4 point
        releases are byte-identical so ~232 exes still collapse hard. CACHED per
        (level,md,solid,rar4_only) recipe across releases, so only the FIRST
        stubborn release pays the fingerprint cost. Returns an ORDERED list of
        exe filenames, or None if it can't run (→ string-sweep fallback).

        rar4_only drops the RAR5/6 binaries BEFORE fingerprinting. For a RAR4
        archive they can't produce the format at all, and running ~100 exes that
        have never been executed on this machine is the expensive part (each
        first launch is scanned by the OS) — roughly ten minutes of pure waste."""
        recipe = self._srr_recipe(srr_file)
        if not recipe:
            return None
        level, md, solid = recipe
        cache = getattr(self, "_fam_cache", None)
        if cache is None:
            cache = self._fam_cache = {}
        key = (level, md, solid, rar4_only)
        if key in cache:
            return cache[key]
        disc = self._disc_source(content_dir, out_root)
        rar_dir = self._find_rar_dir()
        exes = sorted(Path(rar_dir).glob("*_rar*.exe")) if rar_dir else []
        if rar4_only:
            r4 = [e for e in exes
                  if re.match(r"\d{4}-\d{2}-\d{2}_rar[0-4]\d", e.name)]
            if r4:
                exes = r4
        if not disc or not exes:
            cache[key] = None
            return None
        try:
            from rescene.rarstream import RarStream  # type: ignore
        except Exception:
            cache[key] = None
            return None
        dname = Path(disc).name
        tmp = Path(tempfile.mkdtemp(prefix="fam-"))
        probe = tmp / "fam.rar"
        seen: set = set()
        reps: list = []
        try:
            for ex in exes:
                if self._stop.is_set() or self._skip.is_set():
                    return None           # don't cache a partial result
                try:
                    probe.unlink(missing_ok=True)
                except Exception:
                    pass
                try:
                    subprocess.run(
                        [str(ex), "a", f"-m{level}", md,
                         "-s" if solid else "-s-", "-ds", "-mt1", "-o+", "-ep",
                         "-idcd", str(probe), disc],
                        capture_output=True, timeout=180)
                except subprocess.TimeoutExpired:
                    continue
                try:
                    with RarStream(str(probe), packed_file_name=dname,
                                   compressed=True) as rs:
                        sig = hashlib.sha1(rs.read()).digest()
                except Exception:
                    sig = ("ERR", ex.name)
                if sig not in seen:
                    seen.add(sig)
                    reps.append(ex.name)   # exact exe — final≠beta kept apart
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        cache[key] = reps or None
        return cache[key]

    def _sweep_build_exes(self, exe_reps: list, srr_file: str, content_dir: str,
                          out_root: Path, detected):
        """Version-near-miss sweep by EXACT exe. `exe_reps` are one filename per
        distinct-output pack family (a final and its betas are SEPARATE), so
        forcing each via `_build_force` tests the exact winning build — which the
        version-STRING sweep can't (it grabs a beta that near-misses while the
        final matches). Drops builds newer than the release cap, orders
        group-history-first then nearest-release, SFV-verifies each. Returns the
        winning verify dict, or None. A wrong exe fails the piece test in ~1s, so
        the sweep stays bounded even with many families."""
        import datetime

        def _date(fn):
            m = re.match(r"(\d{4})-(\d{2})-(\d{2})_rar", fn)
            return datetime.date(int(m[1]), int(m[2]), int(m[3])) if m else None

        def _vstr(fn):
            m = re.match(r"(\d{4}-\d{2}-\d{2})_rar(\d)(\d\d)", fn)
            return f"{m[1]} {m[2]}.{m[3]}" if m else fn

        cap = getattr(self, "_version_date_cap", None)
        # The NEWER cap ALWAYS applies — a group can't use a build that didn't
        # exist yet. (Widen only re-opens the OLDER end, so it must NOT drop the
        # newer cap; otherwise the sweep grinds 2023/2024 builds on a 2014 file.)
        limit = (cap + datetime.timedelta(days=_VERSION_CAP_MARGIN_DAYS)
                 if cap else None)
        cand = [e for e in exe_reps
                if not (limit and _date(e) and _date(e) > limit)] or list(exe_reps)
        prefs = getattr(self, "_recon_prefs", None) or []
        d0 = self._version_date(detected) if detected else None

        def _key(fn):
            vs = _vstr(fn)
            if vs in prefs:
                return (0, prefs.index(vs))
            d = _date(fn)
            return (1, abs((d - d0).days) if (d and d0) else 1 << 30)

        cand.sort(key=_key)
        # Force the -mt too. rescene rebuilds each file at its PIECE-detected
        # thread count, but the piece is often mt-insensitive so it locks a low
        # -mt while the FULL file needs a higher one (1001: jpg locked -mt2, real
        # is -mt8). So per exe we PIN the whole reconstruct's -mt: method1 (the
        # jpg) and method2 (all files together — the in-context .nds) both build
        # at (this exe, this -mt), exactly reproducing the recipe probe's single
        # `rar a jpg nds` command. Only the true (exe, -mt) CRC-matches the SFV;
        # a wrong one fails the piece test in ~1s, so the sweep stays bounded.
        # NB: _mt_pin, NOT _mt_override — the latter disables method2, which this
        # jpg+content class relies on to reproduce the in-context content.
        mts: list = []
        for m in (self._mt_freq_rank() + [8, 4, 2, 1, 3, 6, 5, 7, 0]):
            if m not in mts:
                mts.append(m)
        self._log(
            f"  Version near-miss: '{detected}' matched the test piece but the "
            f"full archive was off — sweeping {len(cand)} distinct-output "
            f"build(s) × -mt {mts} by EXACT exe (final≠beta), nearest first…",
            "dim")
        sweep_deadline = time.time() + _RECON_TIMEOUT_S
        self._in_version_sweep = True
        try:
            for fn in cand:
                if self._stop.is_set() or self._skip.is_set():
                    return None
                if time.time() > sweep_deadline:
                    self._log("  Version rescue: deadline reached — stopping.",
                              "dim")
                    return None
                self._log(f"    trying {_vstr(fn)}  ({fn}) × -mt…", "dim")
                for mt in mts:
                    if self._stop.is_set() or self._skip.is_set():
                        return None
                    if time.time() > sweep_deadline:
                        self._log("  Version rescue: deadline reached — "
                                  "stopping.", "dim")
                        return None
                    self._clear_produced_volumes(out_root)
                    self._recon_streams = []
                    SrrdbToolAPI._build_force = fn
                    self._mt_pin = mt
                    try:
                        rc = self._srr_reconstruct(
                            srr_file, content_dir, str(out_root),
                            log_rar_pack=False)
                    finally:
                        SrrdbToolAPI._build_force = None
                        self._mt_pin = None
                    if not rc.get("ok"):
                        continue
                    v2 = self._verify_rebuilt_sfv(out_root)
                    if v2["checked"] and not v2["bad"]:
                        self._log(
                            f"  ✓ Version rescue: rebuilt with {_vstr(fn)} "
                            f"-mt{mt}  ({fn}) — all {v2['checked']} volume(s) "
                            "CRC-match the SFV.", "ok")
                        return v2
        finally:
            self._in_version_sweep = False
            SrrdbToolAPI._build_force = None
            self._mt_pin = None
        self._log("  Version rescue exhausted — no distinct-output pack build × "
                  "-mt reproduced the archive exactly. Kept as FAILED.", "warn")
        return None

    def _rescue_version_near_miss(self, srr_file: str, content_dir: str,
                                  out_root: Path):
        """Rescue a near-miss where rescene locked a WinRAR version whose
        full-file compressed size was off by a little.

        rescene's version hunt accepts the FIRST build whose test *piece*
        reproduces the block CRC (main.py:2377), then commits to it and never
        tries another. On low-effort methods (-m1) the compressed piece is tiny
        and several early builds share its CRC, so rescene can lock an ancient
        version (e.g. 3.30, 2004) for a release that was really packed years
        later — the full multi-volume archive then differs by a few header
        bytes. rescene has no way out of this; we do: re-run the reconstruction
        forcing each OTHER pack build in turn (nearest release date first, since
        the true build shares the locked one's compression era) and CRC-verify
        against the SFV.

        Bounded by the per-release deadline. The -mt single-file rescue is
        disabled during the sweep (_in_version_sweep) so each candidate fails
        fast instead of multiplying the work by the thread-count cap. Only ever
        called after a near-miss reconstruction ERROR, so a release that
        rebuilds (or fails for another reason) today is untouched.

        Returns the winning verify-result dict on success, else None."""
        detected = getattr(self, "_last_good_rar", None)

        # PRIMARY: sweep SPECIFIC EXES — one per DISTINCT-OUTPUT family, keeping a
        # final apart from its betas. Only this can crack the beta/final class:
        # the version-STRING sweep below hands rescene a string, which resolves
        # to a beta that near-misses while the FINAL matches (1001_Crosswords:
        # winner rar360.exe, rescene picked rar360b8.exe). The recipe probe
        # proved these are crackable exactly this way.
        rar_dir = self._find_rar_dir()
        try:
            exe_reps = self._pack_family_reps(srr_file, content_dir, out_root)
        except Exception:
            exe_reps = None
        if exe_reps and rar_dir:
            return self._sweep_build_exes(
                exe_reps, srr_file, content_dir, out_root, detected)

        # FALLBACK: version-STRING sweep (no fingerprint available) — can't tell a
        # beta from a final, but better than nothing. `_all_versions` lists one
        # entry per exe (betas collapse to one string), so dedupe to unique, drop
        # builds newer than the release cap, group-history first then nearest.
        allv = list(getattr(self, "_all_versions", None) or [])
        if not allv:
            return None
        seen: set = set()
        candidates = [v for v in allv
                      if v != detected and not (v in seen or seen.add(v))]
        if not candidates:
            return None
        cap = getattr(self, "_version_date_cap", None)
        if cap and not getattr(self, "_version_cap_widen", False):
            import datetime
            limit = cap + datetime.timedelta(days=_VERSION_CAP_MARGIN_DAYS)
            candidates = [v for v in candidates
                          if not (self._version_date(v)
                                  and self._version_date(v) > limit)] or candidates
        prefs = [v for v in (getattr(self, "_recon_prefs", None) or [])
                 if v in candidates]
        rest = [v for v in candidates if v not in prefs]
        d0 = self._version_date(detected) if detected else None
        if d0 is not None:
            rest.sort(key=lambda v: (self._version_date(v) is None,
                                     abs((self._version_date(v) - d0).days)
                                     if self._version_date(v) else 1 << 30))
        candidates = prefs + rest
        self._log(
            f"  Version near-miss: '{detected}' matched the test piece but the "
            f"full archive was off — sweeping {len(candidates)} other pack "
            "build(s), nearest release date first…", "dim")
        sweep_deadline = time.time() + _RECON_TIMEOUT_S
        self._in_version_sweep = True
        try:
            for cand in candidates:
                if self._stop.is_set() or self._skip.is_set():
                    return None
                if time.time() > sweep_deadline:
                    self._log("  Version rescue: deadline reached — stopping.",
                              "dim")
                    return None
                self._clear_produced_volumes(out_root)
                self._recon_streams = []
                SrrdbToolAPI._version_force = cand
                self._log(f"    trying {cand}…", "dim")
                try:
                    rc = self._srr_reconstruct(
                        srr_file, content_dir, str(out_root),
                        log_rar_pack=False)
                finally:
                    SrrdbToolAPI._version_force = None
                if not rc.get("ok"):
                    continue
                v2 = self._verify_rebuilt_sfv(out_root)
                if v2["checked"] and not v2["bad"]:
                    self._log(
                        f"  ✓ Version rescue: rebuilt with {cand} — all "
                        f"{v2['checked']} volume(s) CRC-match the SFV.", "ok")
                    return v2
        finally:
            self._in_version_sweep = False
            SrrdbToolAPI._version_force = None
        self._log("  Version rescue exhausted — no pack build reproduced the "
                  "archive exactly. Kept as FAILED.", "warn")
        return None

    def _make_rar_exe(self, rm, build: str):
        """Build a rescene RarExecutable object from the local pack. `build` is
        either an EXACT exe filename ('2005-11-21_rar360.exe' — what the recipe
        sweep returns, so a beta can be told apart from its final) or a version
        STRING like '2012-03-15 4.11'. Used to SEED method2 when no prior
        compressed file exists to borrow a version from. Returns the
        RarExecutable, or None if the exe isn't in the pack / can't be parsed."""
        rar_dir = self._find_rar_dir()
        if not rar_dir:
            return None
        if build.lower().endswith(".exe"):
            fname = build
        else:
            try:
                date, ver = build.split(" ")
                major, minor = ver.split(".")
                fname = f"{date}_rar{major}{minor}.exe"   # final build (no beta)
            except Exception:
                return None
        if not (Path(rar_dir) / fname).exists():
            return None
        try:
            return rm.RarExecutable(rar_dir, fname)
        except Exception:
            return None

    def _seed_method2(self, rm, build, block, blocks, src,
                      in_folder, hints, auto_locate_renamed):
        """Seed rescene's method2 (compress-all-together) with `build` so it can
        run WITHOUT a prior compressed file, and drive it. Returns the
        CompressedRarFileAll, or None if this isn't the stored-extra shape (a
        genuine preceding packed file + ≥2 file blocks) — the caller then uses
        the untouched normal factory. Only invoked with _m2_seed_build armed,
        `block` non-solid, and archived_files empty (see _factory)."""
        fblocks = rm.get_archived_file_blocks(blocks, block)
        file_blocks = [b for b in fblocks if getattr(b, "packed_size", 0)]
        prev = rm.previous_block(block, fblocks)
        # Need something to compress alongside; <2 file blocks ⇒ nothing to do.
        if len(file_blocks) < 2:
            return None
        # Normally the failing file must have a PRECEDING one (stored jpg first /
        # content second) — method2 borrows that file's build. The recipe sweep
        # arms _m2_seed_first because it has PROVEN the recipe, and in its class
        # the leading file is itself in-context (no build reproduces it alone), so
        # method2 has to serve the set from the very first block.
        if prev is block and not getattr(self, "_m2_seed_first", False):
            return None
        exe = self._make_rar_exe(rm, build)
        if exe is None:
            return None
        try:
            exe.args = rm.RarArguments(
                block, os.path.join(rm.get_temp_directory(), "seed.rar"), [src])
        except Exception:
            return None
        # A pinned -mt must be on the args BEFORE method2's first compress:
        # CompressedRarFileAll packs immediately, and an empty `threads` would
        # pack at rar's default count. That pass could size-match by luck and be
        # accepted with the wrong bytes, and at best it's a wasted full compress.
        pin = getattr(self, "_mt_pin", None)
        if pin is not None:
            exe.args.threads = "-mt%d" % pin
        class _Seed:
            solid = False        # keeps the non-solid prepend/append skips valid
        s = _Seed(); s.good_rar = exe
        # Key it OUTSIDE the packed-file namespace when seeding from the first
        # block: rescene fetches a file's data object with
        # archived_files.setdefault(name, factory(...)), so a seed parked under a
        # REAL file name would be handed back as that file's data source (it has
        # no read()) and the rebuild dies. CompressedRarFileAll falls back to
        # "any entry" when the previous block's name isn't a key, which is
        # exactly what a dummy key gives it.
        rm.archived_files["\x00srrdb_m2_seed" if prev is block
                          else prev.file_name] = s
        # Prioritise the group's likely thread counts for method2's internal
        # size sweep (mt8 usually wins for these groups).
        rm.RarArguments.mt_settings = rm.RarMtSettings()
        rm.RarArguments.mt_settings.mt_set = list(
            getattr(self, "_m2_mt_pref", None) or [])
        self._log(
            f"    seeding method2 all-files rebuild with {build} (compressing "
            "the whole set together; skipping the redundant isolated hunt)…",
            "dim")
        m2 = rm.CompressedRarFileAll(
            fblocks, block, blocks, (in_folder, hints, auto_locate_renamed))
        rm.regular_method_failed = m2
        self._m2_seeded = True     # a seed actually happened (see fast-path)
        return m2

    # ── SRR-driven recipe sweep (the in-context wall breaker) ────────────────

    def _srr_sweep_targets(self, srr_file: str) -> dict:
        """Per-RAR-set packing recipe + per-volume stream targets, read from the
        SRR ALONE (no original archives needed).

        {set_prefix: {"order": [name…],            # archive order
                      "files": {name: {"unpacked": int,
                                       "blocks": [(packed_size, file_crc)…]}},
                      "level": int|None, "md": "-mdG", "solid": bool,
                      "dict": bytes, "rar_version": int}}

        The block list is per VOLUME, in order. For a file split across volumes
        every block except the last carries file_crc = CRC32 of THAT volume's
        slice of the compressed stream (RAR4 stores the unpacked-file CRC only in
        the final block) — which is what makes a candidate testable offline."""
        from rescene.rar import RarReader, BlockType  # type: ignore
        sets: dict = {}
        cur = None
        for b in RarReader(str(srr_file)).read_all():
            if b.rawtype == BlockType.SrrRarFile:
                cur = self._rar_set_prefix(getattr(b, "file_name", "") or "")
                sets.setdefault(cur, {"order": [], "files": {}, "level": None,
                                      "md": "-mdG", "solid": False,
                                      "dict": 4 << 20, "rar_version": 0})
            elif b.rawtype == BlockType.RarPackedFile and cur is not None:
                s = sets[cur]
                name = getattr(b, "file_name", "") or ""
                if not name:
                    continue
                if name not in s["files"]:
                    s["order"].append(name)
                    s["files"][name] = {"unpacked": b.unpacked_size,
                                        "blocks": []}
                s["files"][name]["blocks"].append((b.packed_size, b.file_crc))
                if (getattr(b, "flags", 0) or 0) & 0x10:
                    s["solid"] = True
                s["rar_version"] = max(s["rar_version"],
                                       getattr(b, "rar_version", 0) or 0)
                if s["level"] is None:
                    cp = b.get_compression_parameter()
                    if cp and cp != "-m0":
                        try:
                            s["level"] = int(cp[2:])
                        except ValueError:
                            continue
                        s["md"] = b.get_dictionary_size_parameter() or s["md"]
                        s["dict"] = b.get_dict_size() or s["dict"]
        return sets

    def _sweep_set(self, srr_file: str):
        """The RAR set worth sweeping: has a compressed member, and the most
        content behind it.

        Single-file sets count too. They can't be packed in-context, so the
        in-context ARGUMENT for sweeping doesn't apply — but the sweep's other
        property does: it settles build × -mt against the SRR's stream CRCs in
        minutes, where rescene's isolated hunt recompresses the whole file per
        candidate and times out. Excluding them meant Pocoyo_Racing,
        My_Ballet_Studio and Tsumiki_Block_Drop_Mania (one packed file each,
        5-8 volumes, so real stream CRCs available) got no sweep at all and were
        left to the hunt that had already failed them."""
        try:
            sets = self._srr_sweep_targets(srr_file)
        except Exception:
            return None
        best = None
        for name, s in sets.items():
            if s["level"] is None or not s["order"]:
                continue
            tot = sum(f["unpacked"] for f in s["files"].values())
            if best is None or tot > best[1]:
                best = (name, tot, s)
        return best[2] if best else None

    def _locate_sweep_sources(self, order, content_dir: str, out_root,
                              srr_file: str = "") -> dict:
        """{packed name -> source path} for every file in the set, searched in the
        content folder, the SRR's own _stored extras, and the local extras store
        (by the SRR's packed CRC+size). {} when ANY file is missing — the sweep
        needs the exact set the group packed, and locating a renamed or
        downloadable source is the normal path's job."""
        index: dict = {}
        for root in (Path(content_dir), Path(out_root) / "_stored"):
            if not root.is_dir():
                continue
            for f in root.rglob("*"):
                if f.is_file():
                    index.setdefault(f.name.lower(), str(f))
        # A packed extra (proof jpg / diz) is often NOT loose in the content
        # folder — the rebuild pulls it from the local extras store by content
        # CRC while it runs. The sweep happens BEFORE that, so it has to resolve
        # the same way or it bails on releases it could otherwise crack
        # (iCarly…EXiMiUS: only the .nds was loose, so the sweep skipped and the
        # doomed 12-minute hunt ran instead).
        packed = self._srr_packed_info(srr_file) if srr_file else {}
        found: dict = {}
        for name in order:
            key = Path(name).name.lower()
            hit = index.get(key)
            if not hit:
                size, crc_hex = packed.get(key, (None, None))
                if crc_hex is not None:
                    hit = self._extras_lookup(crc_hex, size)
                    if hit:
                        self._log(f"  Recipe sweep: sourcing {Path(name).name} "
                                  f"from the extras store (CRC {crc_hex}).",
                                  "dim")
            if not hit:
                return {}
            found[name] = hit
        return found

    def _recipe_sweep(self, srr_file: str, content_dir: str, out_root) -> dict:
        """Find the (exe, -mt) that reproduces this set's compressed streams by
        packing ALL its files in ONE `rar a` — the way the group did — and
        checking the output against the SRR's per-volume stream CRCs.

        This is the measurement rescene never makes: it hunts each file in
        isolation, which cannot reproduce an in-context pack at ANY build. The
        sweep is decisive in both directions — a hit is byte-level proof of the
        recipe (the caller then pins it and rebuilds), and a clean exhaustion
        proves no pack build can do it, so no grind is worth starting.

        EVERY build is tried — deliberately NOT collapsed into "families" first.
        Family dedup fingerprints one small file, and a small file cannot tell
        apart builds that differ only on larger input: on
        Bravissi-Mots_PROPER…EXiMiUS, 29 builds produced a byte-exact 78 KB jpg
        at -mt8 but only TWO distinct .nds streams, so dedup folded the true
        build (3.90) into the 3.60 family, tested 3.60 as its representative,
        missed by 5 bytes, and reported a "genuine wall". Doing the dedup
        properly would mean compressing the real input per build — which IS the
        sweep, so it saves nothing. Cost is controlled instead by era-filtering,
        truncating sources, and sweeping -mt OUTERMOST so the counts that
        actually win (mt8 leads this dataset) cover every build first: a full
        pass over 135 builds measured ~64 s.

        Returns {"exe", "version", "mt"} or None."""
        s = self._sweep_set(srr_file)
        if not s:
            return None
        # A sweep is DETERMINISTIC in (pack, SRR, sources): re-running one that
        # already exhausted the pack just burns the same minutes again. Honour
        # the cached verdict while the pack is unchanged — and re-sweep the
        # moment it grows, since a new build is exactly what could crack it.
        cached = self._sweep_cache_hit(getattr(self, "_release_name", ""))
        if cached is not None and not self._use_db_history(
                "sweep", "recipe %s recorded %s"
                % ("found" if cached.get("found") else "sweep exhausted",
                   cached.get("ts", "earlier"))):
            cached = None
        if cached is not None:
            if cached.get("found"):
                self._log(
                    f"  Recipe (measured {cached.get('ts', 'earlier')}): "
                    f"{cached['found']['version']} -mt{cached['found']['mt']} "
                    "— reusing it instead of re-sweeping.", "dim")
                return dict(cached["found"])
            self._log(
                f"  Recipe sweep already exhausted the pack for this release "
                f"({cached.get('ts', 'earlier')}, pack unchanged since) — not "
                "repeating it. Add WinRAR versions to retry.", "dim")
            return None
        srcs = self._locate_sweep_sources(s["order"], content_dir, out_root,
                                          srr_file)
        if not srcs:
            # Not a verdict — the sources may be resolved later in the rebuild
            # (a downloaded srrdb "add", a renamed file). Say so, and leave the
            # release eligible for the last-resort sweep once they exist.
            self._log("  Recipe sweep: not every packed file is available yet — "
                      "skipping for now (will retry if the rebuild resolves "
                      "them).", "dim")
            self._sweep_skipped = True
            return None
        try:
            from rescene.rarstream import RarStream  # type: ignore
        except Exception:
            return None
        rar_dir = self._find_rar_dir()
        if not rar_dir:
            return None

        # Candidate builds: EVERY exe of the right era (see the docstring on why
        # family dedup is not used here). A RAR4 archive is swept with the RAR4
        # binaries first; RAR5/6 binaries can still emit RAR4 via -ma4, so they
        # follow as a second phase rather than being dropped outright.
        allexe = sorted(p.name for p in Path(rar_dir).glob("*_rar*.exe"))
        rar4_only = bool(s["rar_version"] and s["rar_version"] < 50)
        if rar4_only:
            era = [n for n in allexe
                   if re.match(r"\d{4}-\d{2}-\d{2}_rar[0-4]\d", n)]
            late = [n for n in allexe if n not in set(era)]
            reps = (era or allexe) + late
        else:
            reps = allexe
        if not reps:
            return None
        # Front this group's known-good builds. Scene groups re-use one packing
        # machine, so the build that cracked their last release is overwhelmingly
        # likely here too — and on a big release the sweep is budget-bound, so
        # WHERE in the order the answer sits decides whether it is found at all
        # (Pokemon_Black_Version_2 and Mon_Coach_Personnel both ran out of budget
        # partway through the pack). Ordering only; nothing is dropped.
        prefs = [v for v in (getattr(self, "_recon_prefs", None) or []) if v]
        if prefs:
            front = [n for n in reps if self._exe_version_str(n) in prefs]
            if front:
                front.sort(key=lambda n: prefs.index(self._exe_version_str(n)))
                reps = front + [n for n in reps if n not in set(front)]
                self._log(f"  Recipe sweep: trying this group's known builds "
                          f"first ({', '.join(prefs[:3])}).", "dim")

        level, md, solid = s["level"], s["md"], s["solid"]
        work = Path(tempfile.mkdtemp(prefix="recipe-"))
        try:
            # No volume-split file means the SRR carries no compressed-stream
            # CRC at all, so sizes are the only evidence. Sizes SHORTLIST and
            # the SFV confirms (see _sweep_by_size) — and there is nothing for
            # the truncation calibration to prove, so skip straight past it.
            if not any(len(s["files"][n]["blocks"]) > 1 for n in s["order"]):
                return self._sweep_by_size(srr_file, content_dir, out_root, s,
                                           srcs, work, rar_dir, reps, RarStream)
            cmd_files, checks = self._calibrated_stage(
                work, s, srcs, level, md, solid, Path(rar_dir), reps, RarStream)
            if cmd_files is None:
                self._sweep_skipped = True
                return None
            n_crc = sum(len(c[1]) for c in checks)
            if not n_crc:
                self._log("  Recipe sweep: the staged set produced no checkable "
                          "stream CRC — skipping rather than guessing.", "dim")
                return None
            mts: list = []
            for m in (self._mt_freq_rank() + list(_MT_COMMON)):
                if m not in mts:
                    mts.append(m)
            self._log(
                f"  Recipe sweep: packing all {len(cmd_files)} file(s) together "
                f"(-m{level} {md} {'-s' if solid else '-s-'}) across "
                f"{len(reps)} build(s) × -mt {mts} — matching against "
                f"{n_crc} stream CRC(s) from the SRR…", "dim")
            # A truncated sweep can only check volume one, so any hit has to be
            # re-proved on the real files against every block the SRR describes.
            # Those files must carry their PACKED names: the extras store keeps
            # content-addressed copies (61ce0998_2188997_xms-mswe.jpg), and `rar
            # a -ep` stores whatever basename it is handed — so passing the raw
            # path archives the wrong name and every check then fails with "File
            # not found in the archive", regardless of build. Link (or copy) any
            # mismatched source under its real name first.
            named = work / "named"
            full_files = []
            for n in s["order"]:
                p = Path(srcs[n])
                if p.name.lower() == Path(n).name.lower():
                    full_files.append(str(p))
                    continue
                named.mkdir(parents=True, exist_ok=True)
                link = named / Path(n).name
                if not link.exists():
                    try:
                        os.link(str(p), str(link))     # instant, same volume
                    except OSError:
                        shutil.copy2(str(p), str(link))
                full_files.append(str(link))
            full_checks = self._full_checks(s)
            truncated = any(
                os.path.getsize(f) != os.path.getsize(srcs[n])
                for n, f in zip(s["order"], cmd_files))
            probe = work / "probe.rar"
            deadline = time.time() + _RECIPE_SWEEP_BUDGET_S
            tried = 0
            for mt in mts:
                for fn in reps:
                    if self._stop.is_set() or self._skip.is_set():
                        return None
                    if time.time() > deadline:
                        self._log("  Recipe sweep: budget reached — stopping "
                                  f"after {tried} combo(s).", "dim")
                        return None
                    tried += 1
                    if not self._sweep_compress(Path(rar_dir) / fn, level, md,
                                                solid, mt, cmd_files, probe):
                        continue
                    if self._sweep_matches(probe, checks, RarStream):
                        if truncated and not self._confirm_recipe(
                                Path(rar_dir) / fn, level, md, solid, mt,
                                full_files, full_checks, work / "confirm.rar",
                                RarStream):
                            self._log(f"    {self._exe_version_str(fn)} -mt{mt} "
                                      "matches volume one but not the whole "
                                      "stream — keeping looking.", "dim")
                            continue
                        hit = {"exe": fn, "version": self._exe_version_str(fn),
                               "mt": mt}
                        # Never let a logging hiccup discard a PROVEN recipe.
                        try:
                            self._log(
                                f"  ✓ Recipe sweep: {hit['version']} -mt{mt}  "
                                f"({fn}) reproduces this set's streams "
                                f"byte-exact ({tried} combo(s) tried).", "ok")
                        except Exception:
                            pass
                        return hit
            self._log(
                f"  Recipe sweep: no pack build × -mt reproduces this set "
                f"({tried} combo(s) across {len(reps)} build(s), RAR4 and "
                "RAR5-in-RAR4-mode) — a genuine wall, not a search-order "
                "problem.", "warn")
            self._sweep_exhausted = True    # a CLEAN miss, safe to cache
            return None
        except Exception as e:
            self._log(f"  Recipe sweep error: {e}", "dim")
            return None
        finally:
            shutil.rmtree(str(work), ignore_errors=True)

    def _calibrated_stage(self, work: Path, s: dict, srcs: dict, level, md,
                          solid, rar_dir: Path, reps: list, RarStream):
        """Stage the sources, then PROVE the truncation is deep enough.

        The cut size can only be estimated from the file's OVERALL ratio, and a
        file whose head compresses better than its average blows that estimate:
        Miffys_World…EXiMiUS packs 67 MB → 40 MB (1.66:1), but its first 11.8 MB
        compress 4.55:1, so the staged input yielded 2,589,751 stream bytes when
        the first volume needs 2,811,143. The check then CANNOT pass at any
        build, and the sweep reports a confident "genuine wall" for a release the
        probe cracks in a minute.

        So: compress once and look. If a truncated file didn't produce enough to
        cover its checked prefix, re-cut using the ratio actually OBSERVED (plus
        margin) and try again; give up on truncating rather than on the release.
        Returns (cmd_files, checks), or (None, None) if it can't be staged."""
        scale = 1.0
        probe = work / "calib.rar"
        ref = None
        for attempt in range(4):
            stage = work / f"s{attempt}"
            cmd_files, checks = self._stage_sweep_sources(work, s, srcs,
                                                          scale=scale,
                                                          into=stage)
            need = [(c[0], c[1][-1][0] + c[1][-1][1]) for c in checks if c[1]]
            if not need:
                return cmd_files, checks      # nothing to calibrate against
            # The reference must be a build that can actually RUN this recipe.
            # Picking blindly took the oldest exe in the pack, and RAR 2.50
            # rejects it outright (rc=7) — the probe then "failed", calibration
            # was skipped, and Miffys_World…EXiMiUS swept 232 builds against a
            # checkpoint its shallow cut could never reach: a 14-minute false
            # wall for a release that rebuilds at 3.60 -mt8.
            if ref is None:
                for cand in reps:
                    if self._sweep_compress(rar_dir / cand, level, md, solid,
                                            1, cmd_files, probe):
                        ref = cand
                        break
                if ref is None:
                    return cmd_files, checks  # no build runs it; sweep will say so
            elif not self._sweep_compress(rar_dir / ref, level, md, solid, 1,
                                          cmd_files, probe):
                return cmd_files, checks
            short = 0.0
            for name, want in need:
                try:
                    with RarStream(str(probe), packed_file_name=name,
                                   compressed=True) as rs:
                        got = len(rs.read())
                except Exception:
                    got = 0
                if got < want:
                    short = max(short, want / max(1, got))
            if not short:
                return cmd_files, checks      # deep enough — sweep for real
            if all(Path(f).stat().st_size == Path(srcs[n]).stat().st_size
                   for n, f in zip(s["order"], cmd_files)):
                return cmd_files, checks      # already full files; nothing more
            scale *= max(1.6, short * 1.35)
            self._log(f"  Recipe sweep: staged sources compressed further than "
                      f"estimated — re-cutting {scale:.1f}× deeper so the "
                      "checked prefix is actually produced.", "dim")
        return cmd_files, checks

    def _stage_sweep_sources(self, work: Path, s: dict, srcs: dict,
                             scale: float = 1.0, into: Path = None):
        """Copy the set's sources into `work` in ARCHIVE ORDER, truncating the
        first oversized split file, and return (command file list, checks).

        checks = [(name, [(offset, length, crc32)…], exact_total|None)…]

        Truncation is safe because a compressed stream's prefix depends only on
        the input prefix — cutting the tail leaves the first volume's slice
        byte-identical (measured) — but only the FIRST oversized file may be cut:
        files after it would be compressed in a different context, so their
        streams are no longer comparable and are dropped from `checks`. The
        preceding files keep full verification, and every file stays in the
        command so the in-context relationship is preserved.

        `scale` deepens the cut when the caller measured that the estimate was
        too shallow (see _calibrated_stage) — the file-average ratio understates
        how well a head that compresses unusually well will pack."""
        wsrc = into or (work / "src")
        wsrc.mkdir(parents=True, exist_ok=True)
        cmd_files: list = []
        checks: list = []
        cut_at = None
        for idx, name in enumerate(s["order"]):
            real = Path(srcs[name])
            size = real.stat().st_size
            blocks = s["files"][name]["blocks"]
            total_packed = sum(b[0] for b in blocks)
            need = size
            if cut_at is None and size > _RECIPE_TRUNC_MIN and len(blocks) >= 2:
                ratio = max(1.0, size / max(1, total_packed))
                need = min(size, int(blocks[0][0] * ratio * 1.4 * scale)
                           + int(s["dict"]) + (1 << 20))
                cut_at = idx
            elif cut_at is not None:
                # Past the cut nothing is verifiable, so a file is only here to
                # keep "another file follows" true. Cap it — a second multi-GB
                # disc would otherwise be copied in full for no information.
                need = min(size, 2 * int(s["dict"]) + (1 << 20))
            dst = wsrc / Path(name).name
            if need < size:
                left = need
                with open(real, "rb") as f, open(dst, "wb") as o:
                    while left > 0:
                        chunk = f.read(min(left, 1 << 22))
                        if not chunk:
                            break
                        o.write(chunk)
                        left -= len(chunk)
            else:
                shutil.copy2(str(real), str(dst))
            cmd_files.append(str(dst))
            if cut_at is not None and idx > cut_at:
                continue                 # past the cut: context differs, unusable
            off, crcs = 0, []
            last = len(blocks) - 1
            for k, (psz, crc) in enumerate(blocks):
                # The final block holds the UNPACKED file's CRC, not a stream
                # CRC, so it is never a byte-level check.
                if k != last and (need == size or off + psz <= blocks[0][0]):
                    crcs.append((off, psz, crc))
                off += psz
            if crcs or need == size:
                checks.append((Path(name).name, crcs,
                               total_packed if need == size else None))
        return cmd_files, [c for c in checks if c[1] or c[2] is not None]

    # RAR 3.x/4.x reject -mt above 16 outright ("Unknown option: mt17"); only
    # RAR5+ binaries accept up to 32. Measured, not assumed.
    _RAR4_MT_MAX = 16
    _R5_EXE = re.compile(r"\d{4}-\d{2}-\d{2}_rar[5-9]\d\d(b\d)?\.exe$", re.I)
    _MD_KB = {"a": 64, "b": 128, "c": 256, "d": 512, "e": 1024, "f": 2048,
              "g": 4096}

    def _sweep_compress(self, exe: Path, level: int, md: str, solid: bool,
                        mt: int, files: list, out: Path) -> bool:
        """One `rar a` of the whole set at (exe, -mt). False when the build can't
        run the recipe at all (pre-mt RAR 2.x rejects -mt, RAR3/4 reject -mt>16,
        old builds reject a 4 MB dictionary) — those fail in milliseconds, which
        is why sweeping the whole pack stays cheap.

        A RAR5/6 binary is asked for RAR4 output with -ma4 (and the -md letter
        rewritten to the size form it wants), so a RAR4 archive packed by a
        modern WinRAR — which is also the only way to get -mt above 16 — stays
        reachable instead of being written off as a wall."""
        is_r5 = bool(self._R5_EXE.search(exe.name))
        if not is_r5 and mt > self._RAR4_MT_MAX:
            return False
        args = [str(exe), "a", f"-m{level}"]
        if is_r5:
            kb = self._MD_KB.get(md[3:].lower()) if len(md) > 3 else None
            args += ["-ma4", f"-md{kb}k" if kb else md]
        else:
            args += [md]
        args += ["-s" if solid else "-s-", "-ds", f"-mt{mt}", "-o+", "-ep",
                 "-idcd", str(out)]
        try:
            for stale in list(out.parent.glob(out.stem + ".*")):
                stale.unlink(missing_ok=True)
        except Exception:
            pass
        try:
            r = subprocess.run(args + files, capture_output=True, timeout=900)
        except (subprocess.TimeoutExpired, OSError):
            return False
        return r.returncode == 0 and out.is_file()

    def _full_checks(self, s: dict) -> list:
        """Checks covering EVERY verifiable byte the SRR describes: each file's
        exact total packed size plus a CRC for every non-final block. Used to
        CONFIRM a candidate on untruncated sources — the truncated sweep can only
        see the first volume's slice, and builds that agree there still diverge
        later (Murder_on_the_Titanic: 3.60 -mt8 matches the first 4.7 MB, then
        runs 55 bytes long over the whole stream; 3.90 -mt8 is the real recipe)."""
        checks = []
        for name in s["order"]:
            blocks = s["files"][name]["blocks"]
            off, crcs = 0, []
            for k, (psz, crc) in enumerate(blocks):
                if k != len(blocks) - 1:
                    crcs.append((off, psz, crc))
                off += psz
            checks.append((Path(name).name, crcs, off))
        return checks

    def _confirm_recipe(self, exe: Path, level, md, solid, mt, full_files,
                        full_checks, probe: Path, RarStream) -> bool:
        """Re-pack the FULL sources at a candidate (exe, -mt) and require every
        block to match. A truncated prefilter hit is a shortlist, never a verdict:
        without this the sweep hands the rebuild a recipe that reproduces volume
        one and nothing after it, and method2 rightly refuses it ("Our options
        are exhausted") after a full compress."""
        if not self._sweep_compress(exe, level, md, solid, mt, full_files, probe):
            return False
        return self._sweep_matches(probe, full_checks, RarStream)

    def _sweep_matches(self, probe: Path, checks: list, RarStream) -> bool:
        """True when the probe archive reproduces every verifiable stream: exact
        total packed size for fully-present files, and CRC32-per-volume-slice for
        every split file's non-final blocks."""
        for name, crcs, exact in checks:
            try:
                with RarStream(str(probe), packed_file_name=name,
                               compressed=True) as rs:
                    data = rs.read()
            except Exception:
                return False
            if exact is not None and len(data) != exact:
                return False
            for off, ln, crc in crcs:
                if len(data) < off + ln:
                    return False
                if zlib.crc32(data[off:off + ln]) & 0xFFFFFFFF != crc:
                    return False
        return True

    @staticmethod
    def _exe_version_str(fname: str) -> str:
        """'2005-11-21_rar360.exe' → '2005-11-21 3.60' (rescene's own naming)."""
        m = re.match(r"(\d{4}-\d{2}-\d{2})_rar(\d)(\d\d)", fname)
        return f"{m.group(1)} {m.group(2)}.{m.group(3)}" if m else fname

    # Distinct size-matching candidates worth a full rebuild. Measured on
    # Jewel_Link…PUSSYCAT: 113 combos reproduce both packed sizes and ALL 113
    # emit byte-identical output, so the real number is 1. The cap is only a
    # guard against a set where it isn't.
    _SIZE_SWEEP_MAX_CANDS = 6

    def _sweep_by_size(self, srr_file: str, content_dir: str, out_root: Path,
                       s: dict, srcs: dict, work: Path, rar_dir: str,
                       reps: list, RarStream) -> dict | None:
        """Recipe sweep for a set with NO volume-split file.

        Every file fits in one volume, so each block's `file_crc` is the
        UNPACKED file's CRC and the SRR describes no compressed-stream bytes at
        all. The only per-file evidence left is the exact packed SIZE — the same
        weak test that makes rescene lock the wrong build, so it is used to
        SHORTLIST and never to decide:

          1. pack all files together at each (build, -mt) and keep the combos
             that reproduce EVERY file's packed size exactly;
          2. dedupe those on the actual compressed bytes. This is safe dedup —
             unlike the family-dedup removed in 208503f, every combo here has
             compressed the real content, so merging is an OBSERVATION that two
             combos are indistinguishable, not a prediction that they will be;
          3. rebuild at each new distinct candidate AS IT IS FOUND, letting the
             SFV arbitrate exactly as on every other path.

        Confirming inline rather than after a full pass matters: sweeping the
        whole grid first cost 701 s on Jewel_Link when the answer was the first
        candidate found. A wrong candidate only costs one rebuild, and combos
        that emit bytes already rejected are skipped without one.

        Returns {"exe", "version", "mt", "rebuilt": (rc, verify)} or None."""
        from rescene.rar import RarReader, BlockType  # type: ignore
        level, md, solid = s["level"], s["md"], s["solid"]
        want = {Path(n).name: sum(p for p, _ in s["files"][n]["blocks"])
                for n in s["order"]}
        # Full, correctly-named sources: sizes must be the REAL ones (no
        # truncation), and `rar a -ep` stores whatever basename it is handed.
        named = work / "named"
        files = []
        for n in s["order"]:
            p = Path(srcs[n])
            if p.name.lower() == Path(n).name.lower():
                files.append(str(p))
                continue
            named.mkdir(parents=True, exist_ok=True)
            link = named / Path(n).name
            if not link.exists():
                try:
                    os.link(str(p), str(link))
                except OSError:
                    shutil.copy2(str(p), str(link))
            files.append(str(link))

        mts: list = []
        for m in (self._mt_freq_rank() + list(_MT_COMMON)):
            if m not in mts:
                mts.append(m)
        self._log(
            f"  Recipe sweep: no volume-split file, so the SRR carries no stream "
            f"CRC — matching all {len(files)} file(s) on exact packed SIZE across "
            f"{len(reps)} build(s) × -mt {mts}, then confirming against the SFV.",
            "dim")

        probe = work / "probe.rar"
        deadline = time.time() + _RECIPE_SWEEP_BUDGET_S
        seen: dict = {}          # output hash -> first (exe, mt) that made it
        cands: list = []
        tried = 0
        for mt in mts:
            for fn in reps:
                if self._stop.is_set() or self._skip.is_set():
                    return None
                if time.time() > deadline:
                    self._log("  Recipe sweep: budget reached — confirming the "
                              f"{len(cands)} candidate(s) found so far.", "dim")
                    break
                tried += 1
                if not self._sweep_compress(Path(rar_dir) / fn, level, md,
                                            solid, mt, files, probe):
                    continue
                got: dict = {}
                try:
                    for b in RarReader(str(probe)).read_all():
                        if b.rawtype == BlockType.RarPackedFile:
                            got[b.file_name] = (got.get(b.file_name, 0)
                                                + b.packed_size)
                except Exception:
                    continue
                if got != want:
                    continue
                h = self._sweep_output_hash(probe, s["order"], RarStream)
                if h is None or h in seen:
                    continue          # byte-identical to one already rejected
                seen[h] = (fn, mt)
                cand = {"exe": fn, "version": self._exe_version_str(fn),
                        "mt": mt}
                cands.append(cand)
                self._log(
                    f"    Distinct output #{len(cands)} reproduces every packed "
                    f"size: {cand['version']} -mt{mt} ({fn}) — rebuilding to let "
                    "the SFV decide.", "dim")
                got = self._rebuild_with_recipe(cand, srr_file, content_dir,
                                                out_root)
                if got:
                    self._log(f"  ✓ Recipe sweep: {cand['version']} -mt{mt} "
                              f"({fn}) rebuilds this set to the SFV "
                              f"({tried} combo(s) tried).", "ok")
                    out = dict(cand)
                    out["rebuilt"] = got
                    return out
                if len(cands) >= self._SIZE_SWEEP_MAX_CANDS:
                    self._log("  Recipe sweep: "
                              f"{self._SIZE_SWEEP_MAX_CANDS} distinct outputs "
                              "matched the sizes and none rebuilt — stopping.",
                              "warn")
                    return None
            if time.time() > deadline:
                break

        if not cands:
            self._log(
                f"  Recipe sweep: no pack build × -mt reproduces even the packed "
                f"SIZES of this set ({tried} combo(s)) — a genuine wall.", "warn")
            self._sweep_exhausted = True
            return None
        self._log(f"  Recipe sweep: {len(cands)} distinct output(s) matched every "
                  "packed size but none rebuilt to the SFV — the sizes agree and "
                  "the bytes don't.", "warn")
        return None

    def _sweep_output_hash(self, probe: Path, order: list, RarStream):
        """SHA-256 over every packed file's COMPRESSED bytes, in order."""
        h = hashlib.sha256()
        try:
            for n in order:
                with RarStream(str(probe), packed_file_name=Path(n).name,
                               compressed=True) as rs:
                    while True:
                        chunk = rs.read(1 << 20)
                        if not chunk:
                            break
                        h.update(chunk)
        except Exception:
            return None
        return h.hexdigest()

    def _rebuild_with_recipe(self, recipe: dict, srr_file: str,
                             content_dir: str, out_root: Path):
        """Rebuild the release at a recipe the sweep already PROVED byte-exact.

        The whole reconstruction is pinned to that one (exe, -mt) and method2 is
        armed, so every stream comes out of a single all-files `rar a` — the same
        command the group ran. Pinning matters as much as the build: with -mt free
        rescene would re-derive a thread count per stream from its size-only test
        and drift off the proven recipe. The SFV verify stays the arbiter, so a
        sweep hit that somehow doesn't rebuild is still reported as a failure."""
        self._clear_produced_volumes(out_root)
        self._recon_streams = []
        self._m2_seeded = False
        SrrdbToolAPI._build_force = recipe["exe"]
        self._mt_pin = recipe["mt"]
        self._m2_seed_build = recipe["exe"]
        self._m2_seed_first = True     # the leading extra is in-context too
        self._m2_mt_pref = [recipe["mt"]]
        try:
            rc = self._srr_reconstruct(str(srr_file), content_dir,
                                       str(out_root), log_rar_pack=False)
        finally:
            SrrdbToolAPI._build_force = None
            self._mt_pin = None
            self._m2_seed_build = None
            self._m2_seed_first = False
            self._m2_mt_pref = None
        if not rc.get("ok"):
            self._log(f"  Recipe rebuild failed: {rc.get('error', 'unknown')}",
                      "warn")
            return None
        v2 = self._verify_rebuilt_sfv(out_root)
        if v2["checked"] and not v2["bad"]:
            self._log(
                f"  ✓ Rebuilt with {recipe['version']} -mt{recipe['mt']} "
                f"({recipe['exe']}) — all {v2['checked']} volume(s) CRC-match "
                "the SFV.", "ok")
            self._last_good_rar = recipe["version"]
            rc["verified"] = v2
            return rc, v2
        self._log(f"  Recipe rebuild produced {v2.get('bad', 0)} mismatched "
                  "volume(s) — not accepted.", "warn")
        return None

    def _rescue_recipe_sweep(self, srr_file: str, content_dir: str,
                             out_root: Path):
        """Measure the true packing recipe, then rebuild at it. The one rescue
        that can crack an IN-CONTEXT pack (one `rar a` over several files), which
        no amount of isolated version hunting can reach. Returns
        (reconstruct rc, verify dict) on success, else None."""
        if self._stop.is_set() or self._skip.is_set():
            return None
        self._sweep_skipped = False
        recipe = self._recipe_sweep(srr_file, content_dir, out_root)
        if recipe is None and self._sweep_skipped:
            # The sweep never actually ran (sources not present yet). That is NOT
            # a verdict, so _recipe_found stays "skip" and the last-resort rescue
            # can try again after the rebuild has resolved the missing extra.
            return None
        if not recipe:
            self._recipe_found = recipe    # recorded in the results DB
            return None
        # The size-only path (no volume-split file — see _sweep_by_size) proves
        # a candidate by REBUILDING it, so a hit there arrives already built.
        rebuilt = recipe.pop("rebuilt", None)
        self._recipe_found = recipe        # recorded in the results DB
        if rebuilt is not None:
            return rebuilt
        return self._rebuild_with_recipe(recipe, srr_file, content_dir, out_root)

    def _is_extra_before_content_shape(self, srr_path: str) -> bool:
        """True when a smaller packed extra precedes the bigger content in a
        NON-SOLID set — the shape that gets packed by ONE `rar a` command and so
        compresses in-context. Unlike _is_stored_extra_shape this does NOT require
        the extra to be STORED: a COMPRESSED leading jpg is the same wall (its own
        stream is in-context too, so no build reproduces it in isolation either),
        and it is exactly the case the recipe sweep settles cheaply."""
        try:
            from rescene.rar import RarReader, BlockType  # type: ignore
            sizes: dict = {}
            order: list = []
            solid = False
            for block in RarReader(str(srr_path)).read_all():
                if block.rawtype != BlockType.RarPackedFile:
                    continue
                name = getattr(block, "file_name", "")
                if not name:
                    continue
                if getattr(block, "flags", 0) & 0x10:
                    solid = True
                if name not in sizes:
                    order.append(name)
                sizes[name] = max(sizes.get(name, 0),
                                  getattr(block, "unpacked_size", 0) or 0)
            if solid or len(order) < 2:
                return False
            return sizes[order[0]] < max(sizes.values())
        except Exception:
            return False

    def _is_stored_extra_shape(self, srr_path: str) -> bool:
        """True when the SRR shows a STORED extra BEFORE the sole compressed
        content (non-solid) — the EXiMiUS shape whose content is compressed
        IN-CONTEXT of a stored proof jpg, so isolated detection is DOOMED
        (grinds the whole pack, sometimes past the 30-min deadline — and a
        deadline timeout never even reaches the post-fail rescue). Detecting it
        up-front lets us skip straight to method2, which reproduces the
        in-context content while the stored jpg is a trivial re-store.

        The leading extra MUST be STORED (-m0). A COMPRESSED leading extra (a
        PUSSYCAT-style proof jpg) is a SEPARATE compression that method2 cannot
        conjure if it's a reproduction wall (Rayman/Yo-Kai class) — those are
        left to the normal hunt + near-miss rescues, so the fast-path never
        wastes doomed method2 passes on them. Heuristic: ≥2 distinct packed
        files, non-solid, FIRST packed file is STORED and not the largest (a
        stored extra precedes the bigger compressed content)."""
        try:
            from rescene.rar import RarReader, BlockType  # type: ignore
            sizes: dict = {}
            order: list = []
            first_stored = None
            solid = False
            for block in RarReader(str(srr_path)).read_all():
                if block.rawtype != BlockType.RarPackedFile:
                    continue
                name = getattr(block, "file_name", "")
                if not name:
                    continue
                if getattr(block, "flags", 0) & 0x10:   # RAR4 SOLID file flag
                    solid = True
                if name not in sizes:
                    order.append(name)
                    if first_stored is None:            # the archive-first file
                        try:
                            first_stored = (
                                block.get_compression_parameter() == "-m0")
                        except Exception:
                            first_stored = False
                sz = getattr(block, "unpacked_size", 0) or 0
                sizes[name] = max(sizes.get(name, 0), sz)
            if solid or len(order) < 2 or not first_stored:
                return False
            return sizes[order[0]] < max(sizes.values())
        except Exception:
            return False

    def _reconstruct_with_m2_fastpath(self, srr_file: str, content_dir: str,
                                      out_root):
        """Fast-path around the main reconstruct: when the SRR has a packed extra
        BEFORE the sole compressed content (see _is_stored_extra_shape), the
        normal isolated version hunt is DOOMED, so seed method2 (compress all
        files together, in-context) at the group's known build(s) FIRST and
        SFV-verify — skipping ~10-25 min of pointless grinding (and the timeout
        failures where the hunt never reaches the rescue). Falls back to the full
        normal reconstruct if the shape doesn't match, there's no group build, or
        no build verifies. SFV verify still guards, so a false pass is
        impossible; correctness is unchanged, only speed.

        Ahead of that, an extra-before-content set (stored OR compressed extra)
        gets the RECIPE SWEEP: it measures the true (build, -mt) against the
        SRR's stream CRCs in seconds and settles the release either way, instead
        of letting the isolated hunt grind to the 30-min deadline on a shape it
        can never match. A miss costs only the sweep and falls through to the
        untouched normal path."""
        if self._is_extra_before_content_shape(srr_file):
            self._log(
                "  A smaller extra precedes the content in a non-solid set — "
                "the group packed them in ONE `rar a`, so every stream is "
                "compressed in-context and the isolated hunt cannot match it. "
                "Measuring the real recipe first…", "dim")
            rescued = self._rescue_recipe_sweep(srr_file, content_dir,
                                                Path(out_root))
            if rescued:
                return rescued[0]
            self._clear_produced_volumes(Path(out_root))
            self._recon_streams = []
        prefs = list(getattr(self, "_recon_prefs", None) or [])
        if not (prefs and self._is_stored_extra_shape(srr_file)):
            return self._srr_reconstruct(str(srr_file), content_dir,
                                         str(out_root))
        mts: list[int] = []
        for m in (self._mt_freq_rank() + [8, 4, 16, 2, 1, 6, 3]):
            if m not in mts:
                mts.append(m)
        self._m2_mt_pref = mts[:8]
        self._log(
            "  A STORED extra precedes the sole compressed file — its content "
            "was compressed IN-CONTEXT, so the isolated hunt can't match it. "
            f"Trying method2 (all-files together) at {prefs} FIRST, skipping the "
            "doomed hunt.", "dim")
        deadline = time.time() + _RECON_TIMEOUT_S
        try:
            for build in prefs:
                if self._stop.is_set() or self._skip.is_set():
                    break
                if time.time() > deadline:
                    break
                self._clear_produced_volumes(out_root)
                self._recon_streams = []
                self._m2_seeded = False
                self._m2_seed_build = build
                self._log(f"    method2 all-files with {build}…", "dim")
                try:
                    rc = self._srr_reconstruct(
                        str(srr_file), content_dir, str(out_root))
                finally:
                    self._m2_seed_build = None
                if rc.get("ok") and rc.get("files"):
                    v2 = self._verify_rebuilt_sfv(out_root)
                    if v2["checked"] and not v2["bad"]:
                        self._log(
                            f"  ✓ method2 fast-path: rebuilt with {build} — all "
                            f"{v2['checked']} volume(s) CRC-match (skipped the "
                            "isolated hunt).", "ok")
                        return rc
                # If seeding never actually happened (e.g. the block-level guards
                # didn't match), don't keep re-running the same doomed path per
                # build — bail to the normal reconstruct.
                if not getattr(self, "_m2_seeded", False):
                    break
        finally:
            self._m2_seed_build = None
            self._m2_mt_pref = None
        self._log("  method2 fast-path didn't rebuild — falling back to the "
                  "normal version hunt…", "dim")
        self._clear_produced_volumes(out_root)
        self._recon_streams = []   # drop the fast-path's stream records so the
        # normal hunt + multi-file rescue see a clean slate (no duplicate combos)
        return self._srr_reconstruct(str(srr_file), content_dir, str(out_root))

    def _rescue_stored_extra_method2(self, srr_file: str, content_dir: str,
                                     out_root: Path):
        """Rescue the STORED-EXTRA + SOLE-COMPRESSED-FILE wall: a release packed
        in ONE `rar a <stored jpg> <compressed nds>` command, where the .nds
        was compressed IN-CONTEXT of the stored extra so rescene's isolated
        detection matches no build → 'No good RAR version found', and its own
        method2 (compress-all-together) can't engage because the failing file is
        the only compressed one (no prior file to borrow a version from).

        We drive method2 ourselves: for each of the group's known-good builds we
        arm the factory seed (_m2_seed_build) so method2 compresses ALL files
        together at that build — sweeping the thread count internally by size —
        then CRC-verify the whole set against the SFV. The SFV verify is the
        arbiter, so a false pass is impossible.

        Only ever called after the normal hunt raised 'No good RAR version
        found' with NO version locked AND the set has ≥2 packed files, so any
        release that rebuilds today is untouched. Needs group history to target
        a build. Returns the winning verify dict on success, else None."""
        if getattr(self, "_last_good_rar", None):
            return None
        # The seed can only engage when the failing compressed file has a
        # PRECEDING packed block to borrow a build from — i.e. the STORED-extra
        # shape this rescue is named for. On any other shape _seed_method2
        # returns None and rescene silently falls back to its normal factory,
        # which re-runs the ENTIRE isolated hunt that just failed — once per
        # candidate build. Measured on FabStyle_JAP_NDS-PUSSYCAT (content-first,
        # 268 MB .nds): ~28 min of identical churn per build, ~110 min for the
        # four group builds, and it consumed the budget the recipe sweep below
        # actually needed. Refusing here costs nothing: those passes never had
        # a way to succeed.
        if not self._is_stored_extra_shape(srr_file):
            self._log(
                "  (method2 rescue not applicable — no stored extra precedes "
                "the content, so there is no build for it to borrow; leaving "
                "the budget to the recipe sweep.)", "dim")
            return None
        builds = list(getattr(self, "_recon_prefs", None) or [])
        if not builds:
            return None                      # no group build to seed method2
        sizes = self._srr_packed_sizes(srr_file) or {}
        if len(sizes) < 2:
            return None                      # need a stored extra + the content
        # Thread-count priority for method2's internal size sweep: this dataset's
        # winning -mt first (mt8 dominates), then a common fallback set.
        mts: list[int] = []
        for m in (self._mt_freq_rank() + [8, 4, 16, 2, 1, 6, 3]):
            if m not in mts:
                mts.append(m)
        self._m2_mt_pref = mts[:8]
        self._log(
            f"  Stored-extra method2 rescue: the sole compressed file failed "
            f"detection but the set has {len(sizes)} packed files — the content "
            f"was likely compressed IN-CONTEXT of a stored extra. Compressing "
            f"ALL files together at {builds} (mt priority {self._m2_mt_pref}).",
            "dim")
        deadline = time.time() + _RECON_TIMEOUT_S
        try:
            for build in builds:
                if self._stop.is_set() or self._skip.is_set():
                    return None
                if time.time() > deadline:
                    self._log("  method2 rescue: deadline reached — stopping.",
                              "dim")
                    return None
                self._clear_produced_volumes(out_root)
                self._recon_streams = []
                self._m2_seed_build = build
                self._log(f"    method2 all-files with {build}…", "dim")
                try:
                    rc = self._srr_reconstruct(
                        srr_file, content_dir, str(out_root),
                        log_rar_pack=False)
                finally:
                    self._m2_seed_build = None
                if not rc.get("ok"):
                    continue
                v2 = self._verify_rebuilt_sfv(out_root)
                if v2["checked"] and not v2["bad"]:
                    self._log(
                        f"  ✓ Stored-extra method2 rescue: rebuilt with {build} "
                        f"(all files compressed together) — all {v2['checked']} "
                        "volume(s) CRC-match the SFV.", "ok")
                    return v2
        finally:
            self._m2_seed_build = None
            self._m2_mt_pref = None
        self._log(
            "  Stored-extra method2 rescue exhausted — no group build "
            "reproduced the in-context content. Kept as FAILED.", "warn")
        return None

    def _log_recon_combos(self, style: str = "dim") -> None:
        """Emit the (stream → version -mt) combos rescene locked this release.

        Observational: it only reads self._recon_streams, captured from
        rescene's own log messages. Useful for spotting which stream/thread
        count caused a near-miss, and as raw material for the .srr2 idea."""
        streams = getattr(self, "_recon_streams", None)
        if not streams:
            return
        parts = []
        for fname, ver, mt in streams:
            mt_txt = f"-mt{mt}" if mt is not None else "-mt?"
            parts.append(f"{fname} → {ver or '?'} {mt_txt}")
        self._log("  Packed with: " + "; ".join(parts), style)

    def _srr_reconstruct(self, srr_path: str, content_dir: str, out_dir: str,
                         log_rar_pack: bool = True) -> dict:
        """Reconstruct RARs using the rescene Python API.

        SRRs can describe multiple RAR sets (movie + Subs + …). Sets whose
        source files are missing from the content folder (typically subtitle
        data, which is not inside the video file) are skipped individually so
        they cannot take the rebuildable sets down with them."""
        import io
        from contextlib import redirect_stdout, redirect_stderr

        Path(out_dir).mkdir(parents=True, exist_ok=True)
        # Cleared per RAR set below; drives the per-set fast-fail in _get_pref.
        self._set_good_rar = None
        # Observational log of (stream, version, -mt) combos the hunt locked —
        # see _on_rescene_event. Diagnostics only. Reset once per release in
        # _process_one, NOT here: a release's nested SRRs re-enter this method,
        # and their combos should accumulate into the same release rather than
        # wipe the main set's before the summary logs it.
        if not hasattr(self, "_recon_streams"):
            self._recon_streams = []
        self._pending_mt = None

        rar_dir = self._find_rar_dir()
        # No dated folder/name? Fall back to the SRR's own RAR timestamps so the
        # version cap can still arm on auto-matched / undated releases. Set once
        # per release (idempotent — a later widen leaves _version_date_cap as-is).
        if (getattr(self, "_date_cap_enabled", True)
                and getattr(self, "_version_date_cap", None) is None
                and not getattr(self, "_version_cap_widen", False)):
            d = self._srr_pack_date(srr_path)
            if d:
                self._version_date_cap = d
                self._date_cap_source = "srr"
                self._log("  Version date-cap: no folder date — using the SRR's "
                          f"RAR timestamp {d.isoformat()} (+"
                          f"{_VERSION_CAP_MARGIN_DAYS // 365}yr, nearest first).",
                          "dim")
        if log_rar_pack:
            if rar_dir:
                rar_exes = [f.name for f in Path(rar_dir).iterdir()
                            if f.is_file() and _RESCENE_RAR_RE.match(f.name)]
                self._log(f"  RAR pack: {len(rar_exes)} version(s) — {', '.join(sorted(rar_exes))}", "dim")
            else:
                self._log("  No rar_X.YY.exe found — use Setup RAR versions button if release is compressed", "dim")

        # Per-release wall-clock deadline: legit game compression can take many
        # minutes, but a truly stuck release (rescene spiralling on an
        # incompressible embedded jpg) must not hang the batch forever. Generous
        # so it only ever fires on a genuine stall.
        self._live_procs = []
        self._recon_deadline = time.time() + _RECON_TIMEOUT_S

        # Heartbeat so large files don't look frozen; also enforces the deadline
        done_flag = threading.Event()
        def _hb():
            secs = 0
            killed = False
            while not done_flag.wait(30):
                secs += 30
                self._log(f"  Still reconstructing… ({secs}s)", "dim")
                if not killed and time.time() > self._recon_deadline:
                    killed = True
                    self._log(
                        f"  ⏱ Exceeded {_RECON_TIMEOUT_S // 60} min — aborting this "
                        "release (likely a stuck/incompressible file). Batch continues.",
                        "err",
                    )
                    for p in list(self._live_procs):
                        try:
                            p.kill()
                        except Exception:
                            pass
        threading.Thread(target=_hb, daemon=True).start()

        def _explain(err: str) -> str:
            if "rar5" in err.lower() or "not yet supported" in err.lower():
                return ("RAR5 data encountered — if the release is RAR5 it cannot be "
                        "rebuilt (pyReScene 0.7 limit). If the SRR info above says RAR4, "
                        "a RAR 5.x+ rar.exe in the pack poisoned the test archive — "
                        "click Setup RAR versions to purge 5.x+ binaries, then retry")
            if "rar executable" in err.lower() or "no rar" in err.lower():
                return err + " — use Setup RAR versions button then retry"
            if "no good rar" in err.lower():
                return err + " — exact WinRAR version not in pack; try adding more versions"
            if "still not fine" in err.lower():
                return ("near-miss: the right WinRAR version was found but the "
                        "recompressed size is a few bytes off — the original was "
                        "packed with a different thread count (-mt) than pyReScene "
                        "can reproduce for a single-file archive. Not rebuildable.")
            return err

        buf = io.StringIO()
        try:
            # Fresh rescene instance per release — rescene stores working state
            # in module-level globals (archived_files, temp dirs, repository)
            # that otherwise persist across reconstruct() calls and poison every
            # following release in a batch. Reloading guarantees a clean slate.
            rm = self._fresh_rescene()

            # Work out which RAR sets have their source content available
            try:
                rar_sets = self._srr_rar_sets(srr_path)
            except Exception:
                rar_sets = {}
            content_names = {
                f.name.lower() for f in Path(content_dir).rglob("*") if f.is_file()
            }

            # Some groups pack the NFO/JPG INSIDE the RARs. Those exact files
            # are stored in the SRR and already extracted to _stored — hand
            # them to rescene as sources via hints (absolute paths work:
            # os.path.join drops in_folder for absolute hint values).
            hints: dict = {}
            stored_pool = Path(out_dir) / "_stored"
            if rar_sets and stored_pool.is_dir():
                stored_files = {f.name.lower(): f
                                for f in stored_pool.rglob("*") if f.is_file()}
                for info in rar_sets.values():
                    for p in info["packed"]:
                        nm = Path(p).name.lower()
                        if nm not in content_names and nm in stored_files:
                            hints[p] = str(stored_files[nm])
                if hints:
                    self._log(
                        "  Using SRR-stored file(s) as sources: "
                        + ", ".join(sorted(Path(k).name for k in hints)), "dim",
                    )

            # Extras packed INSIDE the rars but NOT stored in the SRR (srrdb
            # "adds" — e.g. an unconfirmed proof jpg) are still needed as rebuild
            # SOURCES. Fetch them from srrdb by add-id, CRC-verify, and drop into
            # _stored so they behave exactly like SRR-stored sources. Cached on
            # disk → a re-run needs no network. First run needs network; if the
            # network is unavailable / bot-wall blocks it, it's skipped cleanly with a warning
            # and reconstruction proceeds (and fails as before) — never worse.
            if rar_sets:
                # Expected packed (unpacked) sizes — to spot text sources whose
                # SRR-stored copy differs from the copy actually packed.
                try:
                    expected_sizes = self._srr_packed_sizes(srr_path)
                except Exception:
                    expected_sizes = {}
                still_missing = []   # packed source absent from content AND hints
                wrong_size   = []    # text source present but the WRONG size — the
                                     # packed copy differs (e.g. ABSTRAKT nfo: the
                                     # SRR stores the release nfo, but a DIFFERENT
                                     # nfo was packed and uploaded to srrdb adds).
                for info in rar_sets.values():
                    for p in info["packed"]:
                        nm = Path(p).name.lower()
                        if nm in content_names:
                            continue
                        if p not in hints:
                            still_missing.append((p, nm))
                        elif nm.endswith((".nfo", ".diz", ".txt")):
                            exp = expected_sizes.get(nm)
                            try:
                                cur = Path(hints[p]).stat().st_size
                            except Exception:
                                cur = None
                            if exp is not None and cur is not None and cur != exp:
                                wrong_size.append((p, nm))
                # Local extras store first — content-addressed by the SRR's exact
                # packed CRC, offline, and it spares srrdb. Whatever it resolves is
                # removed from the missing/wrong lists so only the rest goes to the
                # network (and the LE-fix / fast-fail below). Inert with no store.
                if (still_missing or wrong_size) and self._extras_db_path.exists():
                    packed_info = self._srr_packed_info(srr_path)
                    r_miss = self._resolve_from_extras(
                        still_missing, False, hints, stored_pool, packed_info)
                    r_wrong = self._resolve_from_extras(
                        wrong_size, True, hints, stored_pool, packed_info)
                    if r_miss:
                        still_missing = [x for x in still_missing if x not in r_miss]
                    if r_wrong:
                        wrong_size = [x for x in wrong_size if x not in r_wrong]
                # Wrong-match guard: a MAIN content source (the ROM/main file, not
                # a small extra) that's absent AND has no same-SIZE file on disk
                # for rescene to auto-locate means this folder isn't the release's
                # content at all — e.g. a stale queue entry pairing a PS5 name with
                # a 3DS folder (ASTRO.BOT → LEGO_Ninjago). rescene would otherwise
                # die mid-rebuild with a cryptic "file does not exist". A renamed
                # main file (same size, different name) is NOT flagged — rescene
                # auto-locates it — so this never breaks a legit renamed rebuild.
                if still_missing:
                    _EXTRA_EXTS = (".nfo", ".diz", ".txt", ".sfv", ".jpg", ".jpeg",
                                   ".png", ".gif", ".ini", ".m3u", ".srs")
                    try:
                        disk_sizes = {f.stat().st_size for f in
                                      Path(content_dir).rglob("*") if f.is_file()}
                    except OSError:
                        disk_sizes = set()
                    orphan_main = [
                        (p, nm) for (p, nm) in still_missing
                        if os.path.splitext(nm)[1].lower() not in _EXTRA_EXTS
                        and expected_sizes.get(nm) is not None
                        and expected_sizes.get(nm) not in disk_sizes]
                    if orphan_main:
                        _n = ", ".join(Path(p).name for p, _ in orphan_main[:3])
                        self._log(
                            f"  ⚠ Main content not found in this folder: {_n}"
                            + ("…" if len(orphan_main) > 3 else "")
                            + " — wrong release match for this content, or the "
                            "content simply isn't here. Skipping (can't rebuild "
                            "without the main file).", "warn")
                        raise RuntimeError(
                            "main content missing — wrong release match")
                if len(still_missing) > _MAX_FETCH_ADDS:
                    # Too many missing sources to be "extras" — a wrong match or a
                    # loose-asset release. One summary line, no per-file logging
                    # (which floods the UI) and no add-fetch hammering.
                    _nm = [Path(p).name for p, _ in still_missing]
                    self._log(
                        f"  ⚠ {len(_nm)} packed source(s) are missing from the "
                        f"content folder (e.g. {', '.join(_nm[:4])}…) — looks like "
                        "the wrong release match for this content, or a release "
                        "whose loose files aren't present. Not rebuildable; "
                        "skipping add-fetch.", "warn")
                elif still_missing or wrong_size:
                    wrong_set = {p for p, _ in wrong_size}
                    to_fetch = still_missing + wrong_size
                    release_nm = Path(srr_path).stem
                    details = self._srrdb_details(release_nm)
                    raw_adds = (details or {}).get("adds", []) if details else []
                    # Key by BASENAME. srrdb serves reconstruction adds from
                    # subfolders, e.g. "[for reconstruction]/ctr-h-anfj.jpg", so
                    # the API name carries a path that must not defeat the match
                    # against the packed file's bare name — but the FULL name is
                    # kept (a["name"]) for the download URL, which needs the path.
                    adds = {}
                    for a in raw_adds:
                        nm_a = a.get("name")
                        if nm_a and a.get("id"):
                            adds.setdefault(Path(nm_a).name.lower(), a)
                    if not details:
                        self._log("  ⚠ Couldn't fetch srrdb release details to "
                                  "locate the missing/mismatched packed source(s) — "
                                  "no network / rate-limited? Rebuild may fail until "
                                  "available.", "warn")
                    elif not adds:
                        self._log("  ⚠ srrdb lists no fetchable 'adds' for this "
                                  "release.", "warn")
                    else:
                        _names = sorted(str(a.get("name", "?")) for a in raw_adds)
                        self._log(f"  srrdb adds available: {len(_names)} "
                                  f"({', '.join(_names[:6])}"
                                  + ("…" if len(_names) > 6 else "") + ")", "dim")
                    fetched = []
                    no_match = []
                    for p, nm in to_fetch:
                        is_wrong = p in wrong_set
                        a = adds.get(nm)
                        if not a:
                            # Genuinely-missing source with no add can't rebuild; a
                            # wrong-size one falls back to the line-ending fix below.
                            if not is_wrong:
                                no_match.append(Path(p).name)
                            continue
                        # Wrong-size sources go to _stored/_adds so re-extraction
                        # (which restores the mismatched SRR copy each run) can't
                        # clobber the correct fetched copy; both are cache-checked.
                        dest = (stored_pool / "_adds" / Path(p).name if is_wrong
                                else stored_pool / Path(p).name)
                        if dest.is_file():
                            hints[p] = str(dest); fetched.append(Path(p).name); continue
                        label = ("mismatched source, fetching exact packed copy"
                                 if is_wrong else "missing source not in SRR")
                        self._log(f"  {label}: {Path(p).name}…", "dim")
                        res = self._download_add(
                            details.get("_resolved_name", release_nm),
                            a["id"], a["name"], dest, a.get("crc"))
                        if res["ok"]:
                            hints[p] = str(dest); fetched.append(Path(p).name)
                            self._log(f"    ✓ fetched {Path(p).name} "
                                      f"({res['size']:,} B, CRC verified)", "ok")
                            # Harvest this CRC-verified add (esp. file_id.diz) so
                            # the group's other releases resolve it offline.
                            self._harvest_extra(dest, Path(p).name)
                        else:
                            self._log(f"    ✗ {Path(p).name}: {res['error']}", "warn")
                    if no_match:
                        self._log(f"  ⚠ No srrdb add for {len(no_match)} missing "
                                  f"source(s): {', '.join(no_match)} — can't rebuild "
                                  "without them.", "warn")
                    if fetched:
                        self._log("  Using srrdb add(s) as sources: "
                                  + ", ".join(sorted(fetched)), "dim")

            # Correct line-ending-mismatched text sources (nfo/diz) so their
            # size matches the packed copy — see _fix_text_source_endings.
            declined_text = self._fix_text_source_endings(
                hints, srr_path, out_dir) if hints else []
            # Fast-fail: a wrong-size text source (nfo/diz) that couldn't be
            # fetched as a srrdb add AND isn't a clean CRLF/LF variant is
            # genuinely unreconstructable — the group packed a different copy
            # than the SRR stored, and it exists nowhere fetchable. rescene
            # would happily recompress the big content file for minutes, then
            # bail on the tiny text block with "Data file is not the correct
            # size". Skip that wasted work and fail now with a clear reason.
            #
            # STRICTLY SAFE: `declined` only lists sources that are already the
            # wrong size on disk (the exact check rescene's _repack does) and
            # that our add-fetch could not resolve — so this can never abort a
            # release that would otherwise rebuild. Gated to single-set releases
            # so it never short-circuits a multi-set SRR where another set could
            # still reconstruct.
            if declined_text and len(rar_sets) <= 1:
                names = ", ".join(declined_text)
                return {
                    "ok": False, "files": [], "output": "",
                    "error": (
                        f"packed {names} differs from the SRR-stored copy and no "
                        "exact copy is available on srrdb — not rebuildable "
                        "(skipped the content recompress).")}

            skip_parts: list[str] = []
            run_parts:  list[str] = []
            if len(rar_sets) > 1:
                self._log(f"  SRR describes {len(rar_sets)} RAR sets:", "dim")
                for prefix, info in rar_sets.items():
                    missing = [p for p in info["packed"]
                               if Path(p).name.lower() not in content_names
                               and p not in hints]
                    if missing:
                        skip_parts.append(prefix)
                        self._log(
                            f"    ✗ {prefix} — skipped, source not in content folder: "
                            f"{', '.join(Path(m).name for m in missing[:4])}"
                            + ("…" if len(missing) > 4 else ""), "warn",
                        )
                    else:
                        run_parts.append(prefix)
                        self._log(f"    ✓ {prefix} ({len(info['volumes'])} volume(s))", "dim")
                if not run_parts:
                    return {"ok": False, "files": [], "output": "",
                            "error": "No RAR set has its source files in the content folder"}

            base_kwargs: dict = dict(
                srr_file=str(srr_path),
                in_folder=str(content_dir),
                out_folder=str(out_dir),
                extract_files=False,
                # Content files are often renamed (hash names, year fixes).
                # Only consulted when the stored name is missing on disk:
                # falls back to matching by file size + extension.
                auto_locate_renamed=True,
                hints=hints,
            )
            if rar_dir:
                base_kwargs["rar_executable_dir"] = rar_dir

            errors: list[str] = []
            ok_any = False
            # One call per reconstructable set (srr_part prefix wildcard), or a
            # single plain call when the SRR has just one set.
            for part in (run_parts if skip_parts else [None]):
                kwargs = dict(base_kwargs)
                if part is not None:
                    kwargs["srr_part"] = f"{part}.*"
                # New set = new WinRAR invocation, possibly a different version.
                # Reset so the fast-fail never carries one set's version into
                # another (rescene's archived_files persists across these calls).
                self._set_good_rar = None
                try:
                    with redirect_stdout(buf), redirect_stderr(buf):
                        result = rm.reconstruct(**kwargs)
                    if result is False:
                        errors.append(_explain(
                            buf.getvalue().strip() or "Reconstruction failed (no reason given)"))
                    else:
                        ok_any = True
                except Exception as e:
                    errors.append(_explain(str(e)))

            output_text = buf.getvalue()
            # Recursive: multi-CD releases reconstruct into CD1/, CD2/ …
            # subfolders (stored paths). Exclude our own working dirs.
            _SKIP_DIRS = {"_stored", "_subs_tmp", "_iso_m2ts_tmp", "Sample"}
            out_files = [
                str(f.relative_to(out_dir)) for f in Path(out_dir).rglob("*")
                if f.is_file() and not (_SKIP_DIRS & set(f.relative_to(out_dir).parts[:-1]))
            ]
            if ok_any:
                return {"ok": True, "files": out_files, "output": output_text,
                        "error": "; ".join(errors) if errors else None}
            return {"ok": False, "files": out_files, "output": output_text,
                    "error": "; ".join(errors) or "Reconstruction failed"}
        except ImportError:
            return {"ok": False, "files": [], "output": "", "error": "rescene not available — pip install pyReScene"}
        except Exception as e:
            return {"ok": False, "files": [], "output": buf.getvalue(), "error": _explain(str(e))}
        finally:
            done_flag.set()
            self._recon_deadline = 0
            self._live_procs = []

    def _reconstruct_nested(self, nsrr: Path, content_dir: str, out_root: Path,
                            stored_dir: Path) -> dict:
        """Rebuild one nested SRR (vobsubs etc.). Sources are assembled by exact
        filename from the content folder and anything already produced in the
        output (inner subs RARs feed outer ones) into a small temp pool —
        subtitle files are tiny, so copying is cheap."""
        try:
            sets = self._srr_rar_sets(str(nsrr))
        except Exception as e:
            return {"ok": False, "error": str(e)}
        needed = {Path(p).name.lower()
                  for info in sets.values() for p in info["packed"]}
        if not needed:
            return {"ok": False, "error": "no packed files described"}
        # If every RAR this SRR describes already exists (content or output),
        # there is nothing to rebuild — common for inner vobsub RARs that were
        # kept alongside the content.
        vol_names = {Path(v).name.lower()
                     for info in sets.values() for v in info["volumes"]}
        existing = {
            f.name.lower()
            for root in (Path(content_dir), out_root)
            for f in root.rglob("*") if f.is_file()
        }
        if vol_names and vol_names <= existing:
            return {"ok": True, "produced": [], "already": True}
        pool = out_root / "_subs_tmp"
        pool.mkdir(parents=True, exist_ok=True)
        try:
            missing = set(needed)
            for root in (Path(content_dir), out_root):
                if not missing:
                    break
                for f in root.rglob("*"):
                    if (f.is_file() and f.name.lower() in missing
                            and pool not in f.parents):
                        shutil.copy2(str(f), str(pool / f.name))
                        missing.discard(f.name.lower())
            if missing:
                return {"ok": False, "error":
                        "source not found: " + ", ".join(sorted(missing)[:4])}
            rel_parent = nsrr.parent.relative_to(stored_dir)
            n_out = out_root / rel_parent
            rc = self._srr_reconstruct(str(nsrr), str(pool), str(n_out),
                                       log_rar_pack=False)
            if rc["ok"]:
                produced = [f for f in rc["files"]
                            if Path(f).suffix.lower() not in META_EXTS]
                if produced:
                    return {"ok": True, "produced": produced}
                return {"ok": False, "error": "nothing produced"}
            return {"ok": False, "error": rc.get("error") or "failed"}
        finally:
            shutil.rmtree(str(pool), ignore_errors=True)

    def _srr_list(self, srr_path: str) -> bool:
        """Log SRR contents. Returns True if RAR5 format (reconstruction not supported)."""
        try:
            from rescene.rar import (RarReader, BlockType, COMPR_STORING,  # type: ignore
                                     RAR5_MARKER_BLOCK)
            stored, rars = [], []
            content: dict[str, tuple[int, bool]] = {}
            has_compressed = False
            is_rar5 = False

            for block in RarReader(str(srr_path)).read_all():
                bt = block.rawtype
                if bt == BlockType.SrrHeader:
                    self._log(f"    Creating app: {getattr(block, 'appname', '?')}", "dim")
                elif bt == BlockType.SrrStoredFile:
                    stored.append(getattr(block, "file_name", "?"))
                elif bt == BlockType.SrrRarFile:
                    rars.append(getattr(block, "file_name", "?"))
                    # Peek at embedded RAR header to detect RAR5 format
                    raw = getattr(block, "_file_data", None) or getattr(block, "data", b"")
                    if raw and raw[:8] == RAR5_MARKER_BLOCK:
                        is_rar5 = True
                elif bt == BlockType.RarPackedFile:
                    fname  = getattr(block, "file_name", "?")
                    size   = getattr(block, "unpacked_size", 0)
                    method = getattr(block, "compression_method", COMPR_STORING)
                    compressed = (method != COMPR_STORING)
                    if compressed:
                        has_compressed = True
                    if fname not in content:
                        content[fname] = (size, compressed)

            if stored:
                self._log(f"    Stored ({len(stored)}): {', '.join(stored)}", "dim")
            if rars:
                self._log(f"    RARs ({len(rars)}): {', '.join(rars[:5])}"
                          + (f" …+{len(rars)-5}" if len(rars) > 5 else ""), "dim")
            if content:
                parts = []
                for fname, (sz, comp) in list(content.items())[:4]:
                    parts.append(f"{fname} ({sz:,} B, {'COMPRESSED' if comp else 'stored'})")
                self._log(f"    Content ({len(content)} unique): {'; '.join(parts)}"
                          + (f" …+{len(content)-4}" if len(content) > 4 else ""), "dim")
                if is_rar5:
                    self._log("    ⛔ RAR5 format — pyReScene 0.7 cannot reconstruct these", "err")
                elif has_compressed:
                    self._log("    ⚠ Compressed RARs — needs exact WinRAR version from pack", "warn")
                else:
                    self._log("    Uncompressed RARs — pyReScene handles natively", "dim")
            return is_rar5
        except Exception as e:
            self._log(f"    (SRR read error: {e})", "dim")
            return False

    # Shims for three Python 3.12+ incompatibilities in pyReScene 0.7, applied
    # before importing resample.srs:
    #   time.clock() removed in 3.8; distutils removed in 3.12 (used by
    #   resample.fpcalc); locale.format() removed in 3.12 (used by rescene's
    #   sep() in the rebuild results display — crashing there strands the
    #   rebuilt sample as a .tmp file before the CRC check and rename).
    _SRS_WRAPPER = (
        "import time, types, shutil, sys, locale; "
        "time.clock = time.perf_counter; "
        "locale.format = locale.format_string; "
        "_ds = types.ModuleType('distutils'); "
        "_dss = types.ModuleType('distutils.spawn'); "
        "_dss.find_executable = shutil.which; "
        "sys.modules.setdefault('distutils', _ds); "
        "sys.modules.setdefault('distutils.spawn', _dss); "
        "from resample.srs import main; main(sys.argv[1:])"
    )

    def _srs_info(self, srs_path: str) -> dict:
        """Run `srs -l` on an SRS file: log its metadata and return parsed fields
        (notably 'type', e.g. STREAM / MKV / AVI)."""
        info: dict = {}
        try:
            ri = subprocess.run(
                [sys.executable, "-c", self._SRS_WRAPPER, str(srs_path), "-l"],
                capture_output=True, text=True, timeout=60,
            )
            for line in (ri.stdout.strip() + "\n" + ri.stderr.strip()).splitlines():
                line = line.strip()
                if not line:
                    continue
                self._log(f"    srs-info: {line}", "dim")
                if m := re.match(r"^SRS Type\s*:\s*(\S+)", line):
                    info["type"] = m.group(1).upper()
                elif m := re.match(r"^Sample Name\s*:\s*(.+)$", line):
                    info["name"] = m.group(1).strip()
                elif m := re.match(r"^Sample Size\s*:\s*([\d,]+)", line):
                    info["size"] = int(m.group(1).replace(",", ""))
                elif m := re.match(r"^Sample CRC\s*:\s*([0-9A-Fa-f]{1,8})", line):
                    info["crc"] = int(m.group(1), 16)
        except Exception:
            pass
        return info

    @staticmethod
    def _crc32_file(path: str) -> int:
        import zlib
        crc = 0
        with open(path, "rb") as f:
            while chunk := f.read(1 << 20):
                crc = zlib.crc32(chunk, crc)
        return crc & 0xFFFFFFFF

    _VERIFY_SKIP_DIRS = {"_stored", "_subs_tmp", "_iso_m2ts_tmp", "Sample"}

    @staticmethod
    def _force_rmtree(path: Path, attempts: int = 5) -> list:
        """Delete a tree, surviving the two Windows failures that make
        delete-source look flaky rather than broken.

        `shutil.rmtree` walks bottom-up and raises on the FIRST file it cannot
        remove — after it has already deleted everything it reached. So a single
        stubborn file does not skip the delete, it leaves a half-emptied source
        folder and one error line, which is exactly the "it deleted most of
        them but left a few" shape. The two causes here:

          * a read-only attribute (scene extras often carry one), which rmtree
            reports as PermissionError, and
          * a handle the just-finished rar.exe, the SFV verify pass, or an AV
            scanner has not released yet — transient, and gone within a second.

        Clearing the attribute handles the first; retrying handles the second.
        Returns the paths that STILL exist afterwards — empty means success, so
        the caller can report what survived instead of guessing."""
        import stat as _stat

        def _clear(fn, p, _exc):
            try:
                os.chmod(p, _stat.S_IWRITE)
                fn(p)
            except Exception:
                pass

        for i in range(attempts):
            try:
                shutil.rmtree(str(path), onerror=_clear)
            except Exception:
                pass
            if not path.exists():
                return []
            time.sleep(0.4 * (i + 1))
        left = [str(p) for p in path.rglob("*") if p.is_file()] \
            if path.exists() else []
        return left or ([str(path)] if path.exists() else [])

    def _verify_rebuilt_sfv(self, out_root: Path) -> dict:
        """CRC32-check the produced RAR volumes against every SFV under out_root.

        rescene returning "ok" does NOT guarantee each volume is byte-exact — a
        near-miss on a trailing compressed stream can still leave a bad last
        volume. This is the only honest success test, and the guard that must
        pass before delete-source is allowed to run.

        Returns {sfv_count, checked, bad:[(name,exp,got)], missing:[name]}."""
        sfvs, seen = [], set()
        for pat in ("*.sfv", "*.SFV"):
            for p in out_root.rglob(pat):
                rp = p.resolve()
                if rp not in seen:
                    seen.add(rp); sfvs.append(p)

        # name -> produced file (skip our own working dirs and metadata)
        index: dict = {}
        for f in out_root.rglob("*"):
            if not f.is_file():
                continue
            if self._VERIFY_SKIP_DIRS & set(f.relative_to(out_root).parts[:-1]):
                continue
            index.setdefault(f.name.lower(), f)

        checked, bad, missing, entry_seen = 0, [], [], set()
        for s in sfvs:
            # An SFV lists volumes relative to its own folder; match by basename.
            for name, exp in self._parse_sfv(str(s)):
                key = name.lower()
                if key in entry_seen:
                    continue
                entry_seen.add(key)
                exp = exp.upper().zfill(8)
                # don't try to CRC-verify the SFV/NFO themselves
                if key.endswith((".sfv", ".nfo")):
                    continue
                f = index.get(key)
                if f is None:
                    missing.append(name); continue
                got = "%08X" % self._crc32_file(str(f))
                checked += 1
                if got != exp:
                    bad.append((name, exp, got))
        return {"sfv_count": len(sfvs), "checked": checked,
                "bad": bad, "missing": missing}

    def _srs_create_sample(self, srs_path: str, video_path: str, out_dir: str) -> dict:
        """Create scene sample from SRS file + full video.
        -y skips the overwrite prompt (which would block waiting for stdin)."""
        if not (_find_script("srs") or self._srs_script):
            return {"ok": False, "error": "srs script not found (installed with pyReScene)"}
        Path(out_dir).mkdir(parents=True, exist_ok=True)
        cmd = [sys.executable, "-c", self._SRS_WRAPPER,
               str(srs_path), str(video_path), "-o", str(out_dir), "-y"]
        self._log(f"    srs: {Path(srs_path).name} + {Path(video_path).name} → {Path(out_dir).name}", "dim")
        try:
            t0 = time.time()
            # Generous timeout: locating the sample in a full Blu-ray stream is a
            # linear scan of the whole file (can be 40+ GB).
            r = subprocess.run(cmd, capture_output=True, text=True, timeout=7200)
            elapsed = time.time() - t0
            output = "\n".join(filter(None, [r.stdout.strip(), r.stderr.strip()]))
            # Strip the terminal progress spinner (written as backspace+char pairs)
            output = re.sub(r"\x08.", "", output).strip()
            for line in output.splitlines():
                if line.strip():
                    self._log(f"    srs: {line}", "dim")
            self._log(f"    srs finished in {elapsed:.0f}s (exit {r.returncode})", "dim")

            out_p = Path(out_dir)
            all_files = [f for f in out_p.iterdir() if f.is_file()] if out_p.exists() else []
            created = [f for f in all_files if f.suffix.lower() != ".tmp"]
            tmp_files = [f for f in all_files if f.suffix.lower() == ".tmp"]

            # pyReScene bug: replace_result() can fail its final os.rename silently,
            # leaving '<sample name>-<random>.tmp' behind while srs still exits 0
            # and reports success. Recover the intended name and rename it ourselves.
            if not created and tmp_files and r.returncode == 0:
                tmp = max(tmp_files, key=lambda f: f.stat().st_mtime)
                m = re.match(r"^(.+)-[A-Za-z0-9_]+\.tmp$", tmp.name)
                if m:
                    target = out_p / m.group(1)
                    if not target.exists():
                        tmp.rename(target)
                        self._log(f"    (renamed lingering temp file → {target.name})", "dim")
                        created = [target]
            if r.returncode == 0:
                # Success — clear any stale temp files from older runs
                for tf in tmp_files:
                    if tf.exists() and tf not in created:
                        tf.unlink(missing_ok=True)
            elif tmp_files:
                # Failure — keep temp files for inspection, just report them
                names = ", ".join(f.name for f in tmp_files if f.exists())
                if names:
                    self._log(f"    (temp file(s) left in Sample dir: {names})", "dim")

            if created and r.returncode == 0:
                return {"ok": True, "output": output, "files": [f.name for f in created]}
            # Full output is already logged above — the error field should be
            # just the reason (last meaningful line), not the whole blob.
            err_line = next(
                (l.strip() for l in reversed(output.splitlines()) if l.strip()), "")
            if r.returncode != 0:
                return {"ok": False, "output": output,
                        "error": err_line or f"srs exited {r.returncode}"}
            return {"ok": False, "output": output,
                    "error": f"srs ran (exit 0) but no file created: {err_line}"}
        except subprocess.TimeoutExpired:
            return {"ok": False, "error": "Sample creation timed out (2-hour limit)"}
        except Exception as e:
            return {"ok": False, "error": str(e)}

    def _find_7z(self) -> str | None:
        for name in ("7z.exe", "7za.exe"):
            p = self._app_dir / name
            if p.exists():
                return str(p)
        return shutil.which("7z") or shutil.which("7za")

    def _list_iso_m2ts(self, iso_path: str) -> dict:
        """List all M2TS streams inside a Blu-ray ISO using 7z.
        Returns {"ok", "seven_zip", "entries": [(path_in_iso, size), …]}."""
        seven_zip = self._find_7z()
        if not seven_zip:
            return {"ok": False, "error": "7z.exe not found in apps/ — needed to read ISO"}
        try:
            r = subprocess.run(
                [seven_zip, "l", "-slt", "-r", iso_path, "*.m2ts"],
                capture_output=True, text=True, timeout=60,
            )
        except subprocess.TimeoutExpired:
            return {"ok": False, "error": "ISO listing timed out"}
        except Exception as e:
            return {"ok": False, "error": f"7z list: {e}"}

        # Parse technical listing blocks: "Path = ...\nSize = ..."
        entries: list[tuple[str, int]] = []
        cur_path = cur_size = None
        for line in r.stdout.splitlines():
            if line.startswith("Path = "):
                cur_path = line[7:].strip()
                cur_size = None
            elif line.startswith("Size = ") and cur_path:
                try:
                    cur_size = int(line[7:].strip())
                except ValueError:
                    pass
            elif not line.strip() and cur_path and cur_size is not None:
                if cur_path.lower().endswith(".m2ts") and cur_size > 0:
                    entries.append((cur_path, cur_size))
                cur_path = cur_size = None
        if cur_path and cur_size and cur_path.lower().endswith(".m2ts"):
            entries.append((cur_path, cur_size))

        if not entries:
            return {"ok": False, "error": "No M2TS streams found in ISO"}
        return {"ok": True, "seven_zip": seven_zip, "entries": entries}

    def _carve_nonscene_sample(self, iso_path: str, out_dir: str,
                               base_name: str, target_bytes: int | None) -> dict:
        """Carve a NON-SCENE preview clip from the start of the ISO's main M2TS
        stream. M2TS is valid from byte 0 (opens with PAT/PMT), so streaming the
        first N bytes via `7z -so` yields a playable clip without extracting the
        full stream. The result is clearly named NONSCENE and is NOT the scene
        sample — it will never CRC-match the SRS."""
        lst = self._list_iso_m2ts(iso_path)
        if not lst["ok"]:
            return {"ok": False, "error": lst["error"]}
        best_path, best_size = max(lst["entries"], key=lambda x: x[1])

        # Size the clip like the real scene sample when known; sane bounds.
        target = target_bytes or (100 << 20)
        target = max(16 << 20, min(target, 300 << 20, best_size))
        target -= target % 192  # end on an M2TS packet boundary

        Path(out_dir).mkdir(parents=True, exist_ok=True)
        out_file = Path(out_dir) / f"NONSCENE-{base_name}-preview.m2ts"
        try:
            proc = subprocess.Popen(
                [lst["seven_zip"], "e", "-so", iso_path, best_path],
                stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
            )
            remaining = target
            with open(out_file, "wb") as f:
                while remaining > 0:
                    chunk = proc.stdout.read(min(1 << 20, remaining))
                    if not chunk:
                        break
                    f.write(chunk)
                    remaining -= len(chunk)
            proc.kill()
            proc.wait(timeout=30)
        except Exception as e:
            return {"ok": False, "error": f"carve failed: {e}"}
        if not out_file.exists() or out_file.stat().st_size == 0:
            return {"ok": False, "error": "no data carved from ISO"}
        return {"ok": True, "path": str(out_file), "name": out_file.name,
                "size": out_file.stat().st_size}

    def _extract_m2ts_from_iso(self, iso_path: str, tmp_dir: str) -> dict:
        """
        List all M2TS streams inside a Blu-ray ISO using 7z, pick the largest
        (main feature), extract it to tmp_dir, and return its path.
        Caller is responsible for deleting tmp_dir afterwards.
        """
        lst = self._list_iso_m2ts(iso_path)
        if not lst["ok"]:
            return {"ok": False, "error": lst["error"]}
        seven_zip = lst["seven_zip"]
        entries = lst["entries"]

        # Pick largest (main feature, not trailers/extras)
        best_path, best_size = max(entries, key=lambda x: x[1])
        best_name = Path(best_path).name
        gb = best_size / 1_073_741_824
        self._log(
            f"  ISO contains {len(entries)} M2TS stream(s); extracting {best_name} ({gb:.1f} GB) — "
            "this may take several minutes…",
            "dim",
        )

        Path(tmp_dir).mkdir(parents=True, exist_ok=True)
        # 2 — extract just that one file (7z 'e' drops the directory structure)
        try:
            r2 = subprocess.run(
                [seven_zip, "e", iso_path, best_path, f"-o{tmp_dir}", "-y"],
                capture_output=True, text=True, timeout=14400,  # 4 h
            )
        except subprocess.TimeoutExpired:
            return {"ok": False, "error": "ISO extraction timed out (4-hour limit)"}
        except Exception as e:
            return {"ok": False, "error": f"7z extract: {e}"}

        extracted = Path(tmp_dir) / best_name
        if not extracted.exists():
            err = r2.stderr.strip() or r2.stdout.strip() or "file not found after extraction"
            return {"ok": False, "error": f"Extraction failed: {err}"}

        self._log(f"  Extracted {best_name} ({extracted.stat().st_size:,} B)", "dim")
        return {"ok": True, "path": str(extracted), "name": best_name}

    # ── Results database ──────────────────────────────────────────────────────
    # Every processed release is recorded to srrdb_results.json automatically:
    # year, platform, group, outcome, detected RAR version. View aggregated
    # stats in the log or export CSV for spreadsheet work.

    _PLATFORMS = ["NSW", "3DS", "NDS", "GBA", "GBC", "N64", "NGC", "GCN",
                  "WIIU", "WII", "XBOX360", "XBOXONE", "XBOX", "PS2", "PS3",
                  "PS4", "PS5", "PSP", "PSV", "VITA", "PSX", "DC", "DREAMCAST"]
    _VIDEO_TOKENS = {"BLURAY", "DVDRIP", "DVDR", "WEB", "HDTV", "X264", "X265",
                     "XVID", "H264", "H265", "BDRIP", "WEBRIP"}

    @property
    def _results_path(self) -> Path:
        return Path(__file__).parent / "srrdb_results.json"

    def _load_results(self) -> list:
        try:
            return json.loads(self._results_path.read_text(encoding="utf-8"))
        except Exception:
            return []

    def _parse_meta(self, release: str, queue_path: str) -> tuple:
        """(group, platform, year) derived from a release name + source folder."""
        tokens = set(re.split(r"[._\-]+", (release or "").upper()))
        group = ""
        m = re.match(r"^.+-([A-Za-z0-9_]{2,15})$", release or "")
        if m:
            group = m.group(1)
        platform = next((p for p in self._PLATFORMS if p in tokens), "")
        if not platform and tokens & self._VIDEO_TOKENS:
            platform = "video"
        year = ""
        if queue_path:
            ym = re.match(r"^(\d{4})[-._]\d{2}[-._]\d{2}", Path(queue_path).name)
            if ym:
                year = ym.group(1)
        return group, platform, year

    def _preferred_rar_versions(self, release: str, queue_path: str) -> list:
        """Known-good RAR versions from past successes of the SAME group —
        same year first, then the group's other years. rescene will try these
        before its date-ordered sweep, so a group's second release matches in
        one attempt instead of scanning the whole pack."""
        try:
            group, _, year = self._parse_meta(release, queue_path)
            if not group:
                return []
            prefs: list = []
            results = self._load_results()
            for want_year in (True, False):
                for r in results:
                    if (r.get("ok") and r.get("rar_version")
                            and r.get("group") == group
                            and (not want_year or r.get("year") == year)
                            and r["rar_version"] not in prefs):
                        prefs.append(r["rar_version"])
            return prefs[:4]
        except Exception:
            return []

    def _pack_signature(self) -> str:
        """Stable fingerprint of the current WinRAR pack: version count + a short
        hash of the sorted exe basenames. Changes when versions are added/removed,
        so a cached version-wall is re-attempted after the pack grows."""
        try:
            rar_dir = self._find_rar_dir()
            names = sorted(f.lower() for f in os.listdir(rar_dir)
                           if _EXE_RE.search(f))
            h = hashlib.sha1("\n".join(names).encode("utf-8")).hexdigest()[:12]
            return f"{len(names)}:{h}"
        except Exception:
            return ""

    def _wall_cache_hit(self, release: str, date_cap_on: bool = True) -> dict | None:
        """Return the prior result record if `release` is a cached version-wall
        that is still valid (same pack signature + wall-cache generation), else
        None. A hit means: don't re-grind — nothing changed that could make it
        rebuildable. A CAPPED wall (only the in-range subset was tried) is honored
        ONLY while the date-cap is still on — turn the cap off and it re-runs the
        whole pack, so the cap's subset is never a permanent exclusion."""
        if not release:
            return None
        try:
            cur_sig = self._pack_signature()
            if not cur_sig:
                return None
            for r in self._load_results():
                if (r.get("release") == release
                        and r.get("wall") == "version"
                        and r.get("wall_gen") == _WALL_CACHE_GEN
                        and r.get("pack_sig") == cur_sig):
                    if r.get("wall_capped") and not date_cap_on:
                        return None   # cap off → try the whole pack
                    return r
        except Exception:
            pass
        return None

    def _prior_versions_tried(self, release: str) -> set:
        """Versions a previous, INCOMPLETE run of `release` already tested (same
        pack). Empty when there's no such record — a clean exhaustion is handled
        by the wall cache instead, and a different pack invalidates the list."""
        if not release:
            return set()
        try:
            cur = self._pack_signature()
            if not cur:
                return set()
            for r in self._load_results():
                if (r.get("release") == release and not r.get("ok")
                        and r.get("tried_pack_sig") == cur):
                    return set(r.get("versions_tried") or [])
        except Exception:
            pass
        return set()

    def _sweep_cache_hit(self, release: str) -> dict | None:
        """The stored recipe-sweep verdict for `release`, or None to sweep now.

        {"found": {exe, version, mt}|None, "ts": …}.

        A HIT and a MISS are not the same kind of claim, and are not cached on
        the same terms:

        * A hit is a measurement — that exe reproduced this set's streams
          byte-exact. That stays true no matter how the pack changes or how the
          sweep's search order is rewritten, so it is honoured across pack
          signatures and generations. The one thing that can invalidate it is
          the exe no longer being in the pack.
        * A miss is only ever "the search I ran didn't find one". A bigger pack
          or a fixed search can both turn it into a hit, so it counts only when
          recorded against the CURRENT pack signature AND generation."""
        if not release:
            return None
        try:
            cur = self._pack_signature()
            for r in self._load_results():
                if r.get("release") != release:
                    continue
                sw = r.get("sweep") or {}
                if not sw:
                    continue
                found = sw.get("found")
                if found:
                    exe = found.get("exe") or ""
                    if exe and (Path(self._find_rar_dir()) / exe).is_file():
                        return {"found": found, "ts": r.get("ts")}
                    continue
                if cur and sw.get("pack_sig") == cur \
                        and sw.get("gen") == _WALL_CACHE_GEN:
                    return {"found": None, "ts": r.get("ts")}
        except Exception:
            pass
        return None

    def _record_result(self, summary: dict, queue_path: str):
        try:
            release = summary.get("release") or ""
            group, platform, year = self._parse_meta(release, queue_path)
            # Classify a VERSION WALL: reconstruction failed, no version ever
            # locked (_last_good_rar None), AND the hunt tried EVERY version it
            # was going to (not cut off early by the deadline). Cache it (with the
            # pack signature) so a re-run skips the grind. `_versions_tried` and
            # `_all_versions` are compared as DISTINCT version STRINGS — betas
            # collapse to one string (e.g. 5.00 + 5.00b1..b8 → "2013-10-12 5.00"),
            # so 232 exes are only ~58 strings; comparing raw lengths would never
            # match. When a date-cap was applied the search set is the in-range
            # subset (`_capped_count` distinct strings), not the whole pack.
            allv = set(getattr(self, "_all_versions", None) or [])
            tried = getattr(self, "_versions_tried", None) or set()
            cap_on = bool(getattr(self, "_version_date_cap", None)
                          and not getattr(self, "_version_cap_widen", False))
            cap_src = getattr(self, "_date_cap_source", None)
            searched = (getattr(self, "_capped_count", None)
                        if cap_on else None) or len(allv)
            # Only a CLEAN finish counts as exhaustion — a deadline-truncated
            # hunt may have logged every distinct version string while some betas
            # of the last few went untested, so never cache it as a wall.
            hit_deadline = bool(getattr(self, "_recon_hit_deadline", False))
            exhausted = bool(allv and len(tried) >= searched and not hit_deadline)
            no_version = getattr(self, "_last_good_rar", None) is None
            # A capped exhaustion is a real wall only when the date is RELIABLE
            # (from the folder). An SRR-timestamp date is noisy, so a capped
            # SRR-dated miss stays widen-on-re-run instead of a cached wall.
            wall = ""
            wall_capped = False
            if not summary.get("ok") and no_version and exhausted:
                if not cap_on:
                    wall = "version"
                elif cap_src == "folder":
                    wall = "version"; wall_capped = True
            rec = {
                "ts":          time.strftime("%Y-%m-%d %H:%M"),
                "release":     release,
                "group":       group,
                "platform":    platform,
                "year":        year,
                "ok":          bool(summary.get("ok")),
                "rars":        summary.get("rars"),
                "sample":      summary.get("sample"),
                "subs":        summary.get("subs"),
                "rar_version": getattr(self, "_last_good_rar", None),
                # Per-stream (file, version, -mt) combos the hunt locked. On an
                # ok=True record these are the byte-exact packing settings — the
                # accumulating dataset behind the .srr2 idea. Stored as
                # [[file, version, mt], …]; mt may be null for pre-mt versions.
                "combos":      [list(c) for c in
                                (getattr(self, "_recon_streams", None) or [])],
                "note":        (summary.get("note") or "")[:120],
                # Version-wall cache (see _wall_cache_hit): "" for anything that
                # isn't a proven version miss over its searched set.
                "wall":        wall,
                "wall_gen":    _WALL_CACHE_GEN if wall else None,
                "pack_sig":    self._pack_signature() if wall else None,
                # A capped wall only tried the in-range subset, so it's skipped on
                # re-run ONLY while the date-cap is still on (see _wall_cache_hit).
                "wall_capped": wall_capped,
                # Was a release-date cap actually applied this run? A capped
                # SRR-dated FAILURE triggers widen-on-re-run (see _process_one).
                "date_capped": cap_on,
            }
            # ── what this run actually TRIED, so a re-run doesn't repeat it ──
            # The recipe sweep's verdict (see _sweep_cache_hit). Recorded for a
            # hit, and for a CLEAN exhaustion; a sweep cut short by the budget or
            # by Stop is deliberately NOT cached, since it proved nothing.
            recipe = getattr(self, "_recipe_found", "skip")
            if recipe != "skip" and (recipe
                                     or getattr(self, "_sweep_exhausted", False)):
                rec["sweep"] = {"pack_sig": self._pack_signature(),
                                "gen": _WALL_CACHE_GEN,
                                "found": recipe or None}
            # Versions the hunt actually tested. On a deadline-truncated failure
            # a re-run reorders these to the BACK (see _process_one), so a second
            # pass explores NEW builds instead of grinding the same ones again.
            if tried and not summary.get("ok"):
                rec["versions_tried"] = sorted(tried)
                rec["tried_pack_sig"] = self._pack_signature()
            results = self._load_results()
            # Re-runs replace the previous record for the same release
            results = [r for r in results if r.get("release") != release]
            results.append(rec)
            self._results_path.write_text(
                json.dumps(results, indent=1), encoding="utf-8")
        except Exception:
            pass

    def results_stats(self) -> dict:
        """Log aggregated testing stats: per group and per year/platform."""
        results = self._load_results()
        if not results:
            self._log("Results DB is empty — process some releases first.", "dim")
            return {"ok": True, "count": 0}
        ok_n = sum(1 for r in results if r["ok"])
        self._log(f"══ Results DB — {len(results)} release(s), "
                  f"{ok_n} ok ({ok_n / len(results):.0%}) ══", "info")

        def _agg(key):
            groups: dict = {}
            for r in results:
                k = r.get(key) or "?"
                g = groups.setdefault(k, {"n": 0, "ok": 0, "vers": set()})
                g["n"] += 1
                g["ok"] += 1 if r["ok"] else 0
                if r.get("rar_version"):
                    g["vers"].add(r["rar_version"])
            return groups

        self._log("Per group:", "info")
        for k, g in sorted(_agg("group").items(),
                           key=lambda kv: -kv[1]["n"]):
            vers = ", ".join(sorted(g["vers"])) or "—"
            cls = "ok" if g["ok"] == g["n"] else ("warn" if g["ok"] else "err")
            self._log(f"  {k:<18} {g['ok']}/{g['n']:<4} rar: {vers}", cls)

        self._log("Per year/platform:", "info")
        combo: dict = {}
        for r in results:
            k = f"{r.get('year') or '????'} {r.get('platform') or '?'}"
            g = combo.setdefault(k, {"n": 0, "ok": 0})
            g["n"] += 1
            g["ok"] += 1 if r["ok"] else 0
        for k, g in sorted(combo.items()):
            self._log(f"  {k:<16} {g['ok']}/{g['n']}",
                      "ok" if g["ok"] == g["n"] else "warn")

        # Known-good packing combos harvested from clean rebuilds — the raw
        # (version, -mt) dataset behind the .srr2 idea. Only ok=True records
        # carry byte-exact settings, so restrict to those and de-dup.
        seen: set = set()
        combo_rows: list = []
        for r in results:
            if not r.get("ok"):
                continue
            for c in (r.get("combos") or []):
                fname, ver, mt = (list(c) + [None, None, None])[:3]
                mt_txt = f"-mt{mt}" if mt is not None else "-mt?"
                key = (r.get("group") or "?", ver or "?", mt_txt,
                       (str(fname).rsplit(".", 1)[-1] or "").lower())
                if key in seen:
                    continue
                seen.add(key)
                combo_rows.append(key)
        if combo_rows:
            self._log(f"Known-good packing combos ({len(combo_rows)}):", "info")
            for grp, ver, mt_txt, ext in sorted(combo_rows):
                self._log(f"  {grp:<18} {ver:<16} {mt_txt:<6} .{ext}", "ok")

        fails = [r for r in results if not r["ok"]]
        if fails:
            self._log(f"Failures ({len(fails)}):", "warn")
            for r in fails:
                self._log(f"  ✗ {r['release']} — {r.get('note') or '?'}", "err")
        return {"ok": True, "count": len(results)}

    _RESULT_COLS = ["ts", "release", "group", "platform", "year", "ok", "rars",
                    "sample", "subs", "rar_version", "combos", "note"]

    @staticmethod
    def _fmt_combos(combos) -> str:
        """Flatten [[file, version, mt], …] to 'file=version -mtN; …' for CSV."""
        out = []
        for c in (combos or []):
            fname, ver, mt = (list(c) + [None, None, None])[:3]
            mt_txt = f"-mt{mt}" if mt is not None else "-mt?"
            out.append(f"{fname}={ver or '?'} {mt_txt}")
        return "; ".join(out)

    def export_results_csv(self) -> dict:
        results = self._load_results()
        if not results:
            return {"ok": False, "error": "Results DB is empty"}
        import csv
        out = Path(__file__).parent / "srrdb_results.csv"
        with open(out, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=self._RESULT_COLS, extrasaction="ignore")
            w.writeheader()
            for r in results:
                row = dict(r)
                row["combos"] = self._fmt_combos(r.get("combos"))
                w.writerow(row)
        self._log(f"Exported {len(results)} record(s) → {out}", "ok")
        return {"ok": True, "path": str(out), "count": len(results)}

    def export_results_xlsx(self) -> dict:
        """Styled spreadsheet: frozen header, autofilter, colour-coded rows."""
        results = self._load_results()
        if not results:
            return {"ok": False, "error": "Results DB is empty"}
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
        except ImportError:
            return {"ok": False, "error": "openpyxl not installed — pip install openpyxl"}

        out = Path(__file__).parent / "srrdb_results.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "srrdb results"

        hdr_fill = PatternFill("solid", fgColor="2A3045")
        hdr_font = Font(bold=True, color="FFFFFF")
        ok_fill = PatternFill("solid", fgColor="E2F0DA")   # soft green
        bad_fill = PatternFill("solid", fgColor="F8D7DA")  # soft red

        ws.append([c.upper() for c in self._RESULT_COLS])
        for cell in ws[1]:
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center")

        for r in sorted(results, key=lambda x: (x.get("year") or "",
                                                x.get("group") or "",
                                                x.get("release") or "")):
            row = [r.get(c, "") for c in self._RESULT_COLS]
            row[self._RESULT_COLS.index("ok")] = "OK" if r.get("ok") else "FAILED"
            row[self._RESULT_COLS.index("combos")] = self._fmt_combos(r.get("combos"))
            ws.append(row)
            fill = ok_fill if r.get("ok") else bad_fill
            for cell in ws[ws.max_row]:
                cell.fill = fill

        widths = {"ts": 16, "release": 55, "group": 14, "platform": 10,
                  "year": 7, "ok": 9, "rars": 7, "sample": 24, "subs": 14,
                  "rar_version": 18, "combos": 48, "note": 60}
        for i, c in enumerate(self._RESULT_COLS, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 14)

        ws.freeze_panes = "A2"                       # header stays put
        ws.auto_filter.ref = ws.dimensions           # sortable/filterable columns
        try:
            wb.save(out)
        except PermissionError:
            return {"ok": False, "error":
                    "srrdb_results.xlsx is open in Excel — close it and retry"}
        self._log(f"Exported {len(results)} record(s) → {out} "
                  "(frozen header, colour-coded, filterable)", "ok")
        return {"ok": True, "path": str(out), "count": len(results)}

    # ── Single-release processing thread ─────────────────────────────────────

    def start_process(self, config: dict) -> bool:
        if self._running:
            self._log("Already running.", "warn")
            return False
        self._stop.clear()
        self._skip.clear()
        self._fresh_all = False   # "always fresh" latches per RUN, not forever
        threading.Thread(target=self._process_one, args=(config,), daemon=True).start()
        return True

    # ── "run it fresh" prompt ────────────────────────────────────────────────
    # Three things a re-run reuses from the results DB: a cached version-wall
    # (skip the release outright), a cached recipe-sweep verdict, and the list
    # of builds a timed-out run already tested (an ordering hint). Each saves
    # real time — and each is only as good as the code that recorded it, so
    # after a fix the stored verdict can be stale in a way nothing detects.
    # Rather than hand-editing srrdb_results.json, pause and offer to ignore it.
    #
    # Deliberately ONE prompt per release (the wall check and the tried-builds
    # hint both fire at job start) and it always AUTO-CONTINUES, so an
    # unattended overnight batch behaves exactly as it does today.
    _FRESH_PROMPT_SECS = 30

    def fresh_prompt_reply(self, choice: str):
        """GUI answer to the pause: 'fresh', 'history', or 'always'."""
        self._fresh_choice = choice
        self._fresh_reply.set()

    def _use_db_history(self, kind: str, detail: str = "") -> bool:
        """True = reuse what the DB stored, False = ignore it and run fresh."""
        if self._fresh_all:
            return False
        if self._fresh_release is not None:      # already decided this release
            return self._fresh_release
        if not getattr(self, "_fresh_prompt_on", False):
            return True
        secs = self._FRESH_PROMPT_SECS
        self._fresh_choice = None
        self._fresh_reply.clear()
        self._emit("fresh_prompt", {"kind": kind, "detail": detail,
                                    "release": self._release_name or "",
                                    "secs": secs})
        self._log(f"  ⏸ Reusing stored history ({detail}). Pausing {secs}s — "
                  "choose 'Run fresh' in the banner to ignore it.", "warn")
        # Poll rather than one long wait so Stop/Skip still land immediately.
        end = time.time() + secs
        while time.time() < end:
            if self._stop.is_set() or self._skip.is_set():
                break
            if self._fresh_reply.wait(0.25):
                break
        self._emit("fresh_prompt_done", {})
        choice = self._fresh_choice
        if choice == "always":
            self._fresh_all = True
            self._log("  ↻ Ignoring stored DB history for the REST of this run.",
                      "warn")
        elif choice == "fresh":
            self._log("  ↻ Ignoring stored DB history — running this release "
                      "from scratch.", "warn")
        else:
            self._log("  ▶ Continuing with the stored history.", "dim")
        self._fresh_release = choice not in ("fresh", "always")
        return self._fresh_release

    def stop_process(self, scope: str = "all"):
        """Interrupt the running work.

        scope="current": abort the release being rebuilt right now and move on
            to the next queued file — the batch keeps going.
        scope="all": abort the current release AND halt the whole batch/queue.

        Either way, kill any in-flight rar.exe immediately so a long compress
        stops within a second instead of at the next inter-file checkpoint."""
        for p in list(self._live_procs):
            try:
                p.kill()
            except Exception:
                pass
        if scope == "current":
            self._skip.set()
            self._log("  ⏭ Skipping current release — moving to the next file…", "warn")
        else:
            self._stop.set()
            self._log("  ⏹ Stopping — aborting current release and halting the batch…", "warn")

    def _process_one(self, config: dict, batch_mode: bool = False):
        if not batch_mode:
            self._running = True
            self._emit("status", {"state": "running"})

        release     = (config.get("release_name") or "").strip()
        candidates  = config.get("candidates", [])   # list of {release, hasNFO, hasSRS}
        max_test    = int(config.get("max_test", 5))
        content_dir      = (config.get("content_dir") or "").strip()
        dest_dir         = (config.get("dest_dir") or "").strip()
        do_sample        = config.get("do_sample", False)
        extract_iso_m2ts = config.get("extract_iso_m2ts", False)
        non_scene_sample = config.get("non_scene_sample", False)
        delete_source    = config.get("delete_source", False)
        # Queue rows in the GUI are keyed by the original folder path — keep it
        # for job events even after auto-match renames or nesting descent.
        queue_path = content_dir

        # Per-job outcome for the queue status and the end-of-batch summary
        summary = {
            "release": release or (Path(queue_path).name if queue_path else "?"),
            "ok": True, "rars": None, "sample": None, "subs": None, "note": "",
        }
        self._last_good_rar = None  # set by the rescene event stream
        self._last_good_exe = None  # exact build (exe) rescene invoked
        self._all_versions = []     # full pack list, captured during the hunt
        self._versions_tried = set()  # versions rescene actually tested this job
        self._recon_hit_deadline = False  # cut off by the deadline (not exhausted)
        self._mt_rank_cache = None   # recompute -mt win-frequency per job
        # "skip" = the recipe sweep hasn't run for this release yet; a dict/None
        # after it has (so the last-resort rescue never repeats the fast-path's
        # sweep, and _record_result can persist the measured recipe).
        self._recipe_found = "skip"
        self._sweep_exhausted = False
        self._sweep_skipped = False
        self._release_name = release
        # Builds a PREVIOUS run already tested before it was cut off. Not a
        # skip-list (a truncated run can leave a version half-tested) — just an
        # ordering hint, so the retry reaches untested builds first instead of
        # re-grinding the same head of the list into the same timeout.
        # One "run it fresh" decision per release (see _use_db_history).
        self._fresh_release = None
        self._fresh_prompt_on = bool(config.get("fresh_prompt", False))
        if config.get("ignore_db_history"):
            self._fresh_all = True
        self._retry_tried_versions = self._prior_versions_tried(release)
        if self._retry_tried_versions:
            n = len(self._retry_tried_versions)
            if self._use_db_history("tried", f"{n} build(s) already tested"):
                self._log(
                    f"  Previous run timed out after testing {n} build(s) — "
                    "trying the untested ones first this time.", "info")
            else:
                self._retry_tried_versions = set()

        self._emit("job_start", {"content_dir": queue_path, "release": release})

        # Version-wall cache: skip a release proven last time to match NO pack
        # version, as long as the pack hasn't grown since. Saves the ~30-min
        # whole-pack re-grind on every re-run. Defeatable per-batch via config.
        if release and config.get("skip_known_walls", True):
            hit = self._wall_cache_hit(release, config.get("date_cap", True))
            if hit and not self._use_db_history(
                    "wall", f"version-wall recorded {hit.get('ts', '?')}"):
                hit = None
            if hit:
                extra = ("all builds ≤ the release date"
                         if hit.get("wall_capped") else "every pack version")
                self._log(
                    f"  ⏭ Known version-wall (no match across {extra} on "
                    f"{hit.get('ts','?')}, pack unchanged since) — skipped. "
                    "Add WinRAR versions"
                    + (" or turn off the date-cap" if hit.get("wall_capped")
                       else "") + " to retry.", "warn")
                summary["ok"] = False
                summary["wall_skipped"] = True
                summary["note"] = "known version-wall — skipped (pack unchanged)"
                self._emit("job_done", {"content_dir": queue_path,
                                        "release": release, "ok": False,
                                        "note": summary["note"]})
                if not batch_mode:
                    self._running = False
                    self._emit("status", {"state": "done"})
                return summary

        # Release-date version cap: try builds near the release date first and
        # drop far-future ones (a group can't use a WinRAR newer than the
        # release). WIDEN to the whole pack when a prior capped run for THIS
        # release already failed — so the fast pass runs first, the thorough
        # pass runs only if needed. Cap disabled entirely via config.
        self._version_date_cap = None
        self._version_cap_widen = False
        self._date_cap_source = None   # "folder" | "srr" — reliability of the date
        self._capped_count = None      # distinct version strings in the capped set
        # NB: derive the date from the dated FOLDER (queue_path) even when the
        # release name isn't confirmed yet (auto-match jobs enter here with an
        # empty `release`) — the cap needs the folder, not the name. A missing
        # folder date is filled later from the SRR's own RAR timestamps (see
        # _srr_reconstruct). Cap disabled entirely via config.
        self._date_cap_enabled = bool(config.get("date_cap", True))
        if self._date_cap_enabled:
            self._version_date_cap = self._release_date(release, queue_path)
            if self._version_date_cap:
                self._date_cap_source = "folder"
            if release:   # widen-on-re-run needs a known release name to look up
                try:
                    prev = next((r for r in self._load_results()
                                 if r.get("release") == release), None)
                    if prev and prev.get("date_capped") and not prev.get("ok"):
                        self._version_cap_widen = True
                        self._log("  Version date-cap: prior capped run failed — "
                                  "widening to the whole pack this time.", "dim")
                except Exception:
                    pass
            if self._version_date_cap and not self._version_cap_widen:
                self._log(
                    "  Version date-cap: trying builds up to "
                    f"~{_VERSION_CAP_MARGIN_DAYS // 365}yr after "
                    f"{self._version_date_cap.isoformat()} first (release-dated), "
                    "nearest first.", "dim")

        # Resolve double-nesting: if the selected folder has no files (only one subdir),
        # descend into it. Handles the case where batch source → release folder → content.
        if content_dir and Path(content_dir).is_dir():
            p = Path(content_dir)
            children = list(p.iterdir())
            if children and all(c.is_dir() for c in children) and len(children) == 1:
                content_dir = str(children[0])
                self._log(f"  (descended into subfolder: {children[0].name})", "dim")

        try:
            if not dest_dir:
                self._log("ERROR: No output folder.", "err"); raise ValueError()

            # Metadata-only scene follow-ups (DIRFIX/NFOFIX) carry only a
            # corrected NFO — no packed content — so there is nothing to
            # rebuild and no content CRC to match. Report cleanly instead of
            # letting the match fail with a generic "not in srrdb" error.
            fixtag = _metadata_only_tag(
                release,
                Path(queue_path).name if queue_path else "",
                Path(content_dir).name if content_dir else "")
            if fixtag:
                self._log(
                    f"  ⊘ {fixtag} — metadata-only scene release (corrected "
                    "NFO/dir name); no packed content, nothing to rebuild.",
                    "warn")
                summary["metadata_only"] = fixtag
                summary["note"] = f"{fixtag} — metadata only, nothing to rebuild"
                raise RuntimeError(summary["note"])

            # Empty-folder guard: nothing to match or rebuild. Report it plainly
            # instead of letting the content-CRC lookup fail with the misleading
            # "likely NON-SCENE" message — common with leftover/half-moved folders.
            if content_dir and Path(content_dir).is_dir() and not any(
                    f.is_file() for f in Path(content_dir).rglob("*")):
                self._log("  ⊘ Folder is empty (no content files) — skipped.",
                          "warn")
                summary["empty"] = True
                summary["note"] = "folder empty — nothing to rebuild"
                raise RuntimeError(summary["note"])

            # If no confirmed release name, score candidates against content folder
            if not release:
                if not content_dir or not Path(content_dir).is_dir():
                    self._log("ERROR: Need a content folder to auto-match candidates.", "err")
                    raise ValueError()
                best = {"release": None, "score": 0.0}
                if candidates:
                    self._log(
                        f"  Auto-matching — testing up to {min(len(candidates), max_test)}"
                        f" of {len(candidates)} candidate(s)…", "dim"
                    )
                    best = self.find_best_match(candidates, content_dir, max_test)
                if best["release"] and best["score"] >= 0.5:
                    release = best["release"]
                    self._log(f"  ✓ Matched: {release}  ({best['score']:.0%} confidence)", "ok")
                else:
                    # Candidates are junk (a bad folder name can still match
                    # unrelated releases by name) or absent — an exact content
                    # CRC lookup beats them all.
                    if candidates:
                        self._log(
                            "  Candidates don't match content — trying exact "
                            "content CRC lookup…", "dim",
                        )
                    hres = self.search_by_content_hash(content_dir)
                    if hres.get("ok") and hres.get("results"):
                        release = hres["results"][0]["release"]
                        self._log(f"  ✓ Matched by content CRC: {release}", "ok")
                    elif candidates:
                        top = best.get("release", "?")
                        self._log(
                            f"  No confident match (best: {top} at {best['score']:.0%}) "
                            "and the content CRC is not in srrdb — the file is likely "
                            "NON-SCENE (P2P/custom rip) or modified. Set the release "
                            "name manually if you know it.", "err",
                        )
                        raise RuntimeError("not in srrdb — possibly non-scene")
                    else:
                        self._log(
                            "  No name match and the content CRC is not in srrdb — "
                            "the file is likely NON-SCENE (P2P/custom rip) or modified.",
                            "err",
                        )
                        raise RuntimeError("not in srrdb — possibly non-scene")

            summary["release"] = release
            out_root = Path(dest_dir) / release
            stored_dir = out_root / "_stored"

            self._log(f"▶ {release}", "info")

            # 1 — Download SRR (check both dot and underscore variants in cache)
            alt_release = release.replace(".", "_") if "." in release else release.replace("_", ".")
            srr_file = out_root / f"{release}.srr"
            srr_file_alt = out_root / f"{alt_release}.srr"
            if srr_file.exists():
                self._log(f"  SRR cached ({srr_file.stat().st_size:,} B)", "dim")
            elif srr_file_alt.exists():
                srr_file = srr_file_alt
                self._log(f"  SRR cached ({srr_file.stat().st_size:,} B) [{alt_release}.srr]", "dim")
            else:
                dl = self.download_srr(release, str(out_root))
                if not dl["ok"]:
                    if dl.get("not_found"):
                        self._log(
                            f"  No SRR on srrdb.com for '{release}' — the database has "
                            "no record of this release. It may be a NON-SCENE file "
                            "(P2P/custom rip) or the release name is wrong.", "err",
                        )
                        raise RuntimeError("no SRR on srrdb — possibly non-scene")
                    self._log(f"  ERROR: {dl['error']}", "err"); raise RuntimeError(dl["error"])
                srr_file = Path(dl["srr_path"])  # use actual saved path (may be alt name)
                if dl.get("cached"):
                    self._log(f"  SRR from local cache ({dl['size']:,} B) — "
                              "no srrdb request", "dim")
                else:
                    self._log(f"  SRR downloaded ({dl['size']:,} B)", "ok")

            if self._stop.is_set() or self._skip.is_set(): raise InterruptedError()

            # 2 — Extract stored files (NFO, SFV, SRS)
            self._log("  Extracting stored files (NFO, SFV, SRS)…", "dim")
            ex = self._srr_extract(str(srr_file), str(stored_dir))
            if ex["ok"]:
                files = ex.get("files", [])
                self._log(f"  Stored: {', '.join(files) if files else 'none found'}", "ok" if files else "dim")
                # Move NFO/SFV up to release root for convenience
                for f in stored_dir.rglob("*"):
                    if f.is_file() and f.suffix.lower() in {".nfo", ".sfv"}:
                        dest_f = out_root / f.name
                        if not dest_f.exists():
                            shutil.copy2(str(f), str(dest_f))
            else:
                self._log(f"  Extract WARN: {ex['error']}", "warn")

            if self._stop.is_set() or self._skip.is_set(): raise InterruptedError()

            # 3 — Reconstruct RARs
            self._log("  SRR info:", "dim")
            is_rar5 = self._srr_list(str(srr_file))

            if is_rar5:
                self._log("  Skipping reconstruction — RAR5 not supported by pyReScene 0.7", "err")
                summary["ok"] = False
                summary["note"] = "RAR5 — not supported"
            elif content_dir and Path(content_dir).is_dir():
                # Show what's in the content folder so mismatches are obvious
                cdir = Path(content_dir)
                citems = sorted(cdir.iterdir(), key=lambda f: (f.is_dir(), f.name.lower()))
                self._log(f"  Content folder ({len(citems)} item(s)):", "dim")
                for ci in citems[:12]:
                    tag = "[dir]" if ci.is_dir() else f"{ci.stat().st_size:,} B"
                    self._log(f"    {tag:>14}  {ci.name}", "dim")
                if len(citems) > 12:
                    self._log(f"    … and {len(citems) - 12} more", "dim")

                self._log("  Reconstructing RARs…", "dim")
                # Clear any RAR volumes left in the output by a PREVIOUS run.
                # rescene refuses to overwrite existing archives ("Operation
                # aborted. Archive already exists.") and would just re-verify the
                # stale (near-miss) volumes — which also starves the -mt rescue of
                # freshly-compressed streams to work with, so it can't engage.
                # _clear_produced_volumes only removes SFV-listed volumes; it
                # never touches _stored (the packed sources) or the SRR/NFO/SFV.
                _stale = self._clear_produced_volumes(out_root)
                if _stale:
                    self._log(f"  Cleared {_stale} stale volume(s) from a previous "
                              "run so this rebuild starts fresh.", "dim")
                # Fresh per-release combo log (main set + any nested SRRs both
                # accumulate into this; see _srr_reconstruct).
                self._recon_streams = []
                # No per-stream -mt pin for the normal run (only the multi-file
                # rescue sets this; see _rescue_multifile_crc / _crf_init).
                self._mt_override = {}
                prefs = self._preferred_rar_versions(release, queue_path)
                self._recon_prefs = prefs   # fed to the version-sweep rescue
                SrrdbToolAPI._pref_versions = prefs
                if prefs:
                    self._log(
                        f"  Known-good RAR cache: trying {', '.join(prefs)} "
                        "first (this group's history)", "info",
                    )
                try:
                    rc = self._reconstruct_with_m2_fastpath(
                        str(srr_file), content_dir, out_root)
                finally:
                    SrrdbToolAPI._pref_versions = []
                for line in (rc.get("output") or "").splitlines():
                    line = line.strip()
                    if line:
                        self._log(f"    {line}", "dim")
                if rc["ok"] and rc.get("files"):
                    # Exclude metadata files already placed there before reconstruction
                    produced = [f for f in rc["files"]
                                if Path(f).suffix.lower() not in META_EXTS]
                    self._log(
                        f"  Produced ({len(produced)}): {', '.join(produced[:8])}"
                        + ("…" if len(produced) > 8 else ""),
                        "ok" if produced else "warn",
                    )
                    summary["rars"] = len(produced)
                    if not produced:
                        self._log("  No archive files produced", "warn")
                        summary["ok"] = False
                        summary["note"] = "no RARs produced"
                else:
                    self._log(f"  Reconstruct ERROR: {rc.get('error', 'unknown')}", "err")
                    summary["ok"] = False
                    summary["note"] = (rc.get("error") or "reconstruct error")[:100]
                    if "already exists" in (rc.get("error") or "").lower():
                        # Stale-output collision surfaced by rescene — clear the
                        # produced artifacts (keeps _stored / cached SRR) so a
                        # re-run starts clean, and say so plainly.
                        try:
                            self._clear_produced_volumes(out_root)
                            for d in ("Sample", "Subs", "Subs_NOT_PRODUCED",
                                      "_subs_tmp", "_iso_m2ts_tmp"):
                                shutil.rmtree(str(out_root / d), ignore_errors=True)
                        except Exception:
                            pass
                        self._log("  ⚠ Stale-output collision — cleared the "
                                  "produced files; re-run this release and it "
                                  "should proceed.", "warn")
                        summary["note"] = "stale output collision — cleared, re-run"
                    # If a version+mt was locked before the failure (near-miss /
                    # "Still not fine"), surface the combo so the suspect thread
                    # count is visible.
                    self._log_recon_combos("warn")
                    # A version WAS detected (it reproduced the first stored file)
                    # but the full solid archive still failed — this is the
                    # solid-stream + thread-count determinism wall, not a
                    # missing version.
                    _err = (rc.get("error") or "").lower()
                    if getattr(self, "_last_good_rar", None) and \
                            "exhausted" in _err:
                        self._log(
                            f"  Note: RAR {self._last_good_rar} reproduced the first "
                            "small file, but the full SOLID archive's main file could "
                            "not be matched — the original's thread count / advanced "
                            "compression settings can't be reproduced. Not a "
                            "missing-version issue.", "warn",
                        )
                        summary["note"] = (
                            f"solid archive unrebuildable (matched {self._last_good_rar} "
                            "on first file only)")
                    elif getattr(self, "_last_good_rar", None) and \
                            ("near-miss" in _err or "still not fine" in _err):
                        # rescene locked a version off the test piece but the
                        # full archive was a few bytes off, then gave up. On
                        # -m1 especially, an ancient build can share the piece
                        # CRC while a later build reproduces the whole archive.
                        # Sweep the other pack builds before writing it off.
                        rescued = self._rescue_version_near_miss(
                            str(srr_file), content_dir, out_root)
                        if rescued:
                            summary["ok"] = True
                            summary["verified"] = rescued
                            summary["rars"] = rescued["checked"]
                            summary["note"] = ""
                            self._log_recon_combos("dim")
                    elif not getattr(self, "_last_good_rar", None) and \
                            "no good rar version" in _err:
                        # Detection failed with NOTHING locked. If the set has a
                        # stored extra alongside the sole compressed file, the
                        # content was likely compressed IN-CONTEXT of that extra
                        # (one `rar a` command) and isolated detection can't
                        # match. Drive method2 (compress all files together).
                        rescued = self._rescue_stored_extra_method2(
                            str(srr_file), content_dir, out_root)
                        if rescued:
                            summary["ok"] = True
                            summary["verified"] = rescued
                            summary["rars"] = rescued["checked"]
                            summary["note"] = ""
                            self._log_recon_combos("dim")

                    # Last resort, whatever the failure was (including a deadline
                    # timeout, which no other rescue can follow): measure the
                    # real recipe against the SRR's stream CRCs and rebuild at
                    # it. Skipped when the fast-path already swept this release.
                    if (not summary["ok"]
                            and getattr(self, "_recipe_found", "skip") == "skip"
                            and not (self._stop.is_set() or self._skip.is_set())):
                        rescued = self._rescue_recipe_sweep(
                            str(srr_file), content_dir, out_root)
                        if rescued:
                            rc, v2 = rescued
                            produced = [f for f in (rc.get("files") or [])
                                        if Path(f).suffix.lower()
                                        not in META_EXTS]
                            summary["ok"] = True
                            summary["verified"] = v2
                            summary["rars"] = len(produced) or v2["checked"]
                            summary["note"] = ""
                            self._log_recon_combos("dim")

                # Nested SRRs (e.g. Subs/xxx.subs.srr) describe extra RAR sets
                # such as vobsubs — sometimes two levels deep (per-CD inner RARs
                # inside an outer subs RAR). Extract deeper SRRs first, then
                # rebuild in passes so inner sets become sources for outer ones.
                if stored_dir.exists():
                    found: list[Path] = []
                    scan = sorted(stored_dir.rglob("*.srr"))
                    while scan:
                        nsrr = scan.pop(0)
                        if nsrr in found:
                            continue
                        found.append(nsrr)
                        try:
                            self._srr_extract(str(nsrr), str(nsrr.parent))
                        except Exception:
                            pass
                        for extra in sorted(nsrr.parent.rglob("*.srr")):
                            if extra not in found and extra not in scan:
                                scan.append(extra)

                    pending = found
                    nested_errors: dict = {}
                    for _pass in range(3):
                        if not pending:
                            break
                        remaining = []
                        progressed = False
                        for nsrr in pending:
                            self._log(f"  Nested SRR: {nsrr.name} — reconstructing…", "dim")
                            nrc = self._reconstruct_nested(
                                nsrr, content_dir, out_root, stored_dir)
                            if nrc["ok"]:
                                if nrc.get("already"):
                                    self._log(
                                        "    Target RAR(s) already present — "
                                        "nothing to rebuild", "dim")
                                else:
                                    self._log(
                                        f"    Produced: {', '.join(nrc['produced'][:6])}"
                                        + ("…" if len(nrc["produced"]) > 6 else ""), "ok")
                                progressed = True
                            else:
                                nested_errors[nsrr] = nrc.get("error") or "?"
                                remaining.append(nsrr)
                        pending = remaining
                        if not progressed:
                            break
                    for nsrr in pending:
                        self._log(
                            f"  Nested SRR {nsrr.name} skipped — {nested_errors[nsrr]}",
                            "warn",
                        )

                # Verify produced volumes against the SFV before trusting the
                # rebuild — and, crucially, before delete-source can run.
                # rescene's "ok" does NOT guarantee byte-exact volumes: a
                # near-miss on a trailing compressed stream leaves a bad last
                # volume that must be caught here, reported FAILED, and the
                # source kept.
                if summary.get("ok") and (summary.get("rars") or 0) > 0:
                    vres = self._verify_rebuilt_sfv(out_root)
                    summary["verified"] = vres
                    if vres["bad"]:
                        bad_ex = vres["bad"][0]
                        self._log(
                            f"  ✗ SFV verify FAILED — {len(vres['bad'])} volume(s) "
                            f"wrong (e.g. {bad_ex[0]}: got {bad_ex[2]}, expected "
                            f"{bad_ex[1]}). Near-miss — trying a per-stream -mt "
                            "sweep before giving up…", "warn",
                        )
                        # Show which stream/thread-count combo the hunt locked —
                        # a wrong trailing volume usually points at one stream's
                        # -mt reproducing the right size but the wrong bytes.
                        self._log_recon_combos("warn")
                        # Multi-file CRC rescue: sweep the suspect stream's -mt
                        # (see _rescue_multifile_crc). Only ever runs here, after
                        # the SFV already failed, so nothing that passes today is
                        # affected.
                        rescued = self._rescue_multifile_crc(
                            str(srr_file), content_dir, out_root)
                        if (not rescued
                                and getattr(self, "_recipe_found", "skip") == "skip"
                                and not (self._stop.is_set()
                                         or self._skip.is_set())):
                            # Measure the real recipe against the SRR's stream
                            # CRCs and rebuild at it. This tier was only wired
                            # into the reconstruct-ERROR path, so a set that
                            # PRODUCED volumes and merely failed the SFV could
                            # never reach it — Puzzler_World_2012…PUSSYCAT is
                            # the case: the probe proves 3.60 -mt8 reproduces
                            # both streams, and the version sweep below even
                            # ran rar360.exe at -mt8, yet still missed, because
                            # it compresses each file in ISOLATION while the
                            # release was packed by ONE `rar a jpg nds`. Only
                            # the sweep issues that command. Run it FIRST: it
                            # is seconds rather than minutes and its verdict is
                            # a measurement, so it can only save the version
                            # sweep work it would otherwise do blind.
                            swept = self._rescue_recipe_sweep(
                                str(srr_file), content_dir, out_root)
                            if swept:
                                rescued = swept[1]      # (rc, verify) → verify
                        if not rescued:
                            # The -mt rescue couldn't help — often because the
                            # 2nd file went to method2 so only ONE stream was
                            # recorded (len<2 bail), and nothing tried another
                            # VERSION. The near-miss may just be a wrong build
                            # locked spuriously off a tiny first file (Petz
                            # class: a 579 KB jpg piece-matches 4.11 but the
                            # archive was really 5.11/4.20/3.80). Sweep the other
                            # pack builds (group history first) on the whole
                            # archive before giving up — deadline-bounded, SFV
                            # verify is the arbiter.
                            rescued = self._rescue_version_near_miss(
                                str(srr_file), content_dir, out_root)
                        if rescued:
                            summary["ok"] = True
                            summary["verified"] = rescued
                            summary["rars"] = rescued["checked"]
                            self._log_recon_combos("dim")
                        else:
                            summary["ok"] = False
                            summary["note"] = ("SFV mismatch: "
                                + ", ".join(n for n, _, _ in vres["bad"][:4]))
                            self._log(
                                "  Near-miss, NOT a rebuild — marked FAILED, "
                                "source kept.", "err")
                    elif vres["missing"]:
                        self._log(
                            f"  ⚠ SFV verify: {vres['checked']} volume(s) CRC-OK, but "
                            f"{len(vres['missing'])} SFV-listed file(s) not produced "
                            f"(e.g. {vres['missing'][0]}) — rebuild incomplete.", "warn",
                        )
                    elif vres["checked"]:
                        self._log(f"  ✓ SFV verify: all {vres['checked']} volume(s) "
                                  "CRC-match the SFV.", "ok")
                        # Record the winning combos on a clean rebuild — this is
                        # exactly the (version, -mt) data the .srr2 idea wants.
                        self._log_recon_combos("dim")
                    else:
                        self._log("  ⚠ SFV verify: no SFV found to check against — "
                                  "cannot confirm the rebuild.", "warn")
            else:
                self._log("  No content folder — NFO/SFV extracted only", "dim")
                summary["note"] = "no content folder — NFO/SFV only"

            if self._stop.is_set() or self._skip.is_set(): raise InterruptedError()

            # 4 — Sample creation
            if do_sample and content_dir and Path(content_dir).is_dir():
                srs_files = (list(stored_dir.rglob("*.srs")) + list(stored_dir.rglob("*.SRS"))
                             if stored_dir.exists() else [])
                if srs_files:
                    _DISC_EXTS = {".iso", ".img", ".bin", ".nrg"}
                    _STREAM_EXTS = {".avi", ".mkv", ".mp4", ".m4v", ".mov",
                                    ".wmv", ".m2ts", ".ts", ".vob"}
                    srs_meta = self._srs_info(str(srs_files[0]))
                    srs_type    = srs_meta.get("type", "")
                    sample_size = srs_meta.get("size")
                    sample_crc  = srs_meta.get("crc")
                    sample_name = srs_meta.get("name")

                    # If a file matching the sample's exact size is already in the
                    # content folder, it may BE the sample — verify CRC and copy.
                    sample_done = False
                    if sample_size:
                        for cand in Path(content_dir).rglob("*"):
                            if not (cand.is_file() and cand.stat().st_size == sample_size):
                                continue
                            crc = self._crc32_file(str(cand))
                            if sample_crc is not None and crc == sample_crc:
                                sample_dir = out_root / "Sample"
                                sample_dir.mkdir(parents=True, exist_ok=True)
                                dest = sample_dir / (sample_name or cand.name)
                                if not dest.exists():
                                    shutil.copy2(str(cand), str(dest))
                                self._log(
                                    f"  Sample already in content folder — CRC verified "
                                    f"({crc:08X}) ✓ copied as {dest.name}", "ok",
                                )
                                sample_done = True
                                summary["sample"] = "verified ✓"
                                break

                    # Prefer video/stream files over disc images — and never use a
                    # file that IS the sample as the rebuild source.
                    stream_media = [
                        f for f in Path(content_dir).rglob("*")
                        if f.is_file() and f.suffix.lower() in _STREAM_EXTS
                        and f.stat().st_size != sample_size
                    ]
                    disc_media = [
                        f for f in Path(content_dir).iterdir()
                        if f.is_file() and f.suffix.lower() in _DISC_EXTS
                    ]
                    if sample_done:
                        stream_media, disc_media = [], []
                    media_file = (sorted(stream_media, key=lambda f: f.stat().st_size, reverse=True) or disc_media or [None])[0]
                    _iso_tmp: Path | None = None
                    if media_file and media_file.suffix.lower() in _DISC_EXTS:
                        if srs_type == "STREAM":
                            # STREAM SRS = raw byte matching. A scan of the ISO
                            # covers every stream on the disc — extracting an
                            # M2TS first adds nothing, so use the ISO directly.
                            self._log(
                                "  STREAM-type SRS — scanning the ISO directly "
                                "(covers all streams on the disc)", "dim",
                            )
                        elif extract_iso_m2ts:
                            _iso_tmp = out_root / "_iso_m2ts_tmp"
                            ext_result = self._extract_m2ts_from_iso(
                                str(media_file), str(_iso_tmp)
                            )
                            if ext_result["ok"]:
                                media_file = Path(ext_result["path"])
                            else:
                                self._log(f"  ISO extract FAILED: {ext_result['error']}", "err")
                                media_file = None
                        else:
                            self._log(
                                f"  Sample skipped — only a disc image ({media_file.suffix}) found; "
                                "tick 'Extract M2TS from ISO' to auto-extract the Blu-ray stream "
                                "(requires disk space ≈ size of the main video stream).",
                                "dim",
                            )
                            media_file = None
                    if media_file:
                        self._log("  Creating sample…", "dim")
                        # Multi-CD releases: the SRS stores byte offsets into the
                        # exact CD the sample was cut from. Wrong CD → rebuild
                        # completes with the right size but wrong CRC (pyReScene
                        # mislabels this "LOL xvid issue"). Try each stream file.
                        attempts = [media_file] + [
                            f for f in stream_media[:4] if f != media_file
                        ]
                        samp = None
                        for i, mf in enumerate(attempts):
                            if i:
                                self._log(
                                    f"  Sample source mismatch — retrying with "
                                    f"{mf.name}…", "dim",
                                )
                            samp = self._srs_create_sample(
                                str(srs_files[0]), str(mf),
                                str(out_root / "Sample"),
                            )
                            if samp["ok"]:
                                break
                            err_l = samp.get("error", "").lower()
                            if not ("rebuild failed" in err_l or "signature" in err_l
                                    or "extract correct amount" in err_l):
                                break  # non-retryable error
                        # Fallback: the sample may have been cut from a different
                        # stream on the disc — one raw scan of the ISO covers all
                        # of its M2TS files at once.
                        if (not samp["ok"] and _iso_tmp and disc_media
                                and "signature" in samp.get("error", "").lower()):
                            self._log(
                                "  Not found in main stream — retrying against the full ISO "
                                "(covers all streams on the disc)…", "dim",
                            )
                            samp = self._srs_create_sample(
                                str(srs_files[0]), str(disc_media[0]),
                                str(out_root / "Sample"),
                            )
                        if samp["ok"]:
                            files_str = ", ".join(samp.get("files", []))
                            self._log(f"  Sample created ✓ — {files_str}", "ok")
                            summary["sample"] = "created ✓"
                        elif "signature" in samp.get("error", "").lower() and disc_media:
                            self._log(
                                "  Sample FAILED: the sample's bytes are not on the disc. "
                                "The group remuxed the sample when cutting it (the SRS "
                                "signature is the remux tool's own header bytes), so this "
                                "sample cannot be rebuilt from the ISO by any byte-matching "
                                "tool. The SRS can still verify an existing sample file.", "err",
                            )
                            summary["sample"] = "unrebuildable (BD remux)"
                            if non_scene_sample:
                                self._log(
                                    "  Creating NON-SCENE preview clip from the disc's "
                                    "main stream instead…", "dim",
                                )
                                carve = self._carve_nonscene_sample(
                                    str(disc_media[0]), str(out_root / "Sample"),
                                    release.lower(), sample_size,
                                )
                                if carve["ok"]:
                                    self._log(
                                        f"  Non-scene preview created — {carve['name']} "
                                        f"({carve['size']:,} B). NOT a scene file; it will "
                                        "never CRC-match the SRS.", "ok",
                                    )
                                    summary["sample"] += " + NONSCENE preview"
                                else:
                                    self._log(
                                        f"  Non-scene preview failed: {carve['error']}", "warn",
                                    )
                        elif "signature" in samp.get("error", "").lower():
                            m_trk = re.search(r"track (\d+)", samp["error"], re.IGNORECASE)
                            trk = int(m_trk.group(1)) if m_trk else 1
                            if trk > 1:
                                self._log(
                                    f"  Sample FAILED: track {trk} of the sample does not "
                                    "exist in the main video — the group's sample contains "
                                    "extra/re-encoded data, so it cannot be rebuilt "
                                    "byte-perfect from the movie file.", "err",
                                )
                                summary["sample"] = "unrebuildable (extra track)"
                            else:
                                self._log(
                                    "  Sample FAILED: the sample's video data was not found "
                                    "in this file — the content may not match this exact "
                                    "release (wrong source or re-encode).", "err",
                                )
                                summary["sample"] = "failed (content mismatch)"
                        else:
                            self._log(f"  Sample FAILED: {samp['error']}", "err")
                            summary["sample"] = "failed"
                    elif not sample_done and not stream_media and not disc_media:
                        self._log("  No media file for sample", "dim")
                        summary["sample"] = "no media file"
                    # Clean up extracted M2TS
                    if _iso_tmp and _iso_tmp.exists():
                        shutil.rmtree(str(_iso_tmp), ignore_errors=True)
                        self._log("  Cleaned up temp M2TS extract", "dim")
                else:
                    txt_placeholders = (list(stored_dir.rglob("*sample*.txt"))
                                        if stored_dir.exists() else [])
                    if txt_placeholders:
                        self._log(
                            "  No SRS — this usenet-sourced SRR stores only a sample "
                            "info .txt placeholder, so the sample cannot be rebuilt", "dim",
                        )
                        summary["sample"] = "no SRS (usenet SRR)"
                    else:
                        self._log("  No SRS in stored files", "dim")
                        summary["sample"] = "no SRS"

            # 5 — Move stored extras (Proof/, Subs/, jpgs …) into the release
            # folder, preserving their relative paths, then drop empty _stored.
            # Runs after sample creation, which reads the SRS from _stored.
            if stored_dir.exists():
                moved = 0
                for f in sorted(stored_dir.rglob("*")):
                    if not f.is_file():
                        continue
                    dest_f = out_root / f.relative_to(stored_dir)
                    if dest_f.exists():
                        f.unlink()  # duplicate — NFO/SFV were already copied up
                        continue
                    # rescene files an extra that lived INSIDE an archive under a
                    # folder named after that archive (_stored/xms-wcre.rar/…).
                    # At the release root that name is the produced VOLUME, so
                    # mkdir would collide with it (WinError 183). Such a file is
                    # already inside the rebuilt archive — leave it in _stored
                    # rather than crashing the job over it.
                    clash = next((p for p in dest_f.parents
                                  if p != out_root and p.is_file()), None)
                    if clash is not None:
                        self._log(f"  Kept {f.name} in _stored — its path "
                                  f"collides with the rebuilt {clash.name} "
                                  "(it is packed inside that archive).", "dim")
                        continue
                    dest_f.parent.mkdir(parents=True, exist_ok=True)
                    shutil.move(str(f), str(dest_f))
                    moved += 1
                for d in sorted((p for p in stored_dir.rglob("*") if p.is_dir()),
                                reverse=True):
                    try:
                        d.rmdir()
                    except OSError:
                        pass
                try:
                    stored_dir.rmdir()
                except OSError:
                    pass
                if moved:
                    self._log(f"  Moved {moved} stored extra(s) into release folder", "dim")

            # 6 — Verify the Subs folder actually contains what its SFV expects.
            # If the subs RAR(s) could not be rebuilt, the folder holds only
            # metadata — rename it so the gap is impossible to miss.
            subs_dir = out_root / "Subs"
            if subs_dir.is_dir():
                expected = []
                for s in list(subs_dir.glob("*.sfv")) + list(subs_dir.glob("*.SFV")):
                    expected += [fn for fn, _ in self._parse_sfv(str(s))]
                if expected:
                    missing_subs = [fn for fn in expected
                                    if not (subs_dir / fn).exists()]
                    subs_ok = not missing_subs
                else:
                    subs_ok = any(
                        re.search(r"\.(rar|r\d\d|\d{3})$", f.name, re.IGNORECASE)
                        for f in subs_dir.iterdir() if f.is_file()
                    )
                if subs_ok:
                    summary["subs"] = "✓"
                else:
                    summary["subs"] = "NOT produced"
                    target = out_root / "Subs_NOT_PRODUCED"
                    try:
                        if target.exists():
                            # previous run already renamed — merge contents
                            for f in subs_dir.iterdir():
                                dst = target / f.name
                                if not dst.exists():
                                    shutil.move(str(f), str(dst))
                            subs_dir.rmdir()
                        else:
                            subs_dir.rename(target)
                        self._log(
                            "  ⚠ Subs NOT produced — subtitle sources missing; "
                            "folder renamed to Subs_NOT_PRODUCED", "warn",
                        )
                    except OSError as e:
                        self._log(f"  ⚠ Subs NOT produced (rename failed: {e})", "warn")

            # 7 — Optionally delete the SOURCE folder (old jpg/nfo included).
            # ONLY after the produced volumes have been CRC-verified against the
            # SFV with zero bad and zero missing — the content then genuinely
            # lives inside byte-exact archives. A near-miss (bad last volume) or
            # an incomplete rebuild keeps the source, so a false "success" can
            # never delete the only copy of the source.
            _v = summary.get("verified") or {}
            fully_verified = bool(_v.get("checked")) and not _v.get("bad") and not _v.get("missing")
            # Every path below states its reason. The old chain had branches
            # that fell through logging NOTHING (rars == 0; a queue_path that
            # is no longer a directory; ok False while fully_verified True), so
            # a source that survived looked like a bug with no evidence for it.
            # An unexplained skip is indistinguishable from a broken delete.
            if not delete_source:
                pass
            elif (summary.get("rars") or 0) <= 0:
                self._log("  Source delete SKIPPED — the rebuild produced no RAR "
                          "volumes, so there is nothing proven to keep instead.",
                          "warn")
            elif not summary.get("ok"):
                self._log("  Source delete SKIPPED — the job did not finish "
                          "cleanly; keeping the source.", "warn")
            elif not fully_verified:
                why = ("no SFV entry was checkable" if not _v.get("checked")
                       else f"{len(_v.get('bad') or [])} volume(s) wrong, "
                            f"{len(_v.get('missing') or [])} missing")
                self._log("  Source delete SKIPPED — rebuild not fully "
                          f"CRC-verified against the SFV ({why}); keeping the "
                          "source.", "warn")
            elif not (queue_path and Path(queue_path).is_dir()):
                self._log(f"  Source delete SKIPPED — source folder is no longer "
                          f"a directory: {queue_path!r}", "warn")
            else:
                try:
                    qp = Path(queue_path).resolve()
                    orp = out_root.resolve()
                    if qp == qp.parent:
                        self._log("  Source delete SKIPPED — refusing to remove "
                                  "a drive root", "warn")
                    elif orp == qp or qp in orp.parents or orp in qp.parents:
                        self._log("  Source delete SKIPPED — source and output "
                                  "folders overlap", "warn")
                    else:
                        left = self._force_rmtree(qp)
                        if not left:
                            self._log(f"  Source folder DELETED (option enabled): "
                                      f"{qp}", "warn")
                        else:
                            self._log(
                                f"  ⚠ Source delete INCOMPLETE — {len(left)} "
                                f"file(s) survived in {qp} (locked or in use). "
                                "e.g. " + ", ".join(Path(x).name for x in left[:4]),
                                "err")
                except Exception as e:
                    self._log(f"  Source delete failed: {e}", "err")

            if summary["ok"]:
                self._log(f"  ✓ Done → {out_root}", "ok")
            else:
                self._log(f"  ⚠ Finished with errors → {out_root}", "warn")
            self._emit("job_done", {"release": release, "content_dir": queue_path,
                                    "ok": summary["ok"]})

        except (ValueError, RuntimeError) as e:
            summary["ok"] = False
            if self._stop.is_set() or self._skip.is_set():
                # rar.exe was killed / spawns refused by a user stop or skip.
                self._log("  ⏹ Release aborted by user.", "warn")
                summary["note"] = "stopped"
                self._emit("job_done", {"release": release, "content_dir": queue_path,
                                        "ok": False, "stopped": True})
            else:
                summary["note"] = summary["note"] or str(e) or "failed"
                jd = {"release": release, "content_dir": queue_path, "ok": False,
                      "note": summary["note"]}
                if summary.get("metadata_only"):
                    jd["metadata_only"] = summary["metadata_only"]
                self._emit("job_done", jd)
        except InterruptedError:
            self._log("  Stopped.", "warn")
            summary["ok"] = False
            summary["note"] = "stopped"
            self._emit("job_done", {"release": release, "content_dir": queue_path, "ok": False, "stopped": True})
        except Exception as e:
            msg = str(e)
            if ((("already exists" in msg.lower())
                 or getattr(e, "winerror", None) == 183)
                    and not summary.get("ok")):
                # Stale-output collision (Windows WinError 183): a leftover from an
                # interrupted prior run blocked a fresh write inside rescene. Clear
                # the PRODUCED artifacts (never _stored / the cached SRR, so a
                # re-run stays offline) and say so plainly instead of surfacing the
                # raw error — a re-run then starts clean.
                #
                # NEVER when the rebuild already succeeded: a name clash raised by
                # a POST-rebuild step (Winx_Club…EXiMiUS — moving the stored extra
                # `_stored/xms-wcre.rar/xms-wcre.jpg` tried to mkdir a DIRECTORY
                # over the freshly produced volume of the same name) was landing
                # here and deleting 19 CRC-verified volumes. A verified rebuild is
                # never "stale output".
                try:
                    if "out_root" in locals() and out_root.is_dir():
                        self._clear_produced_volumes(out_root)
                        for d in ("Sample", "Subs", "Subs_NOT_PRODUCED",
                                  "_subs_tmp", "_iso_m2ts_tmp"):
                            shutil.rmtree(str(out_root / d), ignore_errors=True)
                except Exception:
                    pass
                self._log("  ⚠ Stale-output collision (leftover from a prior run) "
                          "— cleared the produced files; re-run this release and "
                          "it should proceed.", "warn")
                summary["note"] = "stale output collision — cleared, re-run"
                summary["ok"] = False
            elif summary.get("ok") and (summary.get("rars") or 0) > 0:
                # The rebuild finished and VERIFIED; only a tidy-up step after it
                # failed. Keep the release — report the blemish, don't discard
                # good volumes over it.
                self._log(f"  ⚠ Rebuild succeeded, but a post-rebuild step "
                          f"failed: {e}", "warn")
                summary["note"] = f"rebuilt; post-step failed: {str(e)[:70]}"
                self._emit("job_done", {"release": release,
                                        "content_dir": queue_path, "ok": True})
                return summary
            else:
                self._log(f"  FAILED: {e}", "err")
                summary["note"] = str(e)[:100]
                summary["ok"] = False
            self._emit("job_done", {"release": release, "content_dir": queue_path, "ok": False})
        finally:
            # A skip only aborts THIS release — clear it so the next queued job
            # runs normally. (_stop stays set on a hard stop: the batch loop
            # checks it and halts.)
            self._skip.clear()
            # Failed jobs leave no folder behind: remove the output release dir
            # when the job failed AND it holds nothing substantial — extras,
            # the cached SRR and stub RAR headers only (every non-metadata file
            # ≤ 64 KB). A folder from an earlier successful run has full-size
            # volumes and is never touched. User-stopped jobs are kept.
            try:
                if (not summary["ok"] and summary.get("note") != "stopped"
                        and "out_root" in locals() and out_root.is_dir()):
                    substantial = any(
                        f.stat().st_size > 65536
                        for f in out_root.rglob("*")
                        if f.is_file() and f.suffix.lower() not in META_EXTS
                    )
                    if not substantial:
                        shutil.rmtree(str(out_root), ignore_errors=True)
                        self._log("  Removed failed output folder", "dim")
            except Exception:
                pass
            # Prune empty directories: a failed job creates the release folder
            # before writing anything; failed sample attempts leave an empty
            # Sample/. rmdir only removes empty dirs, so content is never at risk.
            try:
                if "out_root" in locals() and out_root.is_dir():
                    for d in sorted((p for p in out_root.rglob("*") if p.is_dir()),
                                    reverse=True):
                        try:
                            d.rmdir()
                        except OSError:
                            pass
                    try:
                        out_root.rmdir()
                    except OSError:
                        pass
            except Exception:
                pass
            if not batch_mode:
                self._running = False
                self._emit("status", {"state": "done"})
        self._record_result(summary, queue_path)
        return summary

    # ── Batch processing ──────────────────────────────────────────────────────

    def prefetch_srrs(self, jobs: list) -> bool:
        """Resolve every queued release and download its SRR to the cache path
        NOW — so the later Process run needs no srrdb access at all (rebuilding
        is 100% local). Lets you run the whole network phase in one short online
        window, then rebuild offline. Ambiguous releases are pinned by exact
        content CRC (one API call), never by downloading all candidates."""
        if self._running:
            self._log("Already running.", "warn")
            return False
        self._stop.clear()
        self._skip.clear()
        threading.Thread(target=self._prefetch_thread, args=(jobs,), daemon=True).start()
        return True

    def _prefetch_thread(self, jobs: list):
        self._running = True
        self._emit("status", {"state": "running"})
        self._log(f"Prefetching SRRs for {len(jobs)} release(s) — stay online "
                  "until this finishes…", "info")
        ready = failed = 0
        for i, job in enumerate(jobs):
            if self._stop.is_set():
                break
            path = (job.get("content_dir") or "").strip()
            dest_dir = (job.get("dest_dir") or "").strip()
            release = (job.get("release_name") or "").strip()
            candidates = job.get("candidates", [])
            max_test = int(job.get("max_test", 5))
            name = Path(path).name if path else "?"
            self._emit("batch_progress", {"current": i + 1, "total": len(jobs)})

            # Resolve to a single release name (network-dependent) --------------
            content_dir = path
            try:
                if content_dir and Path(content_dir).is_dir():
                    children = list(Path(content_dir).iterdir())
                    if (children and len(children) == 1 and children[0].is_dir()):
                        content_dir = str(children[0])
            except Exception:
                pass

            resolved = release
            if not resolved and content_dir and Path(content_dir).is_dir():
                # exact content CRC first — one call, name-independent
                hres = self.search_by_content_hash(content_dir)
                if hres.get("ok") and hres.get("results"):
                    resolved = hres["results"][0]["release"]
                    self._log(f"  [{i+1}] {name}: CRC → {resolved}", "dim")
                elif candidates:
                    best = self.find_best_match(candidates, content_dir, max_test)
                    if best.get("release") and best.get("score", 0) >= 0.5:
                        resolved = best["release"]
                        self._log(f"  [{i+1}] {name}: matched → {resolved} "
                                  f"({best['score']:.0%})", "dim")

            if not resolved:
                self._log(f"  [{i+1}] {name}: could not resolve — skipped", "warn")
                self._emit("prefetch_done", {"content_dir": path, "ok": False})
                failed += 1
                continue

            # Download the SRR to the exact path Process expects ---------------
            if not dest_dir:
                self._log(f"  [{i+1}] {resolved}: no output folder set", "err")
                self._emit("prefetch_done", {"content_dir": path, "ok": False})
                failed += 1
                continue
            out_root = Path(dest_dir) / resolved
            existing = list(out_root.glob("*.srr"))
            if existing:
                self._log(f"  [{i+1}] {resolved}: SRR already cached", "dim")
                self._emit("prefetch_done", {"content_dir": path, "release": resolved,
                                             "ok": True})
                ready += 1
                continue
            dl = self.download_srr(resolved, str(out_root))
            if dl.get("ok"):
                src = ("local cache — no srrdb request" if dl.get("cached")
                       else "downloaded")
                self._log(f"  [{i+1}] {resolved}: SRR {src} "
                          f"({dl.get('size', 0):,} B) ✓", "ok")
                self._emit("prefetch_done", {"content_dir": path, "release": resolved,
                                             "ok": True})
                ready += 1
            else:
                self._log(f"  [{i+1}] {resolved}: SRR download failed — "
                          f"{dl.get('error', '?')}", "err")
                self._emit("prefetch_done", {"content_dir": path, "ok": False})
                failed += 1

        self._log(f"Prefetch complete — {ready} ready, {failed} unresolved. "
                  "The network phase is done; Process rebuilds entirely offline.",
                  "ok" if not failed else "warn")
        self._running = False
        self._emit("status", {"state": "done"})

    def start_batch(self, jobs: list) -> bool:
        if self._running:
            self._log("Already running.", "warn")
            return False
        self._stop.clear()
        self._skip.clear()
        self._fresh_all = False   # "always fresh" latches per RUN, not forever
        threading.Thread(target=self._batch_thread, args=(jobs,), daemon=True).start()
        return True

    def _batch_thread(self, jobs: list):
        self._running = True
        self._emit("status", {"state": "running"})
        total = len(jobs)
        results = []
        for i, job in enumerate(jobs):
            if self._stop.is_set():
                break
            self._emit("batch_progress", {"current": i + 1, "total": total})
            self._log(f"\n[{i + 1}/{total}]", "dim")
            results.append(self._process_one(job, batch_mode=True))

        # Metadata-only releases (DIRFIX/NFOFIX) have nothing to rebuild — don't
        # count them as failures against the success ratio; list them apart.
        na = [r for r in results if r.get("metadata_only")]
        empties = [r for r in results
                   if r.get("empty") and not r.get("metadata_only")]
        walls = [r for r in results if r.get("wall_skipped")
                 and not r.get("metadata_only") and not r.get("empty")]
        rebuildable = [r for r in results
                       if not r.get("metadata_only") and not r.get("wall_skipped")
                       and not r.get("empty")]
        ok_n = sum(1 for r in rebuildable if r["ok"])
        tail = f"  ({len(na)} n/a — metadata only)" if na else ""
        if walls:
            tail += f"  ({len(walls)} skipped — known version-walls)"
        if empties:
            tail += f"  ({len(empties)} empty folders)"
        self._log(f"\n══ Batch summary — {ok_n}/{len(rebuildable)} succeeded{tail} ══",
                  "info")
        for r in results:
            if r.get("metadata_only"):
                self._log(f"  ⊘ {r['release']} — "
                          f"{r.get('note') or 'metadata only, nothing to rebuild'}",
                          "dim")
            elif r.get("empty"):
                self._log(f"  ⊘ {r['release']} — empty folder (no content files)",
                          "dim")
            elif r.get("wall_skipped"):
                self._log(f"  ⏭ {r['release']} — known version-wall, skipped "
                          "(add WinRAR versions to retry)", "dim")
            elif r["ok"]:
                parts = []
                if r["rars"] is not None:
                    parts.append(f"{r['rars']} RARs")
                if r["sample"]:
                    parts.append(f"sample: {r['sample']}")
                if r.get("subs"):
                    parts.append(f"subs: {r['subs']}")
                if r["note"]:
                    parts.append(r["note"])
                cls = "warn" if r.get("subs") == "NOT produced" else "ok"
                self._log(f"  ✓ {r['release']} — {', '.join(parts) or 'done'}", cls)
            else:
                detail = r["note"] or "failed"
                if r["sample"]:
                    detail += f" (sample: {r['sample']})"
                if r.get("subs") == "NOT produced":
                    detail += " (subs: NOT produced)"
                self._log(f"  ✗ {r['release']} — {detail}", "err")
        self._running = False
        self._emit("status", {"state": "done"})
